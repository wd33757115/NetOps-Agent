from io import BytesIO
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, Cm
from docx.oxml.ns import qn
from src.apps.official_documents.backend.writing.docx_format import setup_a4_margins, apply_fixed_line_spacing
from .schemas import Content


def render_workspace_document_bytes(content: dict, *, draft=False, template=None) -> bytes:
    from .templates import resolve_template
    data = Content.model_validate(content)
    fmt = (template or resolve_template(data.doc_type))["format"]
    document = Document()
    setup_a4_margins(document)
    for section in document.sections:
        section.page_width = Cm(fmt["page_width_mm"] / 10)
        section.page_height = Cm(fmt["page_height_mm"] / 10)
        for edge in ("top", "right", "bottom", "left"):
            setattr(section, edge + "_margin", Cm(fmt[f"margin_{edge}_mm"] / 10))
    if draft:
        p = document.sections[0].header.paragraphs[0]
        p.text = "草稿 · 仅供核对"
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for block in data.blocks:
        if not any(part.text.strip() for part in block.runs):
            continue
        p = document.add_paragraph()
        apply_fixed_line_spacing(p)
        p.paragraph_format.line_spacing = Pt(fmt["line_height_pt"])
        font, size = fmt["body_font"], fmt["body_size_pt"]
        if block.role in ("title", "issuer"):
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            font, size = fmt["title_font"], fmt["title_size_pt"]
            p.paragraph_format.space_after = Pt(fmt["line_height_pt"] if block.role == "title" else 0)
        elif block.role.startswith("signature_"):
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        elif block.role != "main_recipient":
            p.paragraph_format.first_line_indent = Pt(fmt["body_size_pt"] * fmt["indent_chars"])
        if block.role == "section_heading":
            font = fmt["heading1_font"] if block.level == 1 else fmt["heading2_font"] if block.level == 2 else fmt["body_font"]
            p.paragraph_format.keep_with_next = True
        for part in block.runs:
            run = p.add_run(part.text)
            run.font.name = font
            run.font.size = Pt(size)
            run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), font)
            run.bold = "bold" in part.marks
            run.underline = "underline" in part.marks
    output = BytesIO()
    document.save(output)
    return output.getvalue()
