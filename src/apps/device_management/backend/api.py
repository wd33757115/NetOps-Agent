"""Device Management API."""

from __future__ import annotations

import base64
import binascii
import csv
import io
import json
import logging
import uuid
from datetime import datetime, timezone
from typing import Any

from fastapi import APIRouter, Depends, HTTPException, Query, Request
from pydantic import BaseModel, Field
from sqlalchemy import func, or_

from src.auth.dependencies import require_netops_access
from src.auth.models import CurrentUser

from src.core.collection import CollectionPolicyResolver, ExecutionPlanRepository
from src.platform.inventory.credentials import CredentialResolver
from src.platform.inventory.identity_discovery import (
    discover_device_identity,
    identity_evidence,
    platform_for_identity,
)
from src.platform.inventory.readiness_projection import (
    DeviceReadinessProjectionRepository,
    readiness_source_digest,
)
from src.platform.inventory.repository import DeviceRepository, normalize_device_status
from src.core.evidence import evidence_id_for_command_snapshot
from src.core.patrol.batch_manifest import PatrolBatchManifestRepository
from src.core.patrol.handlers import StartPatrolCommand, StartPatrolHandler
from src.core.patrol.modes import PRODUCTION_MODE, normalize_patrol_mode
from src.core.patrol.production_readiness import PatrolProductionReadinessService
from src.gateway.audit_service import write_audit_log
from src.infrastructure.db.models import (
    CollectionExecutionPlanRecord,
    CommandSnapshot,
    NetworkChangeRecord,
    NetworkEventOccurrenceRecord,
    NetworkEventRecord,
    PatrolCollectionResult,
    PatrolRun,
    RawEvidenceRecord,
    StateObservationRecord,
    WorkflowRun,
    WorkflowStepRecord,
)
from src.infrastructure.db.postgres import get_db_session
from src.observability.read_metrics import observe_read

logger = logging.getLogger(__name__)


def _enqueue_patrol_batch_dispatch(patrol_batch_id: str) -> str:
    from src.core.celery_tasks.tasks import dispatch_patrol_batch_task

    task = dispatch_patrol_batch_task.apply_async(
        args=[patrol_batch_id],
        queue="netops.patrol",
    )
    return str(task.id)

async def device_access(request: Request):
    from src.auth.netops_access import requires_netops_configuration
    from src.apps.device_management.backend.access import app_user
    if request.url.path.startswith("/api/apps/device-management/") or requires_netops_configuration(request.method, request.url.path):
        return await app_user(request)
    from src.apps.inspection.backend.access import app_user as inspection_user
    try:
        return await inspection_user(request)
    except HTTPException as exc:
        if exc.status_code != 403 or "/patrol/" in request.url.path or request.url.path.endswith("/effective-plan"):
            raise
        return await app_user(request)


router = APIRouter(
    prefix="/api/v1/devices",
    tags=["Device Management"],
    dependencies=[Depends(device_access)],
)


class DevicePlanPreviewRequest(BaseModel):
    action: str = "patrol"
    mode: str = "production"


class DeviceUpsertRequest(BaseModel):
    device_name: str = Field(..., min_length=1)
    management_ip: str = Field(..., min_length=1)
    username: str | None = None
    credential_ref: str | None = None
    vendor: str | None = None
    model: str | None = None
    family: str | None = None
    platform: str | None = None
    device_type: str | None = None
    os_version: str | None = None
    serial_number: str | None = None
    site: str | None = None
    role: str | None = None
    status: str = "active"
    output_path: str | None = None
    command_profile: str | None = None
    attributes: dict[str, Any] | None = None
    groups: list[str] | None = None


class DeviceUpdateRequest(BaseModel):
    device_name: str | None = None
    management_ip: str | None = None
    username: str | None = None
    credential_ref: str | None = None
    vendor: str | None = None
    model: str | None = None
    family: str | None = None
    platform: str | None = None
    device_type: str | None = None
    os_version: str | None = None
    serial_number: str | None = None
    site: str | None = None
    role: str | None = None
    status: str | None = None
    output_path: str | None = None
    command_profile: str | None = None
    attributes: dict[str, Any] | None = None
    groups: list[str] | None = None


class DeviceStatusRequest(BaseModel):
    status: str = Field(..., min_length=1)


class DeviceGroupsRequest(BaseModel):
    groups: list[str] = Field(default_factory=list)


class DeviceDynamicGroupRequest(BaseModel):
    name: str = Field(..., min_length=1)
    selector: dict[str, Any]
    description: str | None = None
    status: str = "active"


class CredentialInspectRequest(BaseModel):
    credential_ref: str | None = None


class DeviceImportRequest(BaseModel):
    content: str = Field(..., min_length=1)
    format: str = "csv"
    filename: str | None = None
    dry_run: bool = True


class DeviceBulkFilters(BaseModel):
    device_id: str | None = None
    device_name: str | None = None
    ip: str | None = None
    group: str | None = None
    vendor: str | None = None
    model: str | None = None
    family: str | None = None
    platform: str | None = None
    device_type: str | None = None
    site: str | None = None
    role: str | None = None
    status: str | None = None
    connection_status: str | None = None


class DeviceBulkUpdate(BaseModel):
    command_profile: str | None = None
    credential_ref: str | None = None
    status: str | None = None
    platform: str | None = None
    device_type: str | None = None
    site: str | None = None
    role: str | None = None
    groups: list[str] | None = None
    group_mode: str = "replace"


class DeviceBulkRequest(BaseModel):
    device_ids: list[str] | None = None
    dynamic_group_id: str | None = None
    filters: DeviceBulkFilters | None = None
    update: DeviceBulkUpdate


class DeviceBulkSelectionRequest(BaseModel):
    device_ids: list[str] | None = None
    dynamic_group_id: str | None = None
    filters: DeviceBulkFilters | None = None


class DevicePatrolLaunchRequest(BaseModel):
    device_ids: list[str] | None = None
    dynamic_group_id: str | None = None
    filters: DeviceBulkFilters | None = None
    mode: str = "production"
    save_baseline: bool = False
    ticket_id: str | None = None
    launch_mode: str = "direct"
    request_id: str | None = Field(default=None, min_length=1, max_length=128)
    workflow_name: str = "patrol-change-event"
    scope_contract: dict[str, Any] | None = None
    scope_exclusions: dict[str, dict[str, Any]] | None = None
    scope_not_applicable: dict[str, dict[str, Any]] | None = None


DEVICE_IMPORT_COLUMNS: list[dict[str, Any]] = [
    {
        "name": "device_name",
        "required": True,
        "aliases": ["name", "hostname"],
        "description": "设备名称，建议使用网管系统中的唯一名称。",
        "example": "SW1",
    },
    {
        "name": "management_ip",
        "required": True,
        "aliases": ["ip", "mgmt_ip", "management_address"],
        "description": "管理地址，用于连接设备和幂等匹配。",
        "example": "10.0.0.1",
    },
    {
        "name": "vendor",
        "required": False,
        "aliases": [],
        "description": "厂商，例如 Cisco、H3C。",
        "example": "Cisco",
    },
    {
        "name": "model",
        "required": False,
        "aliases": [],
        "description": "精确型号。",
        "example": "WS-C2960XR-24TS-I",
    },
    {
        "name": "family",
        "required": False,
        "aliases": [],
        "description": "设备族。",
        "example": "IOS-C2960",
    },
    {
        "name": "platform",
        "required": False,
        "aliases": [],
        "description": "采集平台类型。",
        "example": "ios",
    },
    {
        "name": "device_type",
        "required": False,
        "aliases": ["type"],
        "description": "设备类型，例如 router、access_switch、firewall。",
        "example": "access_switch",
    },
    {
        "name": "os_version",
        "required": False,
        "aliases": [],
        "description": "系统版本。",
        "example": "15.2(7)E10",
    },
    {
        "name": "serial_number",
        "required": False,
        "aliases": ["serial"],
        "description": "序列号。",
        "example": "FOC1234",
    },
    {
        "name": "site",
        "required": False,
        "aliases": [],
        "description": "站点或机房。",
        "example": "xian",
    },
    {
        "name": "role",
        "required": False,
        "aliases": [],
        "description": "设备角色。",
        "example": "access",
    },
    {
        "name": "status",
        "required": False,
        "aliases": [],
        "description": "设备状态，active=启用，disabled=停用。留空默认 active；旧值 inactive 会自动转换为 disabled。",
        "example": "active",
    },
    {
        "name": "credential_ref",
        "required": False,
        "aliases": [],
        "description": "凭据引用，例如 env:NETOPS_PASSWORD。",
        "example": "env:NETOPS_PASSWORD",
    },
    {
        "name": "command_profile",
        "required": False,
        "aliases": [],
        "description": "采集命令集名称或版本引用。",
        "example": "browser-test-patrol",
    },
    {
        "name": "output_path",
        "required": False,
        "aliases": [],
        "description": "设备原始巡检输出目录，留空使用系统默认位置。",
        "example": "",
    },
    {
        "name": "groups",
        "required": False,
        "aliases": [],
        "description": "逗号分隔的设备分组。",
        "example": "core,xian",
    },
    {
        "name": "attributes_json",
        "required": False,
        "aliases": [],
        "description": "扩展属性 JSON 对象，用于采集参数、身份发现证据等高级配置。",
        "example": '{"collector_options":{"read_timeout":60}}',
    },
]

DEVICE_IMPORT_COLUMN_NAMES = [column["name"] for column in DEVICE_IMPORT_COLUMNS]
DEVICE_IMPORT_REQUIRED_COLUMNS = [
    column["name"] for column in DEVICE_IMPORT_COLUMNS if column["required"]
]
DEVICE_IMPORT_HEADER_ALIASES = {
    alias: column["name"]
    for column in DEVICE_IMPORT_COLUMNS
    for alias in [column["name"], *column.get("aliases", [])]
}


def _redact_device(device: dict[str, Any]) -> dict[str, Any]:
    data = dict(device)
    data.pop("password", None)
    return data


def _audit_device_action(
    request: Request | None,
    action: str,
    resource_type: str,
    resource_id: str | None,
    detail: dict[str, Any] | None = None,
) -> None:
    user = getattr(getattr(request, "state", None), "current_user", None)
    if user is not None and not isinstance(user, CurrentUser):
        user = None
    write_audit_log(
        action=action,
        user_id=user.user_id if user else None,
        username=user.username if user else None,
        resource_type=resource_type,
        resource_id=resource_id,
        detail=detail,
        ip_address=request.client.host if request and request.client else None,
    )


def _split_csv_list(value: str | None) -> list[str]:
    return [item.strip() for item in (value or "").split(",") if item.strip()]


def _parse_json_field(
    value: str | None, *, row_number: int, errors: list[str]
) -> dict[str, Any] | None:
    text = (value or "").strip()
    if not text:
        return None
    try:
        data = json.loads(text)
    except json.JSONDecodeError as exc:
        errors.append(f"第 {row_number} 行 attributes_json 不是合法 JSON：{exc.msg}")
        return None
    if not isinstance(data, dict):
        errors.append(f"第 {row_number} 行 attributes_json 必须是 JSON 对象")
        return None
    return data


def _build_import_row(
    normalized: dict[str, Any],
    *,
    row_number: int,
    row_errors: list[str] | None = None,
) -> dict[str, Any]:
    errors: list[str] = list(row_errors or [])
    management_ip = normalized.get("management_ip") or ""
    device_name = normalized.get("device_name") or ""
    if not device_name:
        errors.append(f"第 {row_number} 行缺少必填列 device_name")
    if not management_ip:
        errors.append(f"第 {row_number} 行缺少必填列 management_ip")
    try:
        status = normalize_device_status(normalized.get("status") or "active")
        if status == "deleted":
            errors.append(f"第 {row_number} 行 status 只能是 active 或 disabled")
    except ValueError:
        status = normalized.get("status") or "active"
        errors.append(f"第 {row_number} 行 status 只能是 active 或 disabled")
    attributes = _parse_json_field(
        normalized.get("attributes_json"), row_number=row_number, errors=errors
    )
    payload = {
        "device_name": device_name,
        "management_ip": management_ip,
        "username": normalized.get("username") or None,
        "credential_ref": normalized.get("credential_ref") or None,
        "vendor": normalized.get("vendor") or None,
        "model": normalized.get("model") or None,
        "family": normalized.get("family") or None,
        "platform": normalized.get("platform") or None,
        "device_type": normalized.get("device_type") or None,
        "os_version": normalized.get("os_version") or None,
        "serial_number": normalized.get("serial_number") or None,
        "site": normalized.get("site") or None,
        "role": normalized.get("role") or None,
        "status": status,
        "output_path": normalized.get("output_path") or None,
        "command_profile": normalized.get("command_profile") or None,
        "attributes": attributes,
        "groups": _split_csv_list(normalized.get("groups")),
    }
    return {
        "row_number": row_number,
        "valid": not errors,
        "errors": errors,
        "device": payload,
    }


def _canonical_import_headers(fieldnames: list[Any]) -> tuple[list[str], list[str]]:
    canonical_headers: list[str] = []
    errors: list[str] = []
    seen: set[str] = set()
    unknown: list[str] = []
    duplicates: list[str] = []
    for value in fieldnames:
        raw = str(value or "").strip()
        if not raw:
            continue
        key = raw.lower()
        canonical = DEVICE_IMPORT_HEADER_ALIASES.get(key)
        if canonical is None:
            unknown.append(raw)
            continue
        if canonical in seen:
            duplicates.append(raw)
            continue
        seen.add(canonical)
        canonical_headers.append(canonical)
    missing = [name for name in DEVICE_IMPORT_REQUIRED_COLUMNS if name not in seen]
    if missing:
        errors.append(f"缺少必填列：{', '.join(missing)}")
    if unknown:
        errors.append(f"存在未识别列：{', '.join(unknown)}")
    if duplicates:
        errors.append(f"存在重复列或别名冲突：{', '.join(duplicates)}")
    return canonical_headers, errors


def _raise_import_header_error(errors: list[str]) -> None:
    raise HTTPException(
        status_code=400,
        detail={
            "message": "设备导入表头校验失败",
            "errors": errors,
            "required_columns": DEVICE_IMPORT_REQUIRED_COLUMNS,
            "supported_columns": DEVICE_IMPORT_COLUMNS,
            "template": _device_import_template_payload(),
        },
    )


def _parse_csv_import_rows(content: str) -> list[dict[str, Any]]:
    reader = csv.DictReader(io.StringIO(content.strip()))
    if not reader.fieldnames:
        raise HTTPException(status_code=400, detail="CSV 必须包含表头")
    canonical_headers, header_errors = _canonical_import_headers(reader.fieldnames)
    if header_errors:
        _raise_import_header_error(header_errors)
    rows: list[dict[str, Any]] = []
    for index, raw in enumerate(reader, start=2):
        row_errors: list[str] = []
        if raw.get(None):
            row_errors.append(f"第 {index} 行列数超过表头，请检查逗号或引号")
        normalized = {
            DEVICE_IMPORT_HEADER_ALIASES.get(
                str(key or "").strip().lower(), str(key or "").strip().lower()
            ): (value.strip() if isinstance(value, str) else value)
            for key, value in raw.items()
            if key is not None
        }
        normalized = {key: normalized.get(key, "") for key in canonical_headers}
        rows.append(_build_import_row(normalized, row_number=index, row_errors=row_errors))
    return rows


def _parse_xlsx_import_rows(content: str) -> list[dict[str, Any]]:
    try:
        raw = base64.b64decode(content, validate=True)
    except (ValueError, binascii.Error) as exc:
        raise HTTPException(status_code=400, detail="xlsx content must be base64 encoded") from exc
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise HTTPException(status_code=500, detail="openpyxl is required for xlsx import") from exc
    try:
        workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    except Exception as exc:
        raise HTTPException(status_code=400, detail=f"invalid xlsx file: {exc}") from exc
    sheet = workbook.active
    rows_iter = sheet.iter_rows(values_only=True)
    try:
        header_row = next(rows_iter)
    except StopIteration as exc:
        raise HTTPException(status_code=400, detail="XLSX 必须包含表头") from exc
    headers = [str(value or "").strip() for value in header_row]
    if not any(headers):
        raise HTTPException(status_code=400, detail="XLSX 必须包含表头")
    canonical_headers, header_errors = _canonical_import_headers(headers)
    if header_errors:
        _raise_import_header_error(header_errors)
    rows: list[dict[str, Any]] = []
    for index, values in enumerate(rows_iter, start=2):
        if not any(value is not None and str(value).strip() for value in values):
            continue
        normalized: dict[str, Any] = {}
        row_errors: list[str] = []
        if len(values) > len(headers) and any(
            value is not None and str(value).strip() for value in values[len(headers) :]
        ):
            row_errors.append(f"第 {index} 行列数超过表头，请检查多余单元格")
        for header, value in zip(headers, values):
            canonical = DEVICE_IMPORT_HEADER_ALIASES.get(str(header or "").strip().lower())
            if canonical:
                normalized[canonical] = str(value).strip() if value is not None else ""
        normalized = {key: normalized.get(key, "") for key in canonical_headers}
        rows.append(_build_import_row(normalized, row_number=index, row_errors=row_errors))
    return rows


def _parse_import_rows(payload: DeviceImportRequest) -> list[dict[str, Any]]:
    import_format = (payload.format or "csv").strip().lower()
    if import_format == "csv":
        return _parse_csv_import_rows(payload.content)
    if import_format in {"xlsx", "excel"}:
        return _parse_xlsx_import_rows(payload.content)
    raise HTTPException(status_code=400, detail="only csv and xlsx import are supported")


def _device_import_template_payload() -> dict[str, Any]:
    sample_row = {column["name"]: column["example"] for column in DEVICE_IMPORT_COLUMNS}
    buffer = io.StringIO()
    writer = csv.DictWriter(buffer, fieldnames=DEVICE_IMPORT_COLUMN_NAMES, lineterminator="\n")
    writer.writeheader()
    writer.writerow(sample_row)
    return {
        "filename": "device-import-template.csv",
        "format": "csv",
        "columns": DEVICE_IMPORT_COLUMNS,
        "required_columns": DEVICE_IMPORT_REQUIRED_COLUMNS,
        "csv": buffer.getvalue(),
    }


def _preview_import(payload: DeviceImportRequest, repository: DeviceRepository) -> dict[str, Any]:
    rows = _parse_import_rows(payload)
    preview_rows: list[dict[str, Any]] = []
    create_count = 0
    update_count = 0
    invalid_count = 0
    for row in rows:
        existing = None
        if row["valid"]:
            device = row["device"]
            existing = repository.get_device_by_ip_or_name(device["management_ip"])
            if not existing:
                exact_name_matches = repository.get_devices_by_filter(
                    device_name=device["device_name"],
                    status=None,
                )
                existing = exact_name_matches[0] if exact_name_matches else None
            row["action"] = "update" if existing else "create"
            row["existing_device_id"] = existing.get("device_id") if existing else None
            if existing:
                update_count += 1
            else:
                create_count += 1
        else:
            row["action"] = "invalid"
            row["existing_device_id"] = None
            invalid_count += 1
        preview_rows.append(row)
    return {
        "success": True,
        "total": len(preview_rows),
        "create_count": create_count,
        "update_count": update_count,
        "invalid_count": invalid_count,
        "rows": preview_rows,
    }


def _resolve_bulk_devices(
    payload: DeviceBulkRequest | DeviceBulkSelectionRequest,
    repository: DeviceRepository,
    *,
    active_only: bool = False,
) -> list[dict[str, Any]]:
    devices_by_id: dict[str, dict[str, Any]] = {}
    for device_id in payload.device_ids or []:
        matches = repository.get_devices_by_filter(device_id=device_id, status=None)
        if matches and (not active_only or matches[0].get("status", "active") == "active"):
            devices_by_id[matches[0]["device_id"]] = matches[0]
    if payload.dynamic_group_id:
        group, group_devices = repository.resolve_dynamic_group_devices(payload.dynamic_group_id)
        if group is None:
            raise HTTPException(status_code=404, detail="dynamic group not found")
        if group.get("status") != "active":
            raise HTTPException(status_code=400, detail="dynamic group is not active")
        for device in group_devices:
            if not active_only or device.get("status", "active") == "active":
                devices_by_id[device["device_id"]] = device
    if payload.filters:
        filters = payload.filters.model_dump()
        if active_only:
            filters["status"] = "active"
        for device in repository.get_devices_by_filter(**filters):
            devices_by_id[device["device_id"]] = device
    return sorted(
        devices_by_id.values(),
        key=lambda item: (item.get("device_name") or "", item.get("ip") or ""),
    )


def _production_parser_coverage(
    device: dict[str, Any], commands: list[str]
) -> list[dict[str, Any]]:
    from src.core.device_ops.tools import PatrolTool

    return PatrolTool._production_coverage(device, commands)




def _plan_check_item_ids(plan: Any) -> tuple[str, ...]:
    return tuple(
        str(item)
        for item in (getattr(plan, "check_item_ids", ()) or ())
        if str(item)
    )


def _preview_check_item_effective_plan(
    *,
    device: dict[str, Any],
    plan: Any,
) -> dict[str, Any] | None:
    """Resolve the authoritative check-item plan behind a device preview."""

    check_item_ids = _plan_check_item_ids(plan)
    if not check_item_ids:
        return None

    from src.core.check_items.planning import (
        CheckItemEffectivePlanService,
        EffectivePlanRequest,
        build_inventory_snapshot,
        build_target_manifest,
    )

    device_id = str(device.get("device_id") or device.get("id") or "").strip()
    if not device_id:
        raise ValueError("check-item effective plan requires a device_id")
    profile_version_id = (
        str(getattr(plan, "profile_version_id", "") or "").strip() or None
    )
    selector = {"device_ids": [device_id]}
    target = build_target_manifest(devices=[device], selector=selector)
    inventory = build_inventory_snapshot([device])
    effective = CheckItemEffectivePlanService().preview(
        EffectivePlanRequest(
            patrol_batch_id=f"device-plan-preview-{uuid.uuid4()}",
            devices=(dict(device),),
            check_item_ids=check_item_ids,
            check_item_ids_by_device={device_id: check_item_ids},
            profile_version_id=profile_version_id,
            profile_version_ids_by_device={
                device_id: profile_version_id,
            },
            target_manifest_digest=str(target["digest"]),
            inventory_snapshot_digest=str(inventory["digest"]),
            current_observed_at=datetime.now(timezone.utc),
            actor="device-effective-plan-preview",
        )
    )
    return {
        "manifest_id": effective.manifest_id,
        "manifest_digest": effective.manifest_digest,
        "manifest": dict(effective.manifest),
        "decisions": [dict(item) for item in effective.decisions],
        "collection_actions": [
            dict(item) for item in effective.collection_actions
        ],
        "executions": [dict(item) for item in effective.executions],
        "terminal_assessments": [
            dict(item) for item in effective.terminal_assessments
        ],
        "asset_gaps": [dict(item) for item in effective.asset_gaps],
    }


def _execution_plan_violations(plan: Any, mode: str) -> list[dict[str, Any]]:
    policy_source = str(getattr(plan, "policy_source", "") or "").strip()
    commands = list(getattr(plan, "commands", []) or [])
    check_item_ids = _plan_check_item_ids(plan)
    profile_version_id = str(getattr(plan, "profile_version_id", "") or "").strip()
    normalized_mode = normalize_patrol_mode(mode)
    if normalized_mode != PRODUCTION_MODE and commands:
        return []
    violations = []
    if normalized_mode == PRODUCTION_MODE and not profile_version_id:
        violations.append(
            {
                "failure_code": "active_command_profile_required",
                "message": "production patrol requires an active command profile version",
                "policy_source": policy_source or "unknown",
                "profile_name": getattr(plan, "profile_name", None),
            }
        )
    if not commands and not check_item_ids:
        violations.append(
            {
                "failure_code": "collection_commands_required",
                "message": (
                    "patrol requires check items or commands from an active "
                    "profile binding"
                ),
                "policy_source": policy_source or "unknown",
                "profile_name": getattr(plan, "profile_name", None),
            }
        )
    return violations


def _coverage_summary(
    *,
    mode: str,
    commands: list[str],
    coverage: list[dict[str, Any]],
    connection_issues: list[dict[str, Any]],
    plan_violations: list[dict[str, Any]] | None = None,
    runtime_asset_health: dict[str, Any] | None = None,
) -> dict[str, Any]:
    normalized_mode = normalize_patrol_mode(mode)
    plan_violations = plan_violations or []
    eligible_count = sum(1 for item in coverage if item.get("policy_status") == "eligible")
    ready_count = sum(
        1
        for item in coverage
        if item.get("policy_status") == "eligible" and item.get("coverage_status") == "ready"
    )
    blocking = [
        item
        for item in coverage
        if item.get("policy_status") == "eligible" and item.get("coverage_status") != "ready"
    ]
    status = "ready"
    if normalized_mode != PRODUCTION_MODE:
        status = "blocked" if connection_issues else "not_required"
    elif (
        plan_violations
        or connection_issues
    ):
        status = "blocked"
    elif blocking:
        status = "degraded"
    return {
        "status": status,
        "mode": normalized_mode,
        "command_count": len(commands),
        "eligible_count": eligible_count,
        "ready_count": ready_count,
        "blocked_count": len(blocking),
        "coverage_rate": ready_count / eligible_count if eligible_count else 1.0,
        "violations": blocking,
        "plan_violations": plan_violations,
        "connection_issues": connection_issues,
        "runtime_asset_health": runtime_asset_health or {"status": "not_checked", "ready": False},
    }


def _connection_preflight(device: dict[str, Any]) -> list[dict[str, Any]]:
    """Return blocking collector issues visible before a patrol is launched."""
    issues: list[dict[str, Any]] = []
    if not (device.get("ip") or device.get("management_ip")):
        issues.append(
            {
                "failure_code": "missing_management_ip",
                "message": "management IP is required",
            }
        )
    if device.get("password"):
        return issues
    credential_ref = str(device.get("credential_ref") or "").strip()
    if not credential_ref:
        issues.append(
            {
                "failure_code": "missing_credential",
                "message": "credential_ref is required before patrol launch",
            }
        )
        return issues
    try:
        if not CredentialResolver().resolve_password(credential_ref):
            issues.append(
                {
                    "failure_code": "credential_unavailable",
                    "message": f"credential_ref did not resolve to a password: {credential_ref}",
                }
            )
    except Exception as exc:
        issues.append(
            {
                "failure_code": "credential_unavailable",
                "message": str(exc),
            }
        )
    return issues


def _latency_ms(value: Any) -> int | None:
    if value is None:
        return None
    try:
        return int(float(value) * 1000)
    except (TypeError, ValueError):
        return None


def _first_issue(issues: list[dict[str, Any]]) -> tuple[str | None, str | None]:
    if not issues:
        return None, None
    issue = issues[0]
    return (
        str(issue.get("failure_code") or "collector_preflight_failed"),
        str(issue.get("message") or ""),
    )


def _iso(value: Any) -> str | None:
    return value.isoformat() if hasattr(value, "isoformat") else None


def _find_nested_value(payload: Any, key: str) -> Any:
    if isinstance(payload, dict):
        if key in payload:
            return payload[key]
        for value in payload.values():
            found = _find_nested_value(value, key)
            if found is not None:
                return found
    elif isinstance(payload, list):
        for value in payload:
            found = _find_nested_value(value, key)
            if found is not None:
                return found
    return None


def _execution_plan_record_dict(row: CollectionExecutionPlanRecord) -> dict[str, Any]:
    return {
        "execution_plan_id": row.id,
        "device_id": row.device_id,
        "action": row.action,
        "mode": row.mode,
        "profile_name": row.profile_name,
        "profile_version_id": row.profile_version_id,
        "binding_id": row.binding_id,
        "policy_source": row.policy_source,
        "credential_ref": row.credential_ref,
        "collector_options": row.collector_options_json or {},
        "selector": row.selector_json or {},
        "plan": dict(row.plan_json or {}),
        "command_count": row.command_count,
        "plan_sha256": row.plan_sha256,
        "status": row.status,
        "launch_source": row.launch_source,
        "patrol_batch_id": row.patrol_batch_id,
        "workflow_run_id": row.workflow_run_id,
        "celery_task_id": row.celery_task_id,
        "failure_code": row.failure_code,
        "error": row.error,
        "created_at": _iso(row.created_at),
        "launched_at": _iso(row.launched_at),
        "updated_at": _iso(row.updated_at),
    }


def _list_execution_plans_for_device(device_id: str, *, limit: int = 20) -> list[dict[str, Any]]:
    limit = max(1, min(int(limit or 20), 100))
    with get_db_session() as db:
        rows = (
            db.query(CollectionExecutionPlanRecord)
            .filter(CollectionExecutionPlanRecord.device_id == device_id)
            .order_by(CollectionExecutionPlanRecord.created_at.desc())
            .limit(limit)
            .all()
        )
        return [_execution_plan_record_dict(row) for row in rows]


def _device_evidence_timeline(
    device: dict[str, Any],
    *,
    limit: int = 10,
    workflow_run_id: str | None = None,
    patrol_batch_id: str | None = None,
) -> dict[str, Any]:
    device_id = str(device.get("device_id") or "")
    limit = max(1, min(int(limit or 10), 50))
    with get_db_session() as db:
        workflow_query = db.query(WorkflowRun).filter(
            WorkflowRun.template_name == "patrol-change-event",
            WorkflowRun.context["device_id"].as_string() == device_id,
        )
        if workflow_run_id:
            workflow_query = workflow_query.filter(WorkflowRun.id == workflow_run_id)
        if patrol_batch_id:
            workflow_query = workflow_query.filter(
                WorkflowRun.context["patrol_batch_id"].as_string() == patrol_batch_id
            )
        workflow_runs = (
            workflow_query
            .order_by(WorkflowRun.created_at.desc())
            .limit(limit)
            .all()
        )
        run_entries: list[tuple[WorkflowRun | None, list[WorkflowStepRecord], str, datetime]] = []
        seen_snapshot_runs: set[str] = set()
        for workflow_run in workflow_runs:
            workflow_steps = (
                db.query(WorkflowStepRecord)
                .filter(WorkflowStepRecord.run_id == workflow_run.id)
                .order_by(WorkflowStepRecord.step_index.asc())
                .all()
            )
            workflow_snapshot_run_id = next(
                (
                    str(value)
                    for step in workflow_steps
                    if (value := _find_nested_value(step.result or {}, "snapshot_run_id"))
                ),
                "",
            )
            if workflow_snapshot_run_id:
                seen_snapshot_runs.add(workflow_snapshot_run_id)
                run_entries.append(
                    (
                        workflow_run,
                        workflow_steps,
                        workflow_snapshot_run_id,
                        workflow_run.created_at,
                    )
                )
        patrol_run_times = {
            str(row[0]): row[1]
            for row in db.query(
                CommandSnapshot.run_id,
                func.max(CommandSnapshot.observed_at).label("latest_observed_at"),
            )
            .filter(
                or_(
                    CommandSnapshot.inventory_device_id == device_id,
                    CommandSnapshot.device_id == device_id,
                )
            )
            .group_by(CommandSnapshot.run_id)
            .order_by(func.max(CommandSnapshot.observed_at).desc())
            .limit(limit)
            .all()
        }
        for run_id, latest_observed_at in (
            db.query(
                PatrolCollectionResult.run_id,
                func.max(
                    func.coalesce(
                        PatrolCollectionResult.observed_at,
                        PatrolCollectionResult.updated_at,
                    )
                ).label("latest_observed_at"),
            )
            .filter(
                or_(
                    PatrolCollectionResult.inventory_device_id == device_id,
                    PatrolCollectionResult.device_id == device_id,
                )
            )
            .group_by(PatrolCollectionResult.run_id)
            .all()
        ):
            current = patrol_run_times.get(str(run_id))
            if current is None or (latest_observed_at and latest_observed_at > current):
                patrol_run_times[str(run_id)] = latest_observed_at
        if workflow_run_id or patrol_batch_id:
            patrol_run_times = {}
        patrol_run_ids = [
            run_id
            for run_id, _observed_at in sorted(
                patrol_run_times.items(),
                key=lambda item: item[1].timestamp() if item[1] else 0.0,
                reverse=True,
            )[:limit]
        ]
        patrol_run_rows = (
            db.query(PatrolRun).filter(PatrolRun.id.in_(patrol_run_ids)).all()
            if patrol_run_ids
            else []
        )
        patrol_runs_by_id = {row.id: row for row in patrol_run_rows}
        patrol_runs = [
            patrol_runs_by_id[run_id] for run_id in patrol_run_ids if run_id in patrol_runs_by_id
        ]
        for patrol_run in patrol_runs:
            if patrol_run.id in seen_snapshot_runs:
                continue
            run_entries.append((None, [], patrol_run.id, patrol_run.started_at))
        if not workflow_run_id:
            from src.apps.inspection.backend.online_reads import device_execution_runs

            existing_ids = {entry[2] for entry in run_entries}
            for direct_run in device_execution_runs(
                db, device_id, batch_id=patrol_batch_id, limit=limit,
            ):
                if direct_run.id not in existing_ids:
                    run_entries.append((None, [], direct_run.id, direct_run.started_at))
        run_entries.sort(key=lambda entry: entry[3].timestamp() if entry[3] else 0, reverse=True)
        items: list[dict[str, Any]] = []
        for run, steps, snapshot_run_id, _created_at in run_entries[:limit]:
            patrol_run = db.get(PatrolRun, snapshot_run_id)
            from src.apps.inspection.backend.online_execution import execution_state
            from src.apps.inspection.backend.online_reads import execution_steps

            state = execution_state(patrol_run)
            run_context = dict(
                (run.context if run is not None else
                 patrol_run.metadata_json if patrol_run else {}) or {}
            )
            execution_plan_id = str(
                run_context.get("execution_plan_id")
                or ((patrol_run.metadata_json or {}).get("execution_plan_id") if patrol_run else "")
                or ""
            )
            execution_plan = None
            if execution_plan_id:
                plan_row = db.get(CollectionExecutionPlanRecord, execution_plan_id)
                if plan_row is not None:
                    execution_plan = _execution_plan_record_dict(plan_row)
            if execution_plan is None and run is not None:
                plan_row = (
                    db.query(CollectionExecutionPlanRecord)
                    .filter(CollectionExecutionPlanRecord.workflow_run_id == run.id)
                    .order_by(CollectionExecutionPlanRecord.created_at.desc())
                    .first()
                )
                if plan_row is not None:
                    execution_plan = _execution_plan_record_dict(plan_row)
                    execution_plan_id = plan_row.id
            snapshots: list[CommandSnapshot] = []
            changes: list[NetworkChangeRecord] = []
            events: list[NetworkEventRecord] = []
            occurrences: list[tuple[NetworkEventOccurrenceRecord, NetworkEventRecord]] = []
            parse_counts: dict[str, int] = {}
            policy_blocking_reason_counts: dict[str, int] = {}
            evidence_fingerprints: list[str] = []
            evidence_by_snapshot: dict[str, RawEvidenceRecord] = {}
            collection_by_execution_item: dict[str, PatrolCollectionResult] = {}
            observations: list[StateObservationRecord] = []
            changes_by_id: dict[str, NetworkChangeRecord] = {}
            observation_count = 0
            change_count = 0
            event_count = 0
            event_occurrence_count = 0
            if snapshot_run_id:
                snapshots = (
                    db.query(CommandSnapshot)
                    .filter(
                        CommandSnapshot.run_id == snapshot_run_id,
                        or_(
                            CommandSnapshot.inventory_device_id == device_id,
                            CommandSnapshot.device_id == device_id,
                        ),
                    )
                    .order_by(
                        CommandSnapshot.observed_at.asc(), CommandSnapshot.canonical_command.asc()
                    )
                    .all()
                )
                snapshot_ids = [row.id for row in snapshots]
                execution_item_ids = [row.execution_item_id for row in snapshots]
                template_ids = [
                    row.parser_template_id for row in snapshots if row.parser_template_id
                ]
                if execution_item_ids:
                    collection_by_execution_item = {
                        row.execution_item_id: row
                        for row in db.query(PatrolCollectionResult)
                        .filter(
                            PatrolCollectionResult.run_id == snapshot_run_id,
                            PatrolCollectionResult.execution_item_id.in_(execution_item_ids),
                        )
                        .all()
                    }
                if snapshot_ids:
                    evidence_by_snapshot = {
                        row.source_snapshot_id: row
                        for row in db.query(RawEvidenceRecord)
                        .filter(RawEvidenceRecord.source_snapshot_id.in_(snapshot_ids))
                        .all()
                        if row.source_snapshot_id
                    }
                    observation_count = (
                        db.query(StateObservationRecord.id)
                        .filter(StateObservationRecord.snapshot_id.in_(snapshot_ids))
                        .count()
                    )
                    observations = (
                        db.query(StateObservationRecord)
                        .filter(StateObservationRecord.snapshot_id.in_(snapshot_ids))
                        .order_by(StateObservationRecord.observed_at.desc())
                        .limit(20)
                        .all()
                    )
                change_filter = (
                    NetworkChangeRecord.run_id == snapshot_run_id,
                    NetworkChangeRecord.device_id == device_id,
                )
                change_count = db.query(NetworkChangeRecord.id).filter(*change_filter).count()
                occurrence_event_ids = db.query(NetworkEventOccurrenceRecord.event_id).filter(
                    NetworkEventOccurrenceRecord.run_id == snapshot_run_id
                )
                no_event_reason_counts = {}
                no_event_count = 0
                event_query = db.query(NetworkEventRecord).filter(
                    or_(
                        NetworkEventRecord.run_id == snapshot_run_id,
                        NetworkEventRecord.id.in_(occurrence_event_ids),
                    ),
                    NetworkEventRecord.device_id == device_id,
                )
                event_count = event_query.count()
                event_occurrence_count = (
                    db.query(NetworkEventOccurrenceRecord.id)
                    .join(
                        NetworkEventRecord,
                        NetworkEventRecord.id == NetworkEventOccurrenceRecord.event_id,
                    )
                    .filter(
                        NetworkEventOccurrenceRecord.run_id == snapshot_run_id,
                        NetworkEventRecord.device_id == device_id,
                    )
                    .count()
                )
                changes = (
                    db.query(NetworkChangeRecord)
                    .filter(*change_filter)
                    .order_by(NetworkChangeRecord.detected_at.desc())
                    .limit(20)
                    .all()
                )
                for change in changes:
                    if change.event_status == "no_event" and change.decision_reason:
                        reason = str(change.decision_reason)
                        policy_blocking_reason_counts[reason] = (
                            policy_blocking_reason_counts.get(reason, 0) + 1
                        )
                    status = str(change.event_status or '').lower()
                    if 'no_event' in status or 'not_created' in status or 'skipped' in status:
                        no_event_count += 1
                        reason = str(
                            getattr(change, "no_event_reason", None)
                            or change.decision_reason
                            or "unknown"
                        )
                        no_event_reason_counts[reason] = no_event_reason_counts.get(reason, 0) + 1
                changes_by_id = {row.id: row for row in changes}
                events = event_query.order_by(NetworkEventRecord.occurred_at.desc()).limit(20).all()
                occurrences = (
                    db.query(NetworkEventOccurrenceRecord, NetworkEventRecord)
                    .join(
                        NetworkEventRecord,
                        NetworkEventRecord.id == NetworkEventOccurrenceRecord.event_id,
                    )
                    .filter(
                        NetworkEventOccurrenceRecord.run_id == snapshot_run_id,
                        NetworkEventRecord.device_id == device_id,
                    )
                    .order_by(NetworkEventOccurrenceRecord.observed_at.desc())
                    .limit(20)
                    .all()
                )

            policy_blocking_reasons = sorted(policy_blocking_reason_counts.keys())
            policy_blocking_reason = policy_blocking_reasons[0] if policy_blocking_reasons else None
            policy_blocking_count = sum(policy_blocking_reason_counts.values())
            no_event_reasons = sorted(no_event_reason_counts.keys())
            no_event_reason = no_event_reasons[0] if no_event_reasons else None
            if snapshots:
                evidence_fingerprints = [
                    str(row.import_fingerprint)
                    for row in snapshots
                    if row.import_fingerprint
                ] or [
                    row.raw_text_hash
                    for row in snapshots
                    if row.raw_text_hash
                ]
            evidence_fingerprint = evidence_fingerprints[0] if evidence_fingerprints else None

            items.append(
                {
                    "workflow_run_id": run.id if run is not None else None,
                    "workflow_status": run.status if run is not None else None,
                    "workflow_created_at": _iso(run.created_at) if run is not None else None,
                    "workflow_completed_at": _iso(run.completed_at) if run is not None else None,
                    "execution_status": state.get("status") or (
                        patrol_run.status if patrol_run and run is None else None
                    ),
                    "execution_phase": state.get("phase"),
                    "task_id": state.get("task_id"),
                    "execution_started_at": state.get("started_at") or (
                        _iso(patrol_run.started_at) if patrol_run and run is None else None
                    ),
                    "execution_completed_at": state.get("completed_at") if state else (
                        _iso(patrol_run.completed_at) if patrol_run and run is None else None
                    ),
                    "execution_error": state.get("error"),
                    "execution_steps": execution_steps(state) if state else [],
                    "execution_plan_id": execution_plan_id or None,
                    "execution_plan": execution_plan,
                    "patrol_batch_id": run_context.get("patrol_batch_id"),
                    "snapshot_run_id": snapshot_run_id,
                    "snapshot_count": len(snapshots),
                    "collection_status": patrol_run.collection_status if patrol_run else None,
                    "parser_readiness_status": (
                        patrol_run.parser_readiness_status if patrol_run else None
                    ),
                    "fact_readiness_status": (
                        patrol_run.fact_readiness_status if patrol_run else None
                    ),
                    "analysis_status": patrol_run.analysis_status if patrol_run else None,
                    "policy_blocking_reasons": policy_blocking_reasons,
                    "policy_blocking_reason": policy_blocking_reason,
                    "policy_blocking_count": policy_blocking_count,
                    "no_event_count": no_event_count,
                    "no_event_reason": no_event_reason,
                    "no_event_reasons": no_event_reasons,
                    "evidence_fingerprint": evidence_fingerprint,
                    "observation_count": observation_count,
                    "change_count": change_count,
                    "event_count": event_count,
                    "event_occurrence_count": event_occurrence_count,
                    "parse_counts": parse_counts,
                    "steps": [
                        {
                            "step_name": step.step_name,
                            "skill_name": step.skill_name,
                            "status": step.status,
                            "started_at": _iso(step.started_at),
                            "completed_at": _iso(step.completed_at),
                            "error": step.error_message,
                        }
                        for step in steps
                    ],
                    "snapshots": [
                        {
                            "snapshot_id": row.id,
                            "evidence_id": (
                                evidence_by_snapshot[row.id].id
                                if row.id in evidence_by_snapshot
                                else evidence_id_for_command_snapshot(row.id)
                            ),
                            "source_type": (
                                evidence_by_snapshot[row.id].source_type
                                if row.id in evidence_by_snapshot
                                else "cli"
                            ),
                            "evidence_type": (
                                evidence_by_snapshot[row.id].evidence_type
                                if row.id in evidence_by_snapshot
                                else "cli_command_output"
                            ),
                            "content_ref": (
                                evidence_by_snapshot[row.id].content_ref
                                if row.id in evidence_by_snapshot
                                else f"command_snapshot:{row.id}"
                            ),
                            "evidence_fingerprint": (
                                row.import_fingerprint or row.raw_text_hash
                            ),
                            "command": row.canonical_command,
                            "parse_status": row.parse_status,
                            "output_status": row.output_status,
                            "parser_template_id": row.parser_template_id,
                            "execution_status": (
                                collection_by_execution_item[row.execution_item_id].execution_status
                                if row.execution_item_id in collection_by_execution_item
                                else None
                            ),
                            "fact_status": (
                                collection_by_execution_item[row.execution_item_id].fact_status
                                if row.execution_item_id in collection_by_execution_item
                                else None
                            ),
                            "normalization_status": (
                                collection_by_execution_item[
                                    row.execution_item_id
                                ].normalization_status
                                if row.execution_item_id in collection_by_execution_item
                                else None
                            ),
                            "parser_candidate_id": (
                                collection_by_execution_item[
                                    row.execution_item_id
                                ].parser_candidate_id
                                if row.execution_item_id in collection_by_execution_item
                                else None
                            ),
                            "raw_text_hash": row.raw_text_hash,
                            "observed_at": _iso(row.observed_at),
                        }
                        for row in snapshots[:20]
                    ],
                    "observations": [
                        {
                            "observation_id": row.id,
                            "snapshot_id": row.snapshot_id,
                            "entity_type": row.entity_type,
                            "entity_key": row.entity_key,
                            "field": row.field_name,
                            "observation_type": row.observation_type,
                            "semantic_type": row.semantic_type,
                            "unit": row.unit,
                            "value": row.value_json,
                            "value_text": row.value_text,
                            "value_number": row.value_number,
                            "confidence": row.confidence,
                            "parser_template_id": row.parser_template_id,
                            "contract_version": row.contract_version,
                            "historical_backfill": bool(
                                (row.metadata_json or {}).get("historical_backfill")
                            ),
                            "backfill_job_id": (row.metadata_json or {}).get("backfill_job_id"),
                            "observed_at": _iso(row.observed_at),
                        }
                        for row in observations
                    ],
                    "changes": [
                        {
                            "change_id": row.id,
                            "entity_type": row.entity_type,
                            "entity_key": row.entity_key,
                            "field": row.field,
                            "change_type": row.change_type,
                            "observation_type": row.observation_type,
                            "old_value": row.old_value_json,
                            "new_value": row.new_value_json,
                            "old_number": row.old_number,
                            "new_number": row.new_number,
                            "event_status": row.event_status,
                            "current_observation_id": row.current_observation_id,
                            "previous_observation_id": row.previous_observation_id,
                            "change_policy_id": row.change_policy_id,
                            "change_policy_version_id": row.change_policy_version_id,
                            "decision_reason": row.decision_reason,
                            "no_event_reason": (
                                row.decision_reason if row.event_status == "no_event" else None
                            ),
                            "comparison_type": row.comparison_type,
                            "policy_snapshot": row.policy_snapshot_json,
                            "rule_bundle_digest": row.rule_bundle_digest,
                            "analysis_revision": row.analysis_revision,
                            "detected_at": _iso(row.detected_at),
                        }
                        for row in changes[:10]
                    ],
                    "events": [
                        {
                            "event_id": row.id,
                            "event_type": row.event_type,
                            "severity": row.severity,
                            "entity_type": row.entity_type,
                            "entity_key": row.entity_key,
                            "status": row.status,
                            "occurred_at": _iso(row.occurred_at),
                            "last_seen_at": _iso(row.last_seen_at),
                            "occurrence_count": row.occurrence_count,
                            "source_change_id": row.source_change_id,
                            "old_value": (
                                changes_by_id[row.source_change_id].old_value_json
                                if row.source_change_id in changes_by_id
                                else None
                            ),
                            "new_value": (
                                changes_by_id[row.source_change_id].new_value_json
                                if row.source_change_id in changes_by_id
                                else None
                            ),
                            "old_number": (
                                changes_by_id[row.source_change_id].old_number
                                if row.source_change_id in changes_by_id
                                else None
                            ),
                            "new_number": (
                                changes_by_id[row.source_change_id].new_number
                                if row.source_change_id in changes_by_id
                                else None
                            ),
                            "change_policy_id": (
                                changes_by_id[row.source_change_id].change_policy_id
                                if row.source_change_id in changes_by_id
                                else None
                            ),
                            "change_policy_version_id": (
                                changes_by_id[row.source_change_id].change_policy_version_id
                                if row.source_change_id in changes_by_id
                                else None
                            ),
                            "policy_snapshot": (
                                changes_by_id[row.source_change_id].policy_snapshot_json
                                if row.source_change_id in changes_by_id
                                else None
                            ),
                            "event_rule_id": (
                                ((row.payload_json or {}).get("rule") or {}).get("rule_id")
                            ),
                            "event_rule_version": (
                                ((row.payload_json or {}).get("rule") or {}).get("version")
                            ),
                            "historical_backfill": bool(
                                (row.payload_json or {}).get("historical_backfill")
                            ),
                        }
                        for row in events[:10]
                    ],
                    "event_occurrences": [
                        {
                            "occurrence_id": occurrence.id,
                            "event_id": event.id,
                            "event_type": event.event_type,
                            "occurrence_kind": occurrence.occurrence_kind,
                            "source_change_id": occurrence.source_change_id,
                            "run_id": occurrence.run_id,
                            "severity": occurrence.severity,
                            "observed_at": _iso(occurrence.observed_at),
                            "event_rule_id": (
                                ((occurrence.payload_json or {}).get("rule") or {}).get("rule_id")
                            ),
                            "event_rule_version": (
                                ((occurrence.payload_json or {}).get("rule") or {}).get("version")
                            ),
                        }
                        for occurrence, event in occurrences[:20]
                    ],
                }
            )
    return {
        "device": _redact_device(device),
        "items": items,
        "limit": limit,
    }


def _patrol_readiness(
    device: dict[str, Any],
    mode: str = PRODUCTION_MODE,
    *,
    policy_resolver: CollectionPolicyResolver | None = None,
) -> dict[str, Any]:
    try:
        normalized_mode = normalize_patrol_mode(mode)
        plan = (policy_resolver or CollectionPolicyResolver()).build_plan(
            device=device,
            action="patrol",
            mode=normalized_mode,
        )
        commands = plan.command_texts()
        coverage = (
            _production_parser_coverage(device, commands)
            if normalized_mode == PRODUCTION_MODE
            else []
        )
        connection_issues = _connection_preflight(device)
        plan_violations = _execution_plan_violations(plan, normalized_mode)
        runtime_asset_health = _runtime_rule_asset_health(
            mode=normalized_mode,
            coverage=coverage, plan=plan,
        )
        summary = _coverage_summary(
            mode=normalized_mode,
            commands=commands,
            coverage=coverage,
            connection_issues=connection_issues,
            plan_violations=plan_violations,
            runtime_asset_health=runtime_asset_health,
        )
        return {
            **summary,
            "policy_source": plan.policy_source,
            "profile_name": plan.profile_name,
            "profile_version_id": getattr(plan, "profile_version_id", None),
            "check_item_ids": list(_plan_check_item_ids(plan)),
        }
    except Exception as exc:
        return {
            "status": "error",
            "mode": mode,
            "command_count": 0,
            "eligible_count": 0,
            "ready_count": 0,
            "blocked_count": 0,
            "policy_source": None,
            "profile_name": None,
            "profile_version_id": None,
            "check_item_ids": [],
            "violations": [],
            "connection_issues": [],
            "error": str(exc),
        }


def _bulk_update_preview(
    payload: DeviceBulkRequest, repository: DeviceRepository
) -> dict[str, Any]:
    devices = _resolve_bulk_devices(payload, repository)
    update = payload.update
    if update.status is not None:
        try:
            update.status = normalize_device_status(update.status)
        except ValueError as exc:
            raise HTTPException(status_code=400, detail=str(exc)) from exc
        if update.status == "deleted":
            raise HTTPException(status_code=400, detail="use bulk delete to delete devices")
    if update.group_mode not in {"replace", "append"}:
        raise HTTPException(status_code=400, detail="group_mode must be replace or append")
    changed_fields = [
        name
        for name in (
            "command_profile",
            "credential_ref",
            "status",
            "platform",
            "device_type",
            "site",
            "role",
            "groups",
        )
        if getattr(update, name) is not None
    ]
    if not changed_fields:
        raise HTTPException(status_code=400, detail="at least one update field is required")
    rows: list[dict[str, Any]] = []
    for device in devices:
        next_groups = None
        if update.groups is not None:
            if update.group_mode == "append":
                next_groups = list(dict.fromkeys([*(device.get("groups") or []), *update.groups]))
            else:
                next_groups = update.groups
        rows.append(
            {
                "device_id": device["device_id"],
                "device_name": device.get("device_name"),
                "ip": device.get("ip") or device.get("management_ip"),
                "current": _redact_device(device),
                "update": {
                    "command_profile": update.command_profile,
                    "credential_ref": update.credential_ref,
                    "status": update.status,
                    "platform": update.platform,
                    "device_type": update.device_type,
                    "site": update.site,
                    "role": update.role,
                    "groups": next_groups,
                    "group_mode": update.group_mode if update.groups is not None else None,
                },
            }
        )
    return {
        "success": True,
        "target_count": len(rows),
        "changed_fields": changed_fields,
        "rows": rows,
    }


def _patrol_launch_preview(
    payload: DevicePatrolLaunchRequest, repository: DeviceRepository
) -> dict[str, Any]:
    devices = _resolve_bulk_devices(
        DeviceBulkRequest(
            device_ids=payload.device_ids,
            dynamic_group_id=payload.dynamic_group_id,
            filters=payload.filters,
            update=DeviceBulkUpdate(status="preview"),
        ),
        repository,
        active_only=True,
    )
    platform_readiness = _patrol_platform_readiness(payload)
    if not devices:
        return {
            "success": False,
            "target_count": 0,
            "ready_count": 0,
            "error_count": 1,
            "mode": normalize_patrol_mode(payload.mode),
            "save_baseline": payload.save_baseline,
            "launch_mode": payload.launch_mode,
            "workflow_name": payload.workflow_name,
            "platform_readiness": platform_readiness,
            "quality_readiness": _patrol_target_quality_readiness([]),
            "rows": [],
            "errors": [
                {
                    "error": "当前目标条件没有匹配到可巡检设备",
                    "failure_code": "no_patrol_targets",
                    "plan_violations": [],
                    "coverage_violations": [],
                    "connection_issues": [],
                }
            ],
        }
    if not platform_readiness.get("can_launch"):
        return {
            "success": False,
            "target_count": len(devices),
            "ready_count": 0,
            "error_count": len(devices),
            "mode": normalize_patrol_mode(payload.mode),
            "save_baseline": payload.save_baseline,
            "launch_mode": payload.launch_mode,
            "workflow_name": payload.workflow_name,
            "platform_readiness": platform_readiness,
            "quality_readiness": _patrol_target_quality_readiness([]),
            "rows": [],
            "errors": [
                {
                    "error": "巡检运行环境未通过生产准入检查",
                    "failure_code": "patrol_platform_not_ready",
                    "platform_violations": platform_readiness.get("blockers") or [],
                    "plan_violations": [],
                    "coverage_violations": [],
                    "connection_issues": [],
                }
            ],
        }
    resolver = CollectionPolicyResolver()
    rows: list[dict[str, Any]] = []
    errors: list[dict[str, Any]] = []
    skipped: list[dict[str, Any]] = []
    plans_by_device: dict[str, Any] = {}
    for device in devices:
        try:
            plan = resolver.build_plan(
                device=device,
                action="patrol",
                mode=normalize_patrol_mode(payload.mode),
            )
            plans_by_device[str(device.get("device_id") or "")] = plan
            commands = plan.command_texts()
            plan_violations = _execution_plan_violations(plan, payload.mode)
            plan_check_item_ids = _plan_check_item_ids(plan)
            check_item_mode = bool(plan_check_item_ids)
            parser_coverage = (
                _production_parser_coverage(device, commands)
                if (
                    normalize_patrol_mode(payload.mode) == PRODUCTION_MODE
                    and not check_item_mode
                )
                else []
            )
            from src.core.patrol.coverage_policy import decide_parser_coverage

            coverage_decision = (
                {
                    "decision": "proceed",
                    "failure_code": None,
                    "violations": [],
                    "reason": (
                        "parser readiness is evaluated by the frozen "
                        "check-item runtime bundle"
                    ),
                }
                if check_item_mode
                else decide_parser_coverage(
                    coverage=parser_coverage,
                    mode=payload.mode,
                    execution_plan=plan,
                )
            )
            coverage_violations = list(coverage_decision.get("violations") or [])
            runtime_asset_health = (
                {
                    "status": "effective_plan_required",
                    "ready": True,
                    "collection_ready": True,
                    "analysis_ready": None,
                    "event_evaluation_ready": None,
                    "check_item_ids": list(plan_check_item_ids),
                }
                if check_item_mode
                else _runtime_rule_asset_health(
                    mode=payload.mode,
                    coverage=parser_coverage, plan=plan,
                )
            )
            connection_issues = _connection_preflight(device)
            if plan_violations or connection_issues:
                failure_code = (
                    "collector_preflight_failed"
                    if connection_issues
                    else "command_profile_not_ready"
                    if plan_violations
                    else "command_profile_not_ready"
                )
                if connection_issues and device.get("device_id"):
                    issue_code, issue_message = _first_issue(connection_issues)
                    repository.record_connection_health(
                        str(device["device_id"]),
                        status="preflight_failed",
                        failure_code=issue_code,
                        error=issue_message,
                        source="patrol-preflight",
                    )
                errors.append(
                    {
                        "device_id": device.get("device_id"),
                        "device_name": device.get("device_name"),
                        "ip": device.get("ip") or device.get("management_ip"),
                        "plan": plan.to_metadata(),
                        "error": "device patrol preflight failed",
                        "failure_code": failure_code,
                        "plan_violations": plan_violations,
                        "coverage_violations": coverage_violations,
                        "coverage_decision": coverage_decision,
                        "connection_issues": connection_issues,
                        "runtime_asset_health": runtime_asset_health,
                    }
                )
                continue
            if coverage_decision["decision"] == "block":
                errors.append(
                    {
                        "device_id": device.get("device_id"),
                        "device_name": device.get("device_name"),
                        "ip": device.get("ip") or device.get("management_ip"),
                        "plan": plan.to_metadata(),
                        "error": "parser coverage blocked launch",
                        "failure_code": coverage_decision["failure_code"],
                        "plan_violations": [],
                        "coverage_violations": coverage_violations,
                        "coverage_decision": coverage_decision,
                        "connection_issues": [],
                        "runtime_asset_health": runtime_asset_health,
                    }
                )
                continue
            if coverage_decision["decision"] == "skip":
                skipped.append(
                    {
                        "device_id": device.get("device_id"),
                        "device_name": device.get("device_name"),
                        "ip": device.get("ip") or device.get("management_ip"),
                        "plan": plan.to_metadata(),
                        "error": "parser coverage skipped device",
                        "failure_code": coverage_decision["failure_code"],
                        "plan_violations": [],
                        "coverage_violations": coverage_violations,
                        "coverage_decision": coverage_decision,
                        "connection_issues": [],
                        "runtime_asset_health": runtime_asset_health,
                    }
                )
                continue
            rows.append(
                {
                    "device_id": device["device_id"],
                    "device_name": device.get("device_name"),
                    "ip": device.get("ip") or device.get("management_ip"),
                    "device": _redact_device(device),
                    "plan": plan.to_metadata(),
                    "commands": commands,
                    "command_count": len(plan.commands),
                    "plan_violations": [],
                    "coverage_violations": coverage_violations,
                    "coverage_decision": coverage_decision,
                    "connection_issues": [],
                    "runtime_asset_health": runtime_asset_health,
                }
            )
        except Exception as exc:
            errors.append(
                {
                    "device_id": device.get("device_id"),
                    "device_name": device.get("device_name"),
                    "error": str(exc),
                }
            )
    check_item_ids_by_device = {
        device_id: list(check_item_ids)
        for device_id, plan in plans_by_device.items()
        if (check_item_ids := _plan_check_item_ids(plan))
    }
    profile_version_ids_by_device = {
        device_id: getattr(plan, "profile_version_id", None)
        for device_id, plan in plans_by_device.items()
        if _plan_check_item_ids(plan)
    }
    if check_item_ids_by_device and len(check_item_ids_by_device) != len(devices):
        missing_device_ids = sorted(
            {
                str(device.get("device_id") or "")
                for device in devices
            }
            - set(check_item_ids_by_device)
        )
        errors.append(
            {
                "error": (
                    "a patrol batch cannot mix check-item profiles with "
                    "legacy command-only profiles"
                ),
                "failure_code": "mixed_patrol_profile_governance_mode",
                "device_ids": missing_device_ids,
            }
        )
        rows = []
    scope_preview = None
    if check_item_ids_by_device:
        try:
            from src.core.check_items.catalog import (
                CheckItemCatalogRepository,
            )
            from src.core.check_items.scope import resolve_patrol_scope

            catalog = CheckItemCatalogRepository().load_active()
            scope_preview = resolve_patrol_scope(
                contract=payload.scope_contract,
                devices=devices,
                selected_check_item_ids_by_device=(
                    check_item_ids_by_device
                ),
                check_item_version_ids={
                    item.id: item.version_id
                    for item in catalog.check_items
                },
                exclusions=payload.scope_exclusions,
                not_applicable=payload.scope_not_applicable,
            ).to_dict()
        except ValueError as exc:
            errors.append(
                {
                    "error": str(exc),
                    "failure_code": "patrol_scope_contract_invalid",
                }
            )
            rows = []
    return {
        "success": not errors,
        "target_count": len(devices),
        "ready_count": len(rows),
        "error_count": len(errors),
        "skipped_count": len(skipped),
        "mode": normalize_patrol_mode(payload.mode),
        "save_baseline": payload.save_baseline,
        "launch_mode": payload.launch_mode,
        "workflow_name": payload.workflow_name,
        "platform_readiness": platform_readiness,
        "quality_readiness": _patrol_target_quality_readiness(rows),
        "target_devices": [
            _redact_device(device) for device in devices
        ],
        "check_item_ids_by_device": check_item_ids_by_device,
        "profile_version_ids_by_device": profile_version_ids_by_device,
        "check_item_mode": bool(check_item_ids_by_device),
        "scope_preview": scope_preview,
        "rows": rows,
        "errors": errors,
        "skipped": skipped,
    }


def _patrol_platform_readiness(payload: DevicePatrolLaunchRequest) -> dict[str, Any]:
    return PatrolProductionReadinessService().evaluate(
        workflow_name=payload.workflow_name,
        launch_mode=payload.launch_mode,
    )


def _patrol_target_quality_readiness(rows: list[dict[str, Any]]) -> dict[str, Any]:
    if not rows:
        return {
            "status": "not_evaluated",
            "can_determine_health": False,
            "warning_device_count": 0,
            "parser_warning_count": 0,
            "runtime_asset_warning_count": 0,
            "message": "尚未形成可评估的设备执行计划",
        }
    parser_warning_count = sum(len(row.get("coverage_violations") or []) for row in rows)
    runtime_warning_rows = [
        row
        for row in rows
        if str((row.get("runtime_asset_health") or {}).get("status") or "not_checked")
        not in {"ready", "not_required"}
    ]
    warning_device_ids = {
        str(row.get("device_id") or "")
        for row in rows
        if row.get("coverage_violations") or row in runtime_warning_rows
    }
    degraded = bool(parser_warning_count or runtime_warning_rows)
    return {
        "status": "degraded" if degraded else "ready",
        "can_determine_health": not degraded,
        "warning_device_count": len(warning_device_ids),
        "parser_warning_count": parser_warning_count,
        "runtime_asset_warning_count": len(runtime_warning_rows),
        "message": (
            "可以采集原始证据，但 Parser 或规则资产覆盖不足，最终报告将降级"
            if degraded
            else "目标命令与运行时规则资产已通过质量预检"
        ),
    }


def _patrol_target_snapshot(payload: DevicePatrolLaunchRequest) -> dict[str, Any]:
    if payload.device_ids:
        return {
            "target_type": "selected",
            "target_label": f"手工选择 {len(payload.device_ids)} 台设备",
            "selector": {"device_ids": list(payload.device_ids)},
        }
    if payload.dynamic_group_id:
        return {
            "target_type": "dynamic_group",
            "target_label": f"动态分组 {payload.dynamic_group_id}",
            "selector": {"dynamic_group_id": payload.dynamic_group_id},
        }
    selector = payload.filters.model_dump(exclude_none=True) if payload.filters else {}
    group = str(selector.get("group") or "").strip()
    if group:
        label = f"设备分组 {group}"
    elif selector:
        summary = "，".join(f"{key}={value}" for key, value in list(selector.items())[:3])
        label = f"筛选条件：{summary}"
    else:
        label = "全部启用设备"
    return {
        "target_type": "filters",
        "target_label": label,
        "selector": selector,
    }


@router.get("")
@observe_read(
    "device_list",
    variant=lambda args: (
        "with_readiness" if args.get("include_patrol_readiness") else "without_readiness"
    ),
    result_count=lambda result: len(result.get("items") or []),
)
def list_devices(
    search: str | None = None,
    device_id: str | None = None,
    device_name: str | None = None,
    ip: str | None = None,
    group: str | None = None,
    vendor: str | None = None,
    model: str | None = None,
    family: str | None = None,
    platform: str | None = None,
    device_type: str | None = None,
    site: str | None = None,
    role: str | None = None,
    status: str | None = None,
    connection_status: str | None = None,
    include_patrol_readiness: bool = False,
    readiness_mode: str = PRODUCTION_MODE,
    limit: int = Query(100, ge=1, le=500),
    offset: int = Query(0, ge=0),
):
    repository = DeviceRepository()
    normalized_status = status
    if normalized_status in {"", "all"}:
        normalized_status = None
    elif normalized_status is not None:
        normalized_status = normalize_device_status(normalized_status)
    devices = repository.get_devices_by_filter(
        search=search,
        device_id=device_id,
        device_name=device_name,
        ip=ip,
        group=group,
        vendor=vendor,
        model=model,
        family=family,
        platform=platform,
        device_type=device_type,
        site=site,
        role=role,
        status=normalized_status,
        connection_status=connection_status,
    )
    if normalized_status is None:
        devices = [device for device in devices if device.get("status") != "deleted"]
    total = len(devices)
    page = devices[offset : offset + limit]
    items: list[dict[str, Any]] = []
    readiness_resolver = None
    readiness_projections: dict[str, dict[str, Any]] = {}
    readiness_revision = "projection-unavailable"
    readiness_projection_repository = None
    if include_patrol_readiness:
        readiness_projection_repository = DeviceReadinessProjectionRepository()
        try:
            readiness_revision = readiness_projection_repository.asset_revision()
            readiness_projections = readiness_projection_repository.list_for_devices(
                [str(item.get("device_id") or "") for item in page],
                mode=normalize_patrol_mode(readiness_mode),
            )
        except Exception:
            # Rolling deployments may serve traffic before the projection migration.
            readiness_projection_repository = None
    for item in page:
        redacted = _redact_device(item)
        if include_patrol_readiness:
            device_id_value = str(item.get("device_id") or "")
            source_digest = readiness_source_digest(
                item,
                mode=normalize_patrol_mode(readiness_mode),
                asset_revision=readiness_revision,
            )
            projected = readiness_projections.get(device_id_value)
            if projected and projected.get("source_digest") == source_digest:
                redacted["patrol_readiness"] = projected
            else:
                readiness_resolver = readiness_resolver or CollectionPolicyResolver()
                computed = _patrol_readiness(
                    item,
                    readiness_mode,
                    policy_resolver=readiness_resolver,
                )
                redacted["patrol_readiness"] = computed
                if readiness_projection_repository is not None and device_id_value:
                    try:
                        readiness_projection_repository.upsert(
                            device_id=device_id_value,
                            mode=normalize_patrol_mode(readiness_mode),
                            source_digest=source_digest,
                            readiness=computed,
                        )
                    except Exception:
                        pass
        items.append(redacted)
    return {
        "items": items,
        "total": total,
        "limit": limit,
        "offset": offset,
    }


@router.post("")
def upsert_device(payload: DeviceUpsertRequest, request: Request):
    repository = DeviceRepository()
    values = payload.model_dump()
    try:
        values["status"] = normalize_device_status(values.get("status"))
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    if values["status"] == "deleted":
        raise HTTPException(status_code=400, detail="new devices cannot be created as deleted")
    device_id = repository.upsert_device(**values)
    matches = repository.get_devices_by_filter(device_id=device_id, status=None)
    _audit_device_action(
        request,
        "device_upsert",
        "device",
        device_id,
        {
            "device_name": payload.device_name,
            "management_ip": payload.management_ip,
            "vendor": payload.vendor,
            "model": payload.model,
            "groups": payload.groups or [],
        },
    )
    return {
        "success": True,
        "device_id": device_id,
        "device": _redact_device(matches[0]) if matches else None,
    }


@router.post("/import/preview")
def preview_device_import(payload: DeviceImportRequest):
    return _preview_import(payload, DeviceRepository())


@router.get("/import/template")
def get_device_import_template():
    return _device_import_template_payload()


@router.post("/import/apply")
def apply_device_import(payload: DeviceImportRequest, request: Request):
    repository = DeviceRepository()
    preview = _preview_import(payload, repository)
    applied: list[dict[str, Any]] = []
    for row in preview["rows"]:
        if not row["valid"]:
            continue
        device_id = repository.upsert_device(**row["device"])
        matches = repository.get_devices_by_filter(device_id=device_id, status=None)
        applied.append(
            {
                "row_number": row["row_number"],
                "action": row["action"],
                "device_id": device_id,
                "device": _redact_device(matches[0]) if matches else None,
            }
        )
    result = {
        "success": True,
        "dry_run": False,
        "total": preview["total"],
        "applied_count": len(applied),
        "invalid_count": preview["invalid_count"],
        "create_count": preview["create_count"],
        "update_count": preview["update_count"],
        "rows": preview["rows"],
        "applied": applied,
        "skipped": [row for row in preview["rows"] if not row["valid"]],
    }
    _audit_device_action(
        request,
        "device_import_apply",
        "device_import",
        payload.filename,
        {
            "format": payload.format,
            "filename": payload.filename,
            "total": result["total"],
            "applied_count": result["applied_count"],
            "create_count": result["create_count"],
            "update_count": result["update_count"],
            "invalid_count": result["invalid_count"],
        },
    )
    return result


@router.post("/bulk/preview")
def preview_device_bulk_update(payload: DeviceBulkRequest):
    return _bulk_update_preview(payload, DeviceRepository())


@router.post("/bulk/apply")
def apply_device_bulk_update(payload: DeviceBulkRequest, request: Request):
    repository = DeviceRepository()
    preview = _bulk_update_preview(payload, repository)
    update = payload.update
    applied: list[dict[str, Any]] = []
    for row in preview["rows"]:
        update_values: dict[str, Any] = {
            "command_profile": update.command_profile,
            "credential_ref": update.credential_ref,
            "status": update.status,
            "platform": update.platform,
            "device_type": update.device_type,
            "site": update.site,
            "role": update.role,
        }
        if update.groups is not None:
            update_values["groups"] = row["update"]["groups"]
        update_values = {key: value for key, value in update_values.items() if value is not None}
        device = repository.update_device(row["device_id"], **update_values)
        if device:
            applied.append(_redact_device(device))
    result = {
        "success": True,
        "target_count": preview["target_count"],
        "applied_count": len(applied),
        "changed_fields": preview["changed_fields"],
        "devices": applied,
    }
    _audit_device_action(
        request,
        "device_bulk_update",
        "device",
        None,
        {
            "target_count": result["target_count"],
            "applied_count": result["applied_count"],
            "changed_fields": result["changed_fields"],
            "device_ids": [item.get("device_id") for item in applied[:50]],
            "dynamic_group_id": payload.dynamic_group_id,
        },
    )
    return result


@router.post("/bulk/delete/preview")
def preview_device_bulk_delete(payload: DeviceBulkSelectionRequest):
    devices = _resolve_bulk_devices(payload, DeviceRepository())
    return {
        "success": True,
        "target_count": len(devices),
        "rows": [
            {
                "device_id": device["device_id"],
                "device_name": device.get("device_name"),
                "ip": device.get("ip") or device.get("management_ip"),
                "current": _redact_device(device),
            }
            for device in devices
        ],
    }


@router.post("/bulk/delete")
def apply_device_bulk_delete(payload: DeviceBulkSelectionRequest, request: Request):
    repository = DeviceRepository()
    devices = _resolve_bulk_devices(payload, repository)
    deleted: list[dict[str, Any]] = []
    for device in devices:
        snapshot = repository.delete_device(device["device_id"])
        if snapshot:
            deleted.append(_redact_device(snapshot))
    _audit_device_action(
        request,
        "device_bulk_delete",
        "device",
        None,
        {
            "target_count": len(devices),
            "deleted_count": len(deleted),
            "device_ids": [item.get("device_id") for item in deleted[:50]],
            "dynamic_group_id": payload.dynamic_group_id,
        },
    )
    return {
        "success": True,
        "target_count": len(devices),
        "deleted_count": len(deleted),
        "devices": deleted,
    }


@router.get("/credentials")
def list_device_credentials():
    return {"items": CredentialResolver().list_references()}


@router.post("/credentials/inspect")
def inspect_device_credential(payload: CredentialInspectRequest):
    return CredentialResolver().inspect_reference(payload.credential_ref)


@router.get("/groups")
def list_device_groups():
    return {"items": DeviceRepository().list_groups()}


@router.get("/dynamic-groups")
def list_device_dynamic_groups(status: str | None = None):
    return {"items": DeviceRepository().list_dynamic_groups(status=status)}


@router.post("/dynamic-groups")
def upsert_device_dynamic_group(payload: DeviceDynamicGroupRequest, request: Request):
    try:
        group = DeviceRepository().upsert_dynamic_group(**payload.model_dump())
        _audit_device_action(
            request,
            "device_dynamic_group_upsert",
            "device_dynamic_group",
            group.get("id"),
            {"name": group.get("name"), "status": group.get("status")},
        )
        return {"success": True, "group": group}
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc


@router.patch("/dynamic-groups/{group_id}")
def update_device_dynamic_group(
    group_id: str, payload: DeviceDynamicGroupRequest, request: Request
):
    try:
        group = DeviceRepository().upsert_dynamic_group(group_id=group_id, **payload.model_dump())
        _audit_device_action(
            request,
            "device_dynamic_group_update",
            "device_dynamic_group",
            group_id,
            {"name": group.get("name"), "status": group.get("status")},
        )
        return {"success": True, "group": group}
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc


@router.post("/dynamic-groups/{group_id}/status")
def update_device_dynamic_group_status(
    group_id: str, payload: DeviceStatusRequest, request: Request
):
    try:
        group = DeviceRepository().set_dynamic_group_status(group_id, payload.status)
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    if not group:
        raise HTTPException(status_code=404, detail="dynamic group not found")
    _audit_device_action(
        request,
        "device_dynamic_group_status",
        "device_dynamic_group",
        group_id,
        {"status": payload.status},
    )
    return {"success": True, "group": group}


@router.get("/dynamic-groups/{group_id}/preview")
def preview_device_dynamic_group(group_id: str, limit: int = Query(100, ge=1, le=500)):
    group, devices = DeviceRepository().resolve_dynamic_group_devices(group_id)
    if group is None:
        raise HTTPException(status_code=404, detail="dynamic group not found")
    return {
        "success": True,
        "group": group,
        "total": len(devices),
        "limit": limit,
        "items": [_redact_device(device) for device in devices[:limit]],
    }


@router.post("/patrol/preview")
def preview_device_patrol_launch(payload: DevicePatrolLaunchRequest):
    return _patrol_launch_preview(payload, DeviceRepository())


@router.get("/patrol/readiness")
def get_device_patrol_readiness(
    workflow_name: str = "patrol-change-event",
    launch_mode: str = "direct",
):
    if launch_mode not in {"workflow", "direct"}:
        raise HTTPException(status_code=400, detail="launch_mode must be workflow or direct")
    payload = DevicePatrolLaunchRequest(
        workflow_name=workflow_name,
        launch_mode=launch_mode,
    )
    return _patrol_platform_readiness(payload)


@router.get("/patrol/batches/{batch_id}/events")
async def stream_patrol_batch_events(batch_id: str, request: Request):
    """SSE stream of live patrol batch failures / stage updates."""
    from fastapi.responses import StreamingResponse

    from src.core.patrol.batch_events import iter_patrol_batch_events

    def event_stream():
        yield f"data: {json.dumps({'type': 'subscribed', 'patrol_batch_id': batch_id})}\n\n"
        for event in iter_patrol_batch_events(batch_id=batch_id):
            from src.apps.inspection.backend.access import authorize
            try:
                authorize(request.state.current_user.user_id)
            except HTTPException:
                return
            yield f"data: {json.dumps(event, ensure_ascii=False, default=str)}\n\n"

    return StreamingResponse(event_stream(), media_type="text/event-stream")


@router.post("/patrol/launch")
def launch_device_patrol(payload: DevicePatrolLaunchRequest, request: Request):
    from src.apps.inspection.backend.launch import launch_patrol

    user = getattr(getattr(request, "state", None), "current_user", None)
    actor = str(user.user_id) if isinstance(user, CurrentUser) else None
    result = launch_patrol(payload, actor=actor)
    _audit_device_action(
        request,
        "device_patrol_launch",
        "patrol_batch",
        result["patrol_batch_id"],
        {
            "launch_mode": payload.launch_mode,
            "workflow_name": payload.workflow_name,
            "mode": payload.mode,
            "save_baseline": payload.save_baseline,
            "target_count": result["target_count"],
            "launched_count": result["launched_count"],
            "skipped_count": result["skipped_count"],
            "queued_count": result.get("queued_count", 0),
            "dispatch_status": result.get("dispatch_status"),
            "device_ids": list(payload.device_ids or [])[:50],
        },
    )
    return result


def _launch_patrol(payload, *, actor, batch_id, intent):
    from src.apps.inspection.backend.submission import _launch_patrol as submit
    return submit(payload, actor=actor, batch_id=batch_id, intent=intent)


@router.get("/{device_id}")
def get_device(device_id: str):
    matches = DeviceRepository().get_devices_by_filter(device_id=device_id, status=None)
    if not matches:
        raise HTTPException(status_code=404, detail="device not found")
    return _redact_device(matches[0])


@router.delete("/{device_id}")
def delete_device(device_id: str, request: Request):
    device = DeviceRepository().delete_device(device_id)
    if not device:
        raise HTTPException(status_code=404, detail="device not found")
    _audit_device_action(
        request,
        "device_delete",
        "device",
        device_id,
        {"deleted_device": _redact_device(device)},
    )
    return {"success": True, "device": _redact_device(device)}


@router.patch("/{device_id}")
def update_device(device_id: str, payload: DeviceUpdateRequest, request: Request):
    values = payload.model_dump(exclude_unset=True)
    if values.get("status") is not None:
        try:
            values["status"] = normalize_device_status(values.get("status"))
        except ValueError as exc:
            raise HTTPException(status_code=400, detail=str(exc)) from exc
        if values["status"] == "deleted":
            raise HTTPException(
                status_code=400, detail="use DELETE /devices/{device_id} to delete a device"
            )
    device = DeviceRepository().update_device(device_id, **values)
    if not device:
        raise HTTPException(status_code=404, detail="device not found")
    _audit_device_action(
        request,
        "device_update",
        "device",
        device_id,
        {"changed_fields": sorted(values.keys())},
    )
    return {"success": True, "device": _redact_device(device)}


@router.post("/{device_id}/status")
def update_device_status(device_id: str, payload: DeviceStatusRequest, request: Request):
    try:
        device = DeviceRepository().set_device_status(device_id, payload.status)
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    if not device:
        raise HTTPException(status_code=404, detail="device not found")
    _audit_device_action(
        request,
        "device_status_update",
        "device",
        device_id,
        {"status": payload.status},
    )
    return {"success": True, "device": _redact_device(device)}


@router.post("/{device_id}/groups")
def update_device_groups(device_id: str, payload: DeviceGroupsRequest, request: Request):
    device = DeviceRepository().set_device_groups(device_id, payload.groups)
    if not device:
        raise HTTPException(status_code=404, detail="device not found")
    _audit_device_action(
        request,
        "device_groups_update",
        "device",
        device_id,
        {"groups": payload.groups},
    )
    return {"success": True, "device": _redact_device(device)}


@router.post("/{device_id}/effective-plan")
def preview_device_plan(device_id: str, payload: DevicePlanPreviewRequest):
    matches = DeviceRepository().get_devices_by_filter(device_id=device_id, status=None)
    if not matches:
        raise HTTPException(status_code=404, detail="device not found")
    device = matches[0]
    normalized_mode = normalize_patrol_mode(payload.mode)
    plan = CollectionPolicyResolver().build_plan(
        device=device,
        action=payload.action,
        mode=normalized_mode,
    )
    plan_metadata = plan.to_metadata()
    check_item_plan = _preview_check_item_effective_plan(
        device=device,
        plan=plan,
    )
    if check_item_plan is not None:
        collection_actions = list(
            check_item_plan.get("collection_actions") or []
        )
        commands = list(
            dict.fromkeys(
                str(item.get("command") or "").strip()
                for item in collection_actions
                if str(item.get("command") or "").strip()
            )
        )
        plan_metadata.update(
            {
                "runtime_manifest_id": check_item_plan["manifest_id"],
                "runtime_manifest_digest": check_item_plan[
                    "manifest_digest"
                ],
                "collection_items": [
                    {
                        "command": item.get("command"),
                        "action": payload.action,
                        "collection_item_id": item.get("id"),
                        "fallback_rank": item.get("fallback_rank", 0),
                        "check_item_bindings": list(
                            item.get("bindings") or []
                        ),
                    }
                    for item in collection_actions
                ],
            }
        )
        parser_coverage: list[dict[str, Any]] = []
    else:
        commands = plan.command_texts()
        parser_coverage = (
            _production_parser_coverage(device, commands)
            if normalized_mode == PRODUCTION_MODE
            else []
        )
    connection_issues = _connection_preflight(device)
    plan_violations = _execution_plan_violations(plan, normalized_mode)
    if check_item_plan is not None:
        decisions = list(check_item_plan.get("decisions") or [])
        unresolved = [
            item
            for item in decisions
            if item.get("status") not in {"resolved", "not_applicable"}
        ]
        runtime_asset_health = {
            "status": "blocked" if unresolved else "ready",
            "ready": not unresolved,
            "collection_ready": not unresolved,
            "analysis_ready": not unresolved,
            "event_evaluation_ready": not unresolved,
            "check_item_ids": list(_plan_check_item_ids(plan)),
            "runtime_manifest_id": check_item_plan["manifest_id"],
            "runtime_manifest_digest": check_item_plan[
                "manifest_digest"
            ],
            "decisions": decisions,
            "blockers": unresolved,
        }
        coverage_summary = {
            "status": (
                "blocked"
                if plan_violations or connection_issues or unresolved
                else "ready"
            ),
            "mode": normalized_mode,
            "command_count": len(commands),
            "eligible_count": len(decisions),
            "ready_count": sum(
                item.get("status") == "resolved" for item in decisions
            ),
            "not_applicable_count": sum(
                item.get("status") == "not_applicable"
                for item in decisions
            ),
            "blocked_count": len(unresolved),
            "coverage_rate": (
                sum(item.get("status") == "resolved" for item in decisions)
                / len(
                    [
                        item
                        for item in decisions
                        if item.get("status") != "not_applicable"
                    ]
                )
                if any(
                    item.get("status") != "not_applicable"
                    for item in decisions
                )
                else 1.0
            ),
            "violations": unresolved,
            "plan_violations": plan_violations,
            "connection_issues": connection_issues,
            "runtime_asset_health": runtime_asset_health,
        }
    else:
        runtime_asset_health = _runtime_rule_asset_health(
            mode=normalized_mode,
            coverage=parser_coverage,
        )
        coverage_summary = _coverage_summary(
            mode=normalized_mode,
            commands=commands,
            coverage=parser_coverage,
            connection_issues=connection_issues,
            plan_violations=plan_violations,
            runtime_asset_health=runtime_asset_health,
        )
    from src.core.rule_control.command_four_state import command_four_states

    return {
        "success": True,
        "device": _redact_device(device),
        "plan": plan_metadata,
        "check_item_effective_plan": check_item_plan,
        "commands": commands,
        "parser_coverage": parser_coverage,
        "coverage_summary": coverage_summary,
        "command_asset_states": command_four_states(
            coverage=parser_coverage,
            runtime_asset_health=runtime_asset_health,
        ),
        "runtime_asset_health": runtime_asset_health,
    }


@router.post("/{device_id}/connectivity-test")
def test_device_connectivity_endpoint(device_id: str, request: Request):
    repository = DeviceRepository()
    matches = repository.get_devices_by_filter(device_id=device_id, status=None)
    if not matches:
        raise HTTPException(status_code=404, detail="device not found")
    device = matches[0]
    issues = _connection_preflight(device)
    if issues:
        issue_code, issue_message = _first_issue(issues)
        updated = repository.record_connection_health(
            device_id,
            status="preflight_failed",
            failure_code=issue_code,
            error=issue_message,
            source="connectivity-test",
        )
        return {
            "success": False,
            "failure_code": "collector_preflight_failed",
            "device": _redact_device(updated or device),
            "connection_issues": issues,
        }
    from src.core.device_ops.tools import test_device_connectivity

    result = test_device_connectivity(device, actor_id=request.state.current_user.user_id)
    updated = repository.record_connection_health(
        device_id,
        status="connected" if result.get("success") else "failed",
        failure_code=result.get("failure_code"),
        error=result.get("error"),
        latency_ms=_latency_ms(result.get("elapsed_seconds")),
        source="connectivity-test",
        metadata={
            "device_type": result.get("device_type"),
            "prompt": result.get("prompt"),
        },
    )
    return {
        **result,
        "device": _redact_device(updated or device),
        "connection_issues": [],
    }


@router.post("/{device_id}/identity-discovery")
def discover_device_identity_endpoint(device_id: str, request: Request):
    repository = DeviceRepository()
    matches = repository.get_devices_by_filter(device_id=device_id, status=None)
    if not matches:
        raise HTTPException(status_code=404, detail="device not found")
    device = matches[0]
    issues = _connection_preflight(device)
    if issues:
        issue_code, issue_message = _first_issue(issues)
        updated = repository.record_connection_health(
            device_id,
            status="preflight_failed",
            failure_code=issue_code,
            error=issue_message,
            source="identity-discovery",
        )
        return {
            "success": False,
            "failure_code": "collector_preflight_failed",
            "identity_status": "unknown",
            "device": _redact_device(updated or device),
            "connection_issues": issues,
        }

    result = discover_device_identity(device, actor_id=request.state.current_user.user_id)
    profile_registration_error = None
    metadata = {
        "device_type": result.device_type,
        "prompt": result.prompt,
        "identity": identity_evidence(result.identity),
    }
    updated = repository.record_connection_health(
        device_id,
        status="connected" if result.failure_code != "connection_failed" else "failed",
        failure_code=result.failure_code,
        error=result.error,
        latency_ms=_latency_ms(result.elapsed_seconds),
        source="identity-discovery",
        metadata=metadata,
    )
    device_after_health = updated or device
    if result.identity.status == "known":
        attrs = dict(device_after_health.get("attributes") or {})
        attrs["identity_discovery"] = identity_evidence(result.identity)
        updated_identity = repository.update_device(
            device_id,
            vendor=result.identity.vendor,
            model=result.identity.model,
            family=result.identity.family,
            platform=platform_for_identity(result.identity, device_after_health.get("platform")),
            attributes=attrs,
        )
        device_after_health = updated_identity or device_after_health
        try:
            from src.core.parser.repository import ParserMetadataRepository

            ParserMetadataRepository().upsert_device_profile(
                vendor=result.identity.vendor or "",
                model=result.identity.model or "",
                family=result.identity.family,
                confidence=result.identity.confidence,
                evidence=identity_evidence(result.identity),
                status="confirmed",
            )
        except Exception as exc:
            profile_registration_error = str(exc)
    _audit_device_action(
        request,
        "device_identity_discovery",
        "device",
        device_id,
        {
            "success": result.success,
            "identity_status": result.identity_status,
            "identity": identity_evidence(result.identity),
            "profile_registration_error": profile_registration_error,
        },
    )
    return {
        **result.to_dict(),
        "device": _redact_device(device_after_health),
        "connection_issues": [],
        "warnings": (
            [{"code": "profile_registration_failed", "message": profile_registration_error}]
            if profile_registration_error
            else []
        ),
    }


@router.get("/{device_id}/connection-health")
def list_device_connection_health(
    device_id: str,
    limit: int = Query(50, ge=1, le=200),
):
    repository = DeviceRepository()
    matches = repository.get_devices_by_filter(device_id=device_id, status=None)
    if not matches:
        raise HTTPException(status_code=404, detail="device not found")
    return {
        "device": _redact_device(matches[0]),
        "items": repository.list_connection_health(device_id, limit=limit),
        "limit": limit,
    }


@router.get("/{device_id}/execution-plans")
def list_device_execution_plans(
    device_id: str,
    limit: int = Query(20, ge=1, le=100),
):
    matches = DeviceRepository().get_devices_by_filter(device_id=device_id, status=None)
    if not matches:
        raise HTTPException(status_code=404, detail="device not found")
    return {
        "device": _redact_device(matches[0]),
        "items": _list_execution_plans_for_device(device_id, limit=limit),
        "limit": limit,
    }


@router.get("/{device_id}/evidence-timeline")
def list_device_evidence_timeline(
    device_id: str,
    limit: int = Query(10, ge=1, le=50),
    workflow_run_id: str | None = None,
    patrol_batch_id: str | None = None,
):
    matches = DeviceRepository().get_devices_by_filter(device_id=device_id, status=None)
    if not matches:
        raise HTTPException(status_code=404, detail="device not found")
    return _device_evidence_timeline(
        matches[0],
        limit=limit,
        workflow_run_id=workflow_run_id,
        patrol_batch_id=patrol_batch_id,
    )


def _runtime_rule_asset_health(*, mode, coverage, plan=None):
    if normalize_patrol_mode(mode) != PRODUCTION_MODE:
        return {"status": "not_required", "ready": True, "collection_ready": True,
                "analysis_ready": True, "event_evaluation_ready": True}
    check_item_ids = list(_plan_check_item_ids(plan)) if plan is not None else []
    return {"status": "effective_plan_required" if check_item_ids else "blocked",
            "ready": bool(check_item_ids), "collection_ready": bool(check_item_ids),
            "analysis_ready": None, "event_evaluation_ready": None,
            "check_item_ids": check_item_ids,
            "reason": "就绪性由检查项运行包验证" if check_item_ids else "请配置检查项采集计划，旧命令契约链已退役"}
