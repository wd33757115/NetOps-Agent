"""The block-list import template: 封禁地址 and 白名单 sheets, and how to fill them."""

import io
from datetime import date, timedelta

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

HEADER = PatternFill("solid", fgColor="E8EEF8")
GUIDE = [
    "封禁地址：每行写一个或多个地址（用逗号分隔），支持单个 IP、CIDR 网段、掩码写法"
    "（10.1.1.0/255.255.255.0）和地址范围（192.0.2.10-192.0.2.20 或 192.0.2.10-20）。",
    "到期时间：留空或写“永久”表示一直封禁；只写日期表示当天结束；也可以精确到分钟。"
    "到期后封禁策略不再生效，需要时再执行清理命令删除。",
    "同一地址出现在多行时按最晚的到期时间封禁；相邻和重叠的地址合并成一条地址本成员。",
    "白名单：其中的地址不会被封禁，即使它落在封禁名单的网段里。",
    "单行超过一个 /16（65536 个地址）时需要在页面上勾选“允许大范围封禁”；0.0.0.0/0 不能封禁。",
    "CSV 导入：封禁名单一个文件（IP地址,到期时间,备注），白名单另一个文件（IP地址,备注），"
    "UTF-8 或 GBK 编码。",
]


def _sheet(sheet, header, rows, widths):
    sheet.append(header)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.fill = HEADER
    for row in rows:
        sheet.append(row)
    for letter, width in zip("ABC", widths):
        sheet.column_dimensions[letter].width = width
    sheet.freeze_panes = "A2"


def build():
    month = date.today() + timedelta(days=30)
    book = Workbook()
    block = book.active
    block.title = "封禁地址"
    _sheet(
        block,
        ["IP地址", "到期时间", "备注"],
        [
            ["203.0.113.10", "", "永久封禁的单个地址"],
            ["198.51.100.0/24", month.isoformat(), "封禁到这一天结束"],
            ["192.0.2.10-192.0.2.20", f"{month.isoformat()} 18:00", "地址范围，到期精确到分钟"],
        ],
        [28, 20, 36],
    )
    white = book.create_sheet("白名单")
    _sheet(white, ["IP地址", "备注"], [["198.51.100.53", "合作方 DNS，不封禁"]], [28, 36])
    guide = book.create_sheet("填写说明")
    guide.column_dimensions["A"].width = 110
    for line in GUIDE:
        guide.append([line])
        guide.cell(guide.max_row, 1).alignment = Alignment(wrap_text=True, vertical="top")
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()
