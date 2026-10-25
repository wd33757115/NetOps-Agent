"""A policy request sheet drafted from the resource ledger.

The ledger is asked through the platform tool registry, as the requesting user and under
that user's grant of the resource-ledger App, so this App never reads another App's tables.
The sheet has the columns of the uploaded request template, so the rest of the pipeline
treats it exactly like a request a person filled in.
"""

import io

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

HEADERS = ["seq", "源业务系统", "src_ip", "目的业务系统", "dst_ip", "port", "proto",
           "start_time", "end_time", "action", "long_link"]
PLAN_TOOL = "resource_ledger__plan_access"
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def plan(user, body):
    """Request rows from the ledger; a missing ledger or grant reads as a plain reason."""
    from fastapi import HTTPException

    from src.platform.tools.registry import catalog, invoke

    if PLAN_TOOL not in catalog():
        raise RuntimeError("按台账生成申请表需要先安装并启用“资源台账”应用")
    arguments = {key: getattr(body, key) for key in (
        "source", "destination", "source_subsystem", "destination_subsystem", "ports") if getattr(body, key)}
    try:
        return invoke(PLAN_TOOL, arguments, actor_id=user.user_id, read_only=True)
    except HTTPException as exc:
        if exc.status_code in (401, 403):
            raise RuntimeError("当前账号没有“资源台账”应用的使用权限，无法按台账生成申请表") from None
        raise RuntimeError(f"资源台账：{exc.detail}") from None


def available(user):
    """Whether this user can draft from the ledger: the App is installed and granted to them."""
    from src.platform.tools.registry import available_tools

    return bool(available_tools(actor_id=user.user_id, read_only=True, names={PLAN_TOOL}))


def build(user, body):
    """Rows, notes and the request sheet for one source → destination access."""
    result = plan(user, body)
    rows, notes = result["rows"], result["notes"]
    if not rows:
        raise RuntimeError("台账里没有可开通的访问关系" + (f"：{'；'.join(notes)}" if notes else ""))
    note = f"由资源台账按“{body.source}”到“{body.destination}”生成" + (f"；{'；'.join(notes)}" if notes else "")
    data = workbook(rows, start_time=body.start_time, end_time=body.end_time, long_link=body.long_link, note=note)
    return rows, notes, data


def workbook(rows, *, start_time=None, end_time=None, long_link=None, note=""):
    book = Workbook()
    sheet = book.active
    sheet.title = "策略申请"
    sheet.append(HEADERS)
    for cell in sheet[1]:
        cell.font, cell.fill = Font(bold=True), PatternFill("solid", fgColor="EAF2FE")
    for seq, row in enumerate(rows, 1):
        sheet.append([seq, row["source_system"], row["src_ip"], row["destination_system"], row["dst_ip"],
                      row["port"], row["proto"], start_time or None, end_time or None, "permit",
                      (long_link if row["proto"] == "tcp" else None) or None])
    sheet.append([])
    # The reader stops at this line, so the explanation never becomes a policy row.
    sheet.append([f"策略规则说明：{note}" if note else "策略规则说明"])
    for letter, width in zip("ABCDEFGHIJK", (6, 18, 34, 18, 34, 14, 8, 18, 18, 9, 10)):
        sheet.column_dimensions[letter].width = width
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()
