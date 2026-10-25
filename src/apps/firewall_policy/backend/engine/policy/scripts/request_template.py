"""The policy request sheet: one layout for the download template, ledger drafts and the
editable copy a run exports.

Rows 1–8 are the title, a note and examples; row 9 is the header the reader looks for;
requests start at row 10. Nothing above the header can be read as a request, so the
examples stay in every copy. The second sheet explains each column.
"""

import io

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

HEADERS = ["seq", "源业务系统", "src_ip", "目的业务系统", "dst_ip", "port", "proto",
           "start_time", "end_time", "action", "long_link"]
HEADER_ROW, FIRST_ROW = 9, 10
# Rows that carry borders, text formats and drop-downs for people to fill in.
FILL_ROWS = 300
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
# Text keeps "80,443" and "10.1.1.3-6" from being turned into numbers by Excel.
TEXT_COLUMNS = {"源业务系统", "src_ip", "目的业务系统", "dst_ip", "port", "long_link"}

HELP = {
    "seq": "序号，可留空",
    "源业务系统": "可选，用于展示访问关系",
    "src_ip": "源地址：IP、CIDR 网段、掩码写法或地址范围（10.1.1.3-10.1.1.6 或 10.1.1.3-6），"
    "多个用英文逗号分隔",
    "目的业务系统": "可选，用于展示访问关系",
    "dst_ip": "目的地址，写法同 src_ip",
    "port": "目的端口：单端口、范围（9000-9001），多个用英文逗号分隔；icmp 不填",
    "proto": "tcp、udp、tcp/udp 或 icmp",
    "start_time": "限时开通的开始时间，与 end_time 同时填写或同时留空，精确到分钟",
    "end_time": "结束时间；只写日期表示当天结束",
    "action": "只支持 permit",
    "long_link": "长连接老化小时数（0–24000，0 为永不老化），或写“是”（168 小时）；只对 TCP 生效",
}
EXAMPLES = [
    [1, "网银系统", "10.2.0.1,10.2.0.11", "网银核心", "10.2.16.11-10.2.16.12", "9000-9001", "tcp",
     None, None, "permit", None, "多个地址用英文逗号分隔，端口范围用 -"],
    [2, "信贷审批系统", "10.1.4.11", "风险评估系统", "10.1.1.0/24", "80,443", "tcp",
     None, None, "permit", None, "目的写网段，多个端口用逗号"],
    [3, "运维堡垒机", "10.250.1.10-11", "西安生产服务器", "10.1.0.0/20", "22", "tcp",
     None, None, "permit", None, "同一网段内的范围可简写为 10.250.1.10-11"],
    [4, "网银核心", "10.2.16.11", "第三方支付网关", "203.0.113.0/28", "443", "tcp",
     "2026-11-01", "2027-01-31", "permit", "168", "限时开通并开启长连接（168 小时）"],
    [5, "监控平台", "10.250.1.20", "生产服务器", "10.1.0.0/20", None, "icmp",
     None, None, "permit", None, "icmp 不填端口"],
]
GUIDE = [
    ("这张表做什么", "每行一组访问关系：源地址访问目的地址的端口。"
     "下发到哪台防火墙、哪个安全域由拓扑计算，不需要填写。"),
    ("从哪里填", "从“策略申请”第 10 行开始填写；第 4–8 行是示例，不会被读取，可以保留。"),
    *[(name, text) for name, text in HELP.items()],
    ("有效期", "一张表只能有一个有效期，设备上建以工单号命名的时间段。"
     "有效期和长连接目前支持华为 USG6000 和 H3C F5000。"),
    ("合并", "生成时只合并放行范围完全一致的策略，合并结果经过等价校验；也可以在页面上关闭合并。"),
]

BLUE, DEEP, PALE, MUTED = "1F4E79", "2F5597", "EEF3FA", "7F8794"
FONT = "Microsoft YaHei"
THIN = Side(style="thin", color="D6DCE5")
GRID = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def _span(sheet, row, text, *, font, fill, height, wrap=False):
    sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=len(HEADERS) + 1)
    cell = sheet.cell(row, 1, text)
    cell.font, cell.fill = font, PatternFill("solid", fgColor=fill)
    cell.alignment = Alignment(vertical="center", wrap_text=wrap, indent=1)
    sheet.row_dimensions[row].height = height


def workbook(rows=(), *, note=None):
    """The request sheet as xlsx bytes; ``rows`` are dicts keyed by HEADERS (seq optional)."""
    book = Workbook()
    sheet = book.active
    sheet.title = "策略申请"
    sheet.sheet_view.showGridLines = False
    _span(sheet, 1, "防火墙策略申请表", font=Font(name=FONT, size=16, bold=True, color="FFFFFF"),
          fill=BLUE, height=34)
    _span(sheet, 2, note or "从第 10 行开始填写，一行一组访问关系；"
          "下发到哪台防火墙、哪个安全域由拓扑计算。"
          "各列写法见“填写说明”工作表，把鼠标停在第 9 行的列名上也能看到。",
          font=Font(name=FONT, size=10, color="44546A"), fill=PALE, height=36, wrap=True)
    _span(sheet, 3, "填写示例（第 4–8 行仅供参考，不会被读取）",
          font=Font(name=FONT, size=10, bold=True, color=BLUE), fill="DCE6F2", height=22)
    example_font = Font(name=FONT, size=10, italic=True, color=MUTED)
    for offset, values in enumerate(EXAMPLES):
        for column, value in enumerate(values, 1):
            cell = sheet.cell(4 + offset, column, value)
            cell.font, cell.border = example_font, GRID
            cell.fill = PatternFill("solid", fgColor="F7F9FC")
            cell.alignment = Alignment(vertical="center")
        sheet.row_dimensions[4 + offset].height = 20

    for column, name in enumerate(HEADERS, 1):
        cell = sheet.cell(HEADER_ROW, column, name)
        cell.font = Font(name=FONT, size=10, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=DEEP)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = GRID
        cell.comment = Comment(HELP[name], "NetOps Agent", width=260, height=90)
    sheet.row_dimensions[HEADER_ROW].height = 24

    body = Font(name=FONT, size=10)
    last = FIRST_ROW + max(FILL_ROWS, len(rows)) - 1
    for row in range(FIRST_ROW, last + 1):
        for column in range(1, len(HEADERS) + 1):
            cell = sheet.cell(row, column)
            cell.font, cell.border = body, GRID
            cell.alignment = Alignment(vertical="center")
            if HEADERS[column - 1] in TEXT_COLUMNS:
                cell.number_format = "@"
    for index, values in enumerate(rows):
        for column, name in enumerate(HEADERS, 1):
            value = index + 1 if name == "seq" and not values.get("seq") else values.get(name)
            sheet.cell(FIRST_ROW + index, column, None if value in ("", None) else value)

    proto = DataValidation(type="list", formula1='"tcp,udp,tcp/udp,icmp"', allow_blank=True)
    action = DataValidation(type="list", formula1='"permit"', allow_blank=True)
    for validation, letter in ((proto, "G"), (action, "J")):
        validation.add(f"{letter}{FIRST_ROW}:{letter}{last}")
        sheet.add_data_validation(validation)

    for letter, width in zip("ABCDEFGHIJKL", (6, 16, 36, 16, 36, 14, 9, 13, 13, 9, 10, 40)):
        sheet.column_dimensions[letter].width = width
    sheet.freeze_panes = f"A{FIRST_ROW}"
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth, sheet.page_setup.fitToHeight = 1, 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True

    guide = book.create_sheet("填写说明")
    guide.sheet_view.showGridLines = False
    guide.column_dimensions["A"].width, guide.column_dimensions["B"].width = 16, 90
    title = guide.cell(1, 1, "填写说明")
    title.font = Font(name=FONT, size=14, bold=True, color=BLUE)
    guide.row_dimensions[1].height = 28
    for index, (name, text) in enumerate(GUIDE, 3):
        key, value = guide.cell(index, 1, name), guide.cell(index, 2, text)
        key.font = Font(name=FONT, size=10, bold=True, color="44546A")
        value.font = Font(name=FONT, size=10)
        for cell in (key, value):
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = Border(bottom=THIN)
        guide.row_dimensions[index].height = 32 if len(text) > 60 else 20

    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()
