"""IP block list → per-firewall address books and deny rules, cleanup commands and a check sheet."""

import argparse
import csv
import io
import json
import re
import sys
from datetime import datetime
from itertools import islice
from pathlib import Path

from block_config import VENDORS, render, verify
from block_plan import (
    BLOCK_COLUMNS,
    BOOK_LIMIT,
    MAX_ROWS,
    WHITELIST_COLUMNS,
    build,
    count,
    label,
    read_rows,
    table,
    until,
)
from openpyxl import Workbook, load_workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font
from policy_window import FORMAT

PREVIEW_LINES = 400
PREVIEW_HEAD = 160
# Excel holds at most 32767 characters in a cell; ticket scripts are split below that.
CELL = 30000
BOLD = Font(bold=True)


def safe_name(value):
    if not re.fullmatch(r"[\w.-]{1,80}", value) or value in {".", ".."}:
        raise ValueError("工单号和设备名只能包含文字、数字、下划线、点或短横线（最多 80 字符）")
    return value


def decode(data):
    for encoding in ("utf-8-sig", "gb18030"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise ValueError("CSV 文件需为 UTF-8 或 GBK 编码")


def capped(rows):
    rows = list(islice(rows, MAX_ROWS + 1001))
    if len(rows) > MAX_ROWS + 1000:
        raise ValueError(f"名单超过 {MAX_ROWS} 行，请拆分后分批导入")
    return rows


def sheets(path):
    """[(sheet title, [(row number, values)])] of a workbook, or the one table of a CSV."""
    if path.suffix.lower() == ".xlsx":
        book = load_workbook(path, read_only=True, data_only=True)
        try:
            return [
                (sheet.title, capped(enumerate(sheet.iter_rows(values_only=True), 1)))
                for sheet in book.worksheets
            ]
        finally:
            book.close()
    text = decode(path.read_bytes())
    first = text.split("\n", 1)[0]
    delimiter = "\t" if "\t" in first and "," not in first else ","
    return [(None, capped(enumerate(csv.reader(io.StringIO(text), delimiter=delimiter), 1)))]


def parsed(rows, columns, sheet):
    try:
        return table(rows, columns)
    except ValueError as exc:
        raise ValueError(f"{sheet}：{exc}") from None


def preview(text):
    """Commands as shown on the page: long book listings are elided, the rules always shown."""
    lines = text.splitlines()
    policy = next(
        (i for i, line in enumerate(lines) if line.startswith("security-policy")), len(lines)
    )
    if len(lines) <= PREVIEW_LINES or policy <= PREVIEW_HEAD:
        return text, False
    marker = f"…… 省略 {policy - PREVIEW_HEAD} 行地址本成员，完整命令请下载封禁命令包 ……"
    return "\n".join(lines[:PREVIEW_HEAD] + [marker] + lines[policy:]) + "\n", True


def chunks(text):
    part, size = [], 0
    for line in text.splitlines():
        if part and size + len(line) + 1 > CELL:
            yield "\n".join(part)
            part, size = [], 0
        part.append(line)
        size += len(line) + 1
    if part:
        yield "\n".join(part)


def few(items, limit=20):
    items = list(items)
    text = "、".join(str(item) for item in items[:limit])
    return f"{text} 等 {len(items)} 项" if len(items) > limit else text


def outcome(entry):
    """(result, books, explanation) of one valid block row."""
    notes = []
    if entry["removed"]:
        removed = few((label(*i) for i in entry["removed"]), 5)
        notes.append(f"白名单剔除 {count(entry['removed'])} 个地址：{removed}")
    if entry["later"]:
        if None in entry["later"]:
            notes.append("同一地址另有永久封禁的行，按永久封禁")
        else:
            notes.append(f"同一地址另有更晚到期的行，封禁到 {until(max(entry['later']))}")
    if not entry["blocked"]:
        return "未封禁", "", "全部在白名单内"
    return ("部分封禁" if entry["removed"] else "已封禁"), few(entry["books"], 10), "；".join(notes)


def sheet(book, title, header, rows, widths):
    target = book.create_sheet(title)
    for index, width in enumerate(widths):
        target.column_dimensions[chr(65 + index)].width = width
    target.freeze_panes = "A2"
    cells = []
    for value in header:
        item = WriteOnlyCell(target, value=value)
        item.font = BOLD
        cells.append(item)
    target.append(cells)
    for row in rows:
        target.append(row)


def report(path, batch, entries, issues, whitelist, summary, meta):
    book = Workbook(write_only=True)
    sheet(
        book,
        "汇总",
        ["项目", "内容"],
        [
            ("工单号", meta["ticket_id"]),
            ("批次", batch["prefix"]),
            ("生成时间", meta["generated_at"]),
            ("封禁名单", meta["block_name"]),
            ("白名单", meta["whitelist_name"] or "（未单独上传）"),
            ("地址本条目上限", meta["limit"]),
            ("名单行数", summary["rows"]),
            ("已封禁行", summary["blocked_rows"]),
            ("部分封禁行（白名单剔除）", summary["partial_rows"]),
            ("全部在白名单内的行", summary["whitelisted_rows"]),
            ("无效行（未封禁）", summary["error_rows"]),
            ("封禁地址数", summary["addresses"]),
            ("白名单剔除地址数", summary["removed_addresses"]),
            ("地址本成员条数", summary["members"]),
            ("地址本", summary["books"]),
            ("拒绝策略", summary["rules"]),
            ("设备", "、".join(f"{d['name']}（{d['vendor']}）" for d in meta["devices"])),
            ("", ""),
            ("到期时间", "规则 / 时间段 / 地址本 / 成员 / 地址数"),
            *[
                (
                    until(group["expiry"]),
                    f"{'、'.join(group['rules'])} / {group['time_range'] or '无（永久）'} / "
                    f"{len(group['books'])} 个 / {len(group['intervals'])} 条 / "
                    f"{count(group['intervals'])}",
                )
                for group in batch["groups"]
            ],
        ],
        [26, 90],
    )
    sheet(
        book,
        "地址本成员",
        ["地址本", "规则", "到期时间", "序号", "成员", "地址数", "来源行"],
        (
            (
                item["name"],
                item["rule"],
                until(item["expiry"]),
                position + 1,
                label(low, high),
                high - low,
                few(batch["sources"].get((item["name"], position), [])),
            )
            for item in batch["books"]
            for position, (low, high) in enumerate(item["members"])
        ),
        [26, 24, 18, 8, 34, 10, 60],
    )
    warnings = {}
    for issue in issues:
        if issue["level"] == "warning":
            warnings.setdefault((issue["sheet"], issue["row"]), []).append(issue["message"])
    rows = []
    for entry in entries:
        result, books, note = outcome(entry)
        extra = warnings.get((entry["sheet"], entry["row"]), [])
        rows.append(
            (
                entry["row"],
                entry["text"],
                entry["expiry_text"],
                entry["note"],
                result,
                books,
                "；".join(filter(None, [note, *extra])),
            )
        )
    rows += [
        (i["row"], i["value"], i["expiry_text"], i["note"], "未封禁", "", i["message"])
        for i in issues
        if i["level"] == "error"
    ]
    sheet(
        book,
        "名单逐行",
        ["行号", "IP地址", "到期时间", "备注", "结果", "地址本", "说明"],
        sorted(rows, key=lambda row: row[0]),
        [8, 34, 18, 24, 10, 34, 70],
    )
    sheet(
        book,
        "白名单",
        ["行号", "地址", "备注", "剔除地址数", "剔除范围", "涉及封禁行"],
        (
            (
                item["row"],
                item["text"],
                item["note"],
                count(item["removed"]),
                few((label(*i) for i in item["removed"]), 10),
                few(item["rows"]),
            )
            for item in whitelist
        ),
        [8, 34, 24, 12, 50, 50],
    )
    sheet(
        book,
        "问题",
        ["表", "行号", "级别", "地址", "说明"],
        (
            (
                i["sheet"],
                i["row"],
                "错误（未封禁）" if i["level"] == "error" else "提示",
                i["value"],
                i["message"],
            )
            for i in issues
        ),
        [12, 8, 14, 34, 80],
    )
    book.save(path)


def manifest(batch, devices, summary, meta):
    objects = (
        f"新增 {len(batch['time_ranges'])} 个时间段、{len(batch['books'])} 个地址本、"
        f"{len(batch['rules'])} 条拒绝策略"
    )
    expiries = "、".join(until(group["expiry"]) for group in batch["groups"])
    first_book = batch["books"][0]["name"]
    return {
        "version": "1.0",
        "generated_at": meta["generated_at"],
        "ticket_id": meta["ticket_id"],
        "ticket_title": "IP 地址封禁",
        "change_background": f"按封禁名单 {meta['block_name']} 封禁 {summary['rows']} 行来源地址",
        "change_purpose": (
            f"在 {len(devices)} 台防火墙入方向拒绝 {summary['addresses']} 个源地址"
            f"（{summary['members']} 条地址本成员，不限安全域；到期：{expiries}）"
        ),
        "requester": "",
        "requester_dept": "",
        "assignee": "",
        "priority": "P2",
        "due_date": None,
        "device_count": len(devices),
        "change_type": "配置变更",
        "risk_level": "中" if len(devices) <= 3 else "高",
        "batch": batch["prefix"],
        "devices": [
            {
                "device_name": d["name"],
                "ip_address": d.get("ip") or "",
                "vendor": d["vendor"],
                "model": "",
                "version": "",
                "before_summary": "-",
                "after_summary": objects,
            }
            for d in devices
        ],
        "scripts": [
            {
                "device_name": d["name"],
                "vendor": d["vendor"],
                "order": order,
                "commands": text,
                "command_count": len(text.splitlines()),
            }
            for d in devices
            for order, text in enumerate(chunks(d["commands"]), 1)
        ],
        "verification": [
            *[
                {
                    "step": index,
                    "title": f"验证 {d['name']} 封禁策略",
                    "method": (
                        f"display security-policy rule name {batch['rules'][0]['name']}；"
                        f"display current-configuration | include {batch['prefix']}"
                        if VENDORS[d["vendor"]] == "huawei"
                        else "display security-policy ip；"
                        f"display object-group ip address name {first_book}"
                    ),
                    "expected": "拒绝策略位于最前，引用的地址本和时间段与命令一致",
                    "verifier": "运维",
                    "standard": "地址本条目数与封禁核对表一致",
                }
                for index, d in enumerate(devices, 1)
            ],
            {
                "step": len(devices) + 1,
                "title": "封禁效果抽查",
                "method": "从名单中抽取地址，查看策略命中计数或日志",
                "expected": "来自名单地址的访问被拒绝，白名单地址不受影响",
                "verifier": "运维",
                "standard": "符合封禁名单与白名单",
            },
        ],
        "rollback": [
            {
                "step": index,
                "device_name": d["name"],
                "rollback_command": d["cleanup"],
                "expected_effect": "删除本批次的拒绝策略、地址本和时间段，恢复封禁前状态",
                "duration": "5 分钟",
                "executor": "运维",
            }
            for index, d in enumerate(devices, 1)
        ],
    }


def main():
    parser = argparse.ArgumentParser(description="IP 地址封禁")
    parser.add_argument("--request", required=True)
    parser.add_argument("--block", required=True)
    parser.add_argument("--whitelist")
    parser.add_argument("-o", "--output", required=True)
    args = parser.parse_args()
    request = json.loads(Path(args.request).read_text(encoding="utf-8"))
    ticket_id = safe_name(request["ticket_id"])
    limit = int(request.get("book_limit") or BOOK_LIMIT)
    if not 1 <= limit <= BOOK_LIMIT:
        raise ValueError(f"地址本条目上限应在 1–{BOOK_LIMIT} 之间")
    devices = request["devices"]
    if not devices or len({d["name"] for d in devices}) != len(devices):
        raise ValueError("请选择防火墙，设备名不能重复")
    for device in devices:
        safe_name(device["name"])
        if device["vendor"] not in VENDORS:
            raise ValueError(
                f"{device['name']} 的类型 {device['vendor']} 暂不支持封禁，目前支持华为、H3C"
            )
    now = datetime.strptime(request["now"], FORMAT) if request.get("now") else datetime.now()
    allow_broad = bool(request.get("allow_broad"))

    tables = sheets(Path(args.block))
    named = dict(tables)
    block_sheet = "封禁地址" if "封禁地址" in named else tables[0][0]
    block_label = block_sheet or "封禁名单"
    entries, issues = read_rows(
        parsed(named[block_sheet], BLOCK_COLUMNS, block_label),
        block_label,
        now,
        allow_broad=allow_broad,
    )
    sources = []
    if len(tables) > 1:
        title = "白名单" if "白名单" in named else tables[1][0]
        if title != block_sheet:
            sources.append((title, named[title]))
    if args.whitelist:
        for title, rows in sheets(Path(args.whitelist))[:1]:
            sources.append((f"白名单文件{'·' + title if title else ''}", rows))
    whitelist = []
    for title, rows in sources:
        found, problems = read_rows(
            parsed(rows, WHITELIST_COLUMNS, title), title, now, allow_broad=True, whitelist=True
        )
        if problems:
            # A skipped whitelist row could block an address that must stay reachable.
            raise ValueError(
                "白名单有无效行，请修正后重新导入："
                + "；".join(f"{p['sheet']}第 {p['row']} 行 {p['message']}" for p in problems[:5])
            )
        whitelist += found
    if not entries:
        errors = [i for i in issues if i["level"] == "error"][:3]
        raise ValueError(
            "名单中没有有效地址"
            + (
                "：" + "；".join(f"第 {i['row']} 行 {i['message']}" for i in errors)
                if errors
                else ""
            )
        )
    batch = build(entries, whitelist, ticket_id=ticket_id, limit=limit, now=now)
    description = ticket_id if ticket_id.isascii() else batch["prefix"]

    output = Path(args.output)
    package = output / "package"
    package.mkdir(parents=True, exist_ok=True)
    for device in devices:
        commands, cleanup = render(device["vendor"], batch, description)
        verify(device["vendor"], batch, commands, cleanup, limit)
        (package / f"{device['name']}-封禁.txt").write_text(
            commands, encoding="utf-8", newline="\n"
        )
        (package / f"{device['name']}-清理.txt").write_text(cleanup, encoding="utf-8", newline="\n")
        device.update(commands=commands, cleanup=cleanup)

    error_rows = {(i["sheet"], i["row"]) for i in issues if i["level"] == "error"}
    states = [outcome(entry)[0] for entry in entries]
    summary = dict(
        rows=len(entries) + len(error_rows),
        blocked_rows=states.count("已封禁"),
        partial_rows=states.count("部分封禁"),
        whitelisted_rows=states.count("未封禁"),
        error_rows=len(error_rows),
        warning_rows=len({(i["sheet"], i["row"]) for i in issues if i["level"] == "warning"}),
        whitelist_rows=len(whitelist),
        addresses=count(batch["blocked"]),
        removed_addresses=count(batch["removed"]),
        members=sum(len(book["members"]) for book in batch["books"]),
        books=len(batch["books"]),
        rules=len(batch["rules"]),
        time_ranges=len(batch["time_ranges"]),
    )
    meta = dict(
        ticket_id=ticket_id,
        generated_at=now.strftime(FORMAT),
        block_name=request.get("block_name") or Path(args.block).name,
        whitelist_name=request.get("whitelist_name"),
        limit=limit,
        devices=devices,
    )
    report(package / "封禁核对表.xlsx", batch, entries, issues, whitelist, summary, meta)
    ticket = manifest(batch, devices, summary, meta)
    (output / "manifest.json").write_text(json.dumps(ticket, ensure_ascii=False), encoding="utf-8")

    devices_view = []
    for device in devices:
        shown, truncated = preview(device["commands"])
        devices_view.append(
            dict(
                name=device["name"],
                vendor=device["vendor"],
                lines=len(device["commands"].splitlines()),
                preview=shown,
                truncated=truncated,
                cleanup=device["cleanup"],
            )
        )
    result = dict(
        batch=batch["prefix"],
        ticket_id=ticket_id,
        generated_at=meta["generated_at"],
        start=batch["start"],
        book_limit=limit,
        allow_broad=allow_broad,
        inputs=dict(block=meta["block_name"], whitelist=meta["whitelist_name"]),
        summary=summary,
        groups=[
            dict(
                expiry=group["expiry"],
                label=until(group["expiry"]),
                time_range=group["time_range"],
                rules=group["rules"],
                books=group["books"],
                members=len(group["intervals"]),
                addresses=count(group["intervals"]),
            )
            for group in batch["groups"]
        ],
        books=[
            dict(
                name=book["name"],
                rule=book["rule"],
                expiry=book["expiry"],
                members=len(book["members"]),
                addresses=count(book["members"]),
                first=label(*book["members"][0]),
                last=label(*book["members"][-1]),
            )
            for book in batch["books"]
        ],
        devices=devices_view,
        issues=issues[:300],
        issue_total=len(issues),
        whitelist=[
            dict(
                row=item["row"],
                sheet=item["sheet"],
                value=item["text"],
                note=item["note"],
                removed=count(item["removed"]),
                rows=item["rows"][:20],
                row_total=len(item["rows"]),
            )
            for item in whitelist[:200]
        ],
        whitelist_total=len(whitelist),
        ticket=dict(
            ticket_title=ticket["ticket_title"],
            change_background=ticket["change_background"],
            change_purpose=ticket["change_purpose"],
        ),
    )
    (output / "result.json").write_text(json.dumps(result, ensure_ascii=False), encoding="utf-8")


if __name__ == "__main__":
    try:
        main()
    except Exception as exc:
        print(f"封禁生成失败：{exc}", file=sys.stderr)
        raise SystemExit(1)
