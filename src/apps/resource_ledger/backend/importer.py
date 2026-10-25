"""Excel import and export of the whole ledger.

Preview and commit run the same code inside one transaction: preview rolls it back,
commit keeps it only when no row has an error, so the preview shows exactly what the
commit writes. Sheets apply outermost first (sites, network zones, business systems,
blocks, reservations, addresses, hosts, then providers and circuits), and anything a sheet
names that does not exist yet is created and listed in the report, so nothing appears
unannounced. A block's site, zone or subsystem that does not exist yet is created under
what the block inherits.
"""

import io
import re
from datetime import date, datetime

from fastapi import HTTPException
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from sqlalchemy import select

from . import circuit_import, netmath, service
from .common import diff, new_id, now, record, sessions
from .inventory import (
    HOST_NAMES,
    SUBSYSTEM_NAMES,
    SYSTEM_NAMES,
    named,
    port_text,
    sync_host,
    wanted_ips,
    wanted_ports,
)
from .models import (
    Address,
    Host,
    Level,
    Prefix,
    Region,
    Reservation,
    Service,
    Site,
    Space,
    Subsystem,
    System,
    Zone,
)
from .schemas import (
    ADDRESS_KINDS,
    ENVIRONMENTS,
    HOST_KINDS,
    HOST_ROLES,
    HOST_STATES,
    IMPORTANCE,
    PREFIX_STATES,
    ZONE_KINDS,
    ServiceIn,
)

MAX_ROWS = 5000
TOP = "地址资源"
SHEETS = {
    **circuit_import.SITE_SHEET,
    "网络区域": ["地域", "机房", "网络区域", "说明名", "类型", "安全等级", "备注"],
    "业务系统": ["业务系统", "子系统", "简称", "重要性", "主管部门", "业务负责人", "运维负责人", "联系方式", "备注"],
    "网段": ["网段", "层级", "归属", "用途", "VLAN", "网关", "状态", "负责人", "申请工单号", "备注", "路由域"],
    "保留范围": ["起始 IP", "结束 IP", "用途", "路由域"],
    "IP 地址": ["IP", "用途", "主机名", "DNS 名", "备注", "路由域"],
    "主机": ["主机名", "业务系统", "子系统", "IP 地址", "服务端口", "环境", "状态", "类型", "角色", "操作系统",
           "负责人", "上线日期", "下线日期", "标签", "备注"],
    **circuit_import.SHEETS,
}
REQUIRED = {
    **circuit_import.REQUIRED,
    "网络区域": ["机房", "网络区域"], "业务系统": ["业务系统"], "网段": ["网段"],
    "保留范围": ["起始 IP", "结束 IP", "用途"], "IP 地址": ["IP"], "主机": ["主机名", "业务系统"],
}
EXAMPLES = {
    "网络区域": [["西安", "西安机房", "region2", "生产业务区", "生产", "三级", ""]],
    "业务系统": [["风险评估系统", "", "风评", "核心", "风险管理部", "王敏", "周宁", "029-88886666", ""],
              ["风险评估系统", "风险评估子系统1", "", "", "", "", "周宁", "", "Web 与应用"]],
    "网段": [["10.1.0.0/16", TOP, "西部生产地址", "西安、成都生产", "", "", "使用中", "网络组", "", "", ""],
           ["10.1.0.0/17", "地域", "西安", "", "", "", "使用中", "", "", "", ""],
           ["10.1.0.0/18", "机房", "西安机房", "", "", "", "使用中", "", "", "", ""],
           ["10.1.0.0/20", "网络区域", "region2", "", "", "", "使用中", "", "", "", ""],
           ["10.1.0.0/22", "业务系统", "风险评估系统", "", "", "", "使用中", "", "", "", ""],
           ["10.1.1.0/24", "子系统", "风险评估子系统1", "Web 与应用", 101, "10.1.1.254", "使用中", "张伟",
            "CHG-2026-0101", "", ""]],
    "保留范围": [["10.1.1.200", "10.1.1.220", "Web 扩容预留", ""]],
    "IP 地址": [["10.1.1.10", "VIP", "risk-web-vip", "risk.xa.corp", "", ""]],
    "主机": [["risk-app-01", "风险评估系统", "风险评估子系统1", "10.1.1.11", "TCP/80、TCP/443", "生产", "已投产",
            "虚拟机", "应用", "CentOS 7.9", "周宁", "2026-03-01", "", "", ""]],
    **circuit_import.EXAMPLES,
}


def notes(levels):
    names = "、".join(lv.title for lv in levels) or "（还没有规划层级）"
    return [
        "每个工作表第一行是列名，请勿修改；用不到的工作表可以删除。导入顺序：站点 → 网络区域 → 业务系统 → 网段 → 保留范围 → "
        "IP 地址 → 主机 → 运营商 → 运营商网络 → 局端 → 专线类型 → 共享风险 → 专线。",
        "网络区域：机房不存在时需要填写地域，导入会一并新建。",
        "业务系统：子系统一栏为空的行是业务系统本身；填了子系统的行只取运维负责人（作为子系统负责人）和备注。"
        "重要性可填：" + "、".join(IMPORTANCE) + "。",
        f"网段：写成 10.1.1.0/24，必须是网段的起始地址；已存在的网段按本表更新。层级填“{TOP}”或规划层级（{names}）。",
        "归属：填这一级对应的地域、机房、网络区域、业务系统或子系统名称，自定义层级和地址资源填名称；下级网段自动继承上级的归属。",
        "地域和业务系统不存在时直接新建；机房、网络区域、子系统不存在时，按上级网段的地域、机房、业务系统新建。",
        "网段状态可填：" + "、".join(PREFIX_STATES) + "；IP 用途可填：" + "、".join(ADDRESS_KINDS) + "。",
        "IP 地址和保留范围必须落在一个没有子网段的网段内；网关地址和保留范围内的地址不能分配。",
        "主机：IP 地址可填多个，用逗号或顿号分隔；已分配的地址会关联到主机，未分配的地址从主机所属业务系统的网段里分配。",
        "服务端口写成 TCP/80、TCP/8000-8010、UDP/53、ICMP，多个用顿号分隔。主机的 IP 地址和服务端口以本表为准。",
        "主机环境可填：" + "、".join(ENVIRONMENTS) + "；状态可填：" + "、".join(HOST_STATES)
        + "；类型可填：" + "、".join(HOST_KINDS) + "；角色可填：" + "、".join(HOST_ROLES) + "。",
        "路由域：只有一个路由域时不用填。",
        *circuit_import.NOTES,
    ]


def workbook(sheets, levels):
    book = Workbook()
    head = Font(bold=True)
    fill = PatternFill("solid", fgColor="EAF2FE")
    first = True
    for name, columns in SHEETS.items():
        sheet = book.active if first else book.create_sheet()
        first = False
        sheet.title = name
        sheet.append(columns)
        for cell in sheet[1]:
            cell.font, cell.fill = head, fill
        for row in sheets.get(name, []):
            sheet.append(row)
        for index, column in enumerate(columns, 1):
            sheet.column_dimensions[sheet.cell(1, index).column_letter].width = max(12, len(column) * 3)
    sheet = book.create_sheet("说明")
    for line in notes(levels):
        sheet.append([line])
    sheet.column_dimensions["A"].width = 120
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def template():
    with sessions() as session:
        levels = session.scalars(select(Level).order_by(Level.position)).all()
        return workbook(EXAMPLES, levels)


def export(actor):
    """The whole ledger in the import layout, so a corrected copy can be imported back."""
    with sessions() as session:
        lookup = service.Lookup(session)
        domains = service.planners(session, lookup)
        sheets = {name: [] for name in SHEETS}
        for zone in sorted(lookup.objects["zone"].values(), key=lambda z: z.name):
            ref = lookup.zone_ref(zone.id)
            sheets["网络区域"].append([ref["region"], ref["site"], zone.name, zone.title, zone.kind,
                                    zone.security_level, zone.note])
        for system in sorted(lookup.objects["system"].values(), key=lambda s: s.name):
            sheets["业务系统"].append([system.name, None, system.short_name, system.importance, system.department,
                                    system.business_owner, system.ops_owner, system.contact, system.note])
            for sub in sorted((s for s in lookup.objects["subsystem"].values() if s.system_id == system.id),
                              key=lambda s: s.name):
                sheets["业务系统"].append([system.name, sub.name, None, None, None, None, sub.owner, None, sub.note])
        several = len(domains) > 1
        for space, index, planner in domains.values():
            domain = space.name if several else None
            for prefix in index.prefixes:
                level = planner.level(prefix)
                tier = level.title if level else (TOP if index.parent[prefix.id] is None else None)
                owner = planner.title(prefix) if level and level.dimension != "custom" else prefix.label
                sheets["网段"].append([prefix.network, tier, owner, prefix.purpose, prefix.vlan, prefix.gateway,
                                     prefix.status, prefix.owner, prefix.ticket, prefix.note, domain])
            for hold in index.reservations:
                sheets["保留范围"].append([netmath.ip_text(hold.first), netmath.ip_text(hold.last), hold.purpose, domain])
            for address in index.addresses:
                if not address.host_id:
                    sheets["IP 地址"].append([address.ip, address.kind, address.hostname, address.dns_name,
                                           address.note, domain])
        hosts = session.scalars(select(Host).order_by(Host.hostname)).all()
        ips, ports = {}, {}
        for address in session.scalars(select(Address).where(Address.host_id.is_not(None)).order_by(Address.value)):
            ips.setdefault(address.host_id, []).append(address.ip)
        for item in session.scalars(select(Service).order_by(Service.protocol, Service.port_first)):
            ports.setdefault(item.host_id, []).append(port_text(item))
        for host in hosts:
            sheets["主机"].append([
                host.hostname, lookup.name("system", host.system_id),
                lookup.name("subsystem", host.subsystem_id) if host.subsystem_id else None,
                "、".join(ips.get(host.id, [])), "、".join(ports.get(host.id, [])), host.environment, host.status,
                host.kind, host.role, host.os, host.owner, host.online_date, host.offline_date, host.tags, host.note,
            ])
        sheets.update(circuit_import.export_rows(session, several))
        return workbook(sheets, lookup.ordered_levels())


def text(value):
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


def read(data):
    try:
        book = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception:
        raise HTTPException(422, "无法读取这个文件，请上传用导入模板填写的 .xlsx 文件") from None
    found = {}
    for name in SHEETS:
        if name not in book.sheetnames:
            continue
        rows = book[name].iter_rows(values_only=True)
        header = [text(v) for v in next(rows, [])]
        lost = [c for c in REQUIRED[name] if c not in header]
        if lost:
            raise HTTPException(422, f"工作表“{name}”缺少列：{'、'.join(lost)}")
        items = []
        for number, values in enumerate(rows, 2):
            row = {key: text(values[i]) if i < len(values) else "" for i, key in enumerate(header) if key}
            if any(row.values()):
                items.append((number, row))
            if len(items) > MAX_ROWS:
                raise HTTPException(422, f"工作表“{name}”超过 {MAX_ROWS} 行，请分批导入")
        found[name] = items
    if not found:
        raise HTTPException(422, "文件里没有可导入的工作表，请使用导入模板")
    return found


def choice(row, column, options, default):
    value = row.get(column) or default
    if value not in options:
        raise ValueError(f"{column}“{value}”无效，可填：{'、'.join(options)}")
    return value


def split(value):
    return [part for part in re.split(r"[,，、;；\s]+", value or "") if part]


def services_of(value):
    """'TCP/80、UDP/53、ICMP' as service inputs."""
    result = []
    for part in split(value):
        protocol, _, ports = part.partition("/")
        if protocol.upper() not in ("TCP", "UDP", "ICMP"):
            raise ValueError(f"服务端口“{part}”无效，请写成 TCP/80、UDP/53 或 ICMP")
        result.append(ServiceIn(protocol=protocol, ports=ports or None))
    return result


class Pending:
    """A block row written in the first pass and checked once its parents are in place."""

    def __init__(self, number, label, prefix, outcome, changes, later=None):
        self.number, self.label, self.prefix = number, label, prefix
        self.outcome, self.changes, self.later = outcome, changes, later
        self.error = None


class Run(circuit_import.CircuitRows):
    def __init__(self, session, actor):
        self.session, self.actor = session, actor
        self.rows = []
        self.created = {k: [] for k in ("路由域", "地域", "机房", "网络区域", "业务系统", "子系统")}
        self.levels = {lv.title: lv for lv in session.scalars(select(Level)).all()}
        self.pending = []
        self.indexes = {}

    def result(self, sheet, number, label, outcome, message=None, changes=None):
        self.rows.append({"sheet": sheet, "row": number, "label": label, "result": outcome,
                          "message": message, "changes": changes or []})

    def log(self, entity, entity_id, label, action, detail=None):
        record(self.session, self.actor, entity, entity_id, label, action, detail)

    def made(self, kind, row, label, entity_label):
        self.session.add(row)
        self.session.flush()
        self.created.setdefault(kind, []).append(label)
        self.log(kind, row.id, entity_label, "导入")
        return row

    # ------------------------------------------------------------------ objects

    def first(self, statement):
        return self.session.scalars(statement).first()

    def region(self, name):
        return self.first(select(Region).where(Region.name == name)) or self.made(
            "地域", Region(id=new_id(), name=name), name, f"地域 {name}")

    def site(self, name, region_id=None):
        site = self.first(select(Site).where(Site.name == name))
        if site:
            return site
        if not region_id:
            raise ValueError(f"机房“{name}”不存在，需要知道它属于哪个地域才能新建")
        return self.made("机房", Site(id=new_id(), region_id=region_id, name=name, kind="自建"), name, f"机房 {name}")

    def zone(self, site, name):
        return self.first(select(Zone).where(Zone.site_id == site.id, Zone.name == name)) or self.made(
            "网络区域", Zone(id=new_id(), site_id=site.id, name=name, kind="生产"), f"{site.name} / {name}",
            f"网络区域 {name}")

    def system(self, name):
        return self.first(select(System).where(System.name == name)) or self.made(
            "业务系统", System(id=new_id(), name=name, importance="一般"), name, f"业务系统 {name}")

    def subsystem(self, system, name):
        return self.first(select(Subsystem).where(Subsystem.system_id == system.id, Subsystem.name == name)) or self.made(
            "子系统", Subsystem(id=new_id(), system_id=system.id, name=name), f"{system.name} / {name}",
            f"子系统 {system.name} / {name}")

    def space(self, name, create=False):
        if name:
            space = self.first(select(Space).where(Space.name == name))
            if space:
                return space
            if not create:
                raise ValueError(f"路由域“{name}”不存在，请先在“网段”表里导入它的网段")
            return self.made("路由域", Space(id=new_id(), name=name, kind="其他"), name, f"路由域 {name}")
        try:
            return self.session.get(Space, service.resolve_space(self.session))
        except HTTPException as exc:
            raise ValueError(exc.detail) from None

    # ------------------------------------------------------------------ sheets

    def zone_row(self, number, row):
        label = f"{row['机房']} / {row['网络区域']}"
        try:
            region = row.get("地域")
            site = self.first(select(Site).where(Site.name == row["机房"]))
            if site and region:
                actual = self.session.get(Region, site.region_id)
                if actual.name != region:
                    raise ValueError(f"机房“{site.name}”属于地域“{actual.name}”，与填写的“{region}”不一致")
            if not site:
                if not region:
                    raise ValueError(f"机房“{row['机房']}”不存在，请填写地域以便新建")
                site = self.site(row["机房"], self.region(region).id)
            values = {
                "title": row.get("说明名") or None, "kind": choice(row, "类型", ZONE_KINDS, "生产"),
                "security_level": row.get("安全等级") or None, "note": row.get("备注") or None,
            }
        except ValueError as exc:
            return self.result("网络区域", number, label, "错误", str(exc))
        zone = self.first(select(Zone).where(Zone.site_id == site.id, Zone.name == row["网络区域"]))
        if zone:
            changes = diff(zone, values, service.ZONE_NAMES)
            if not changes:
                return self.result("网络区域", number, label, "无变化")
            self.log("网络区域", zone.id, f"网络区域 {zone.name}", "导入", changes)
            return self.result("网络区域", number, label, "更新", changes=changes)
        zone = Zone(id=new_id(), site_id=site.id, name=row["网络区域"], **values)
        self.session.add(zone)
        self.session.flush()
        self.log("网络区域", zone.id, f"网络区域 {zone.name}", "导入")
        self.result("网络区域", number, label, "新增")

    def system_row(self, number, row):
        name, sub = row["业务系统"], row.get("子系统")
        label = f"{name} / {sub}" if sub else name
        try:
            if sub:
                values = {"owner": row.get("运维负责人") or None, "note": row.get("备注") or None}
                names = SUBSYSTEM_NAMES
            else:
                values = {
                    "short_name": row.get("简称") or None, "importance": choice(row, "重要性", IMPORTANCE, "一般"),
                    "department": row.get("主管部门") or None, "business_owner": row.get("业务负责人") or None,
                    "ops_owner": row.get("运维负责人") or None, "contact": row.get("联系方式") or None,
                    "note": row.get("备注") or None,
                }
                names = SYSTEM_NAMES
        except ValueError as exc:
            return self.result("业务系统", number, label, "错误", str(exc))
        system = self.first(select(System).where(System.name == name))
        if sub:
            system = system or self.system(name)
            target = self.first(select(Subsystem).where(Subsystem.system_id == system.id, Subsystem.name == sub))
            if not target:
                target = Subsystem(id=new_id(), system_id=system.id, name=sub, **values)
                self.session.add(target)
                self.session.flush()
                self.log("子系统", target.id, f"子系统 {name} / {sub}", "导入")
                return self.result("业务系统", number, label, "新增")
            entity = ("子系统", f"子系统 {name} / {sub}")
        elif not system:
            system = System(id=new_id(), name=name, **values)
            self.session.add(system)
            self.session.flush()
            self.log("业务系统", system.id, f"业务系统 {name}", "导入")
            return self.result("业务系统", number, label, "新增")
        else:
            target, entity = system, ("业务系统", f"业务系统 {name}")
        changes = diff(target, values, names)
        if not changes:
            return self.result("业务系统", number, label, "无变化")
        self.log(entity[0], target.id, entity[1], "导入", changes)
        self.result("业务系统", number, label, "更新", changes=changes)

    def prefix_row(self, number, row):
        label = row.get("网段") or "（未填网段）"
        later = None
        try:
            first, last, length, network = netmath.parse_network(row["网段"])
            label = network
            space = self.space(row.get("路由域"), create=True)
            tier = row.get("层级") or ""
            level = None
            if tier and tier != TOP:
                level = self.levels.get(tier)
                if not level:
                    raise ValueError(f"层级“{tier}”不存在，可填：{'、'.join([TOP, *self.levels])}")
            owner = (row.get("归属") or "")[:64] or None
            binding = dict.fromkeys(service.FIELDS.values())
            if level and level.dimension != "custom":
                if not owner:
                    raise ValueError(f"“{level.title}”层级需要填写归属的{service.KINDS[level.dimension]}")
                if level.dimension == "region":
                    binding["region_id"] = self.region(owner).id
                elif level.dimension == "system":
                    binding["system_id"] = self.system(owner).id
                else:
                    # Sites, zones and subsystems are found under what the block inherits.
                    later = (level.dimension, owner)
            elif level and not owner:
                raise ValueError(f"“{level.title}”层级需要填写名称")
            status = choice(row, "状态", PREFIX_STATES, "使用中")
            vlan = row.get("VLAN") or None
            if vlan is not None:
                if not vlan.isdigit() or not 1 <= int(vlan) <= 4094:
                    raise ValueError(f"VLAN“{vlan}”无效，应为 1–4094")
                vlan = int(vlan)
            gateway = row.get("网关") or None
            if gateway:
                value = netmath.ip_value(gateway)
                lo, hi = netmath.usable(first, last, length)
                if not lo <= value <= hi:
                    raise ValueError(f"网关 {gateway} 不在 {network} 的可用地址范围内")
                gateway = netmath.ip_text(value)
            values = {
                "level_id": level.id if level else None, **binding,
                "label": owner if not level or level.dimension == "custom" else None,
                "purpose": (row.get("用途") or "")[:200] or None, "vlan": vlan, "gateway": gateway,
                "status": status, "owner": row.get("负责人") or None, "ticket": row.get("申请工单号") or None,
                "note": row.get("备注") or None,
            }
        except ValueError as exc:
            return self.result("网段", number, label, "错误", str(exc))
        prefix = self.first(select(Prefix).where(
            Prefix.space_id == space.id, Prefix.first == first, Prefix.length == length))
        if prefix:
            if later:
                # Settled once the tree is in place; finding the same object changes nothing.
                values.pop(service.FIELDS[later[0]])
            changes, outcome = diff(prefix, values, service.PREFIX_NAMES), "更新"
        else:
            prefix = Prefix(id=new_id(), space_id=space.id, network=network, first=first, last=last, length=length,
                            updated_at=now(), updated_by=self.actor.username, **values)
            self.session.add(prefix)
            changes, outcome = [], "新增"
        self.pending.append(Pending(number, label, prefix, outcome, changes, later))

    def settle(self, kind, name, above):
        """Find or create a block's site, zone or subsystem under what the block inherits."""
        if kind == "site":
            return self.site(name, above.get("region")).id
        if kind == "zone":
            if above.get("site"):
                return self.zone(self.session.get(Site, above["site"]), name).id
            zones = self.session.scalars(select(Zone).where(Zone.name == name)).all()
            if len(zones) == 1:
                return zones[0].id
            raise ValueError(f"网络区域“{name}”需要上级网段先划到机房" if not zones
                             else f"有多个机房都有网络区域“{name}”，需要上级网段先划到机房")
        if above.get("system"):
            return self.subsystem(self.session.get(System, above["system"]), name).id
        subs = self.session.scalars(select(Subsystem).where(Subsystem.name == name)).all()
        if len(subs) == 1:
            return subs[0].id
        raise ValueError(f"子系统“{name}”需要上级网段先划到业务系统" if not subs
                         else f"有多个业务系统都有子系统“{name}”，需要上级网段先划到业务系统")

    def finish_prefixes(self):
        """Bind what had to wait for the tree, then check every imported block in place."""
        self.session.flush()
        lookup = service.Lookup(self.session)
        domains = service.planners(self.session, lookup)
        ordered = sorted(self.pending, key=lambda p: (p.prefix.space_id, p.prefix.first, p.prefix.length))
        for item in ordered:
            if not item.later:
                continue
            kind, name = item.later
            _, index, planner = domains[item.prefix.space_id]
            try:
                above = planner.effective(index.parent[item.prefix.id])
                ident = self.settle(kind, name, above)
            except HTTPException as exc:
                item.error = exc.detail
                continue
            except ValueError as exc:
                item.error = str(exc)
                continue
            lookup.objects[kind][ident] = self.session.get(
                {"site": Site, "zone": Zone, "subsystem": Subsystem}[kind], ident)
            if kind == "site":
                site = lookup.objects["site"][ident]
                lookup.objects["region"].setdefault(site.region_id, self.session.get(Region, site.region_id))
            if kind == "subsystem":
                sub = lookup.objects["subsystem"][ident]
                lookup.objects["system"].setdefault(sub.system_id, self.session.get(System, sub.system_id))
            field = service.FIELDS[kind]
            if getattr(item.prefix, field) != ident:
                item.changes.append({"field": service.PREFIX_NAMES[field], "before": getattr(item.prefix, field),
                                     "after": ident})
                setattr(item.prefix, field, ident)
        self.session.flush()
        for item in ordered:
            if item.error:
                continue
            _, index, planner = domains[item.prefix.space_id]
            try:
                effective = planner.effective(item.prefix)
                for address in index.addresses_in(item.prefix.first, item.prefix.last):
                    if address.host_id and index.deepest(address.value) is item.prefix:
                        host = self.session.get(Host, address.host_id)
                        problem = service.host_problem(lookup, host, effective)
                        if problem:
                            raise service.bad(f"{address.ip} 是主机 {host.hostname} 的地址，{problem}")
            except HTTPException as exc:
                item.error = exc.detail
        for item in sorted(self.pending, key=lambda p: p.number):
            if item.error:
                self.result("网段", item.number, item.label, "错误", item.error)
                continue
            changes = service.readable(lookup, item.changes)
            if item.outcome == "更新":
                item.outcome = "更新" if changes else "无变化"
                if changes:
                    item.prefix.updated_at, item.prefix.updated_by = now(), self.actor.username
            if item.outcome == "新增":
                title = domains[item.prefix.space_id][2].title(item.prefix)
                changes = [{"field": "归属" if item.prefix.level_id else "名称", "before": None, "after": title}] if title else []
            if item.outcome != "无变化":
                self.log("网段", item.prefix.id, f"网段 {item.label}", "导入", changes)
            self.result("网段", item.number, item.label, item.outcome, changes=changes)
        self.indexes = {}

    def index(self, space):
        if space.id not in self.indexes:
            self.indexes[space.id] = service.SpaceIndex(self.session, space.id)
        return self.indexes[space.id]

    def leaf(self, space, first, last):
        index = self.index(space)
        prefix = index.deepest(first)
        if prefix is None or last > prefix.last:
            raise ValueError("不在任何已规划的网段内，请先导入网段")
        if index.children[prefix.id]:
            raise ValueError(f"所在网段 {prefix.network} 已划分子网段，请把地址放进子网段")
        return index, prefix

    def reservation_row(self, number, row):
        label = f"{row.get('起始 IP', '')} – {row.get('结束 IP', '')}"
        try:
            space = self.space(row.get("路由域"))
            first, last = netmath.ip_value(row["起始 IP"]), netmath.ip_value(row["结束 IP"])
            if first > last:
                raise ValueError("起始地址不能大于结束地址")
            index, prefix = self.leaf(space, first, last)
            lo, hi = netmath.usable(prefix.first, prefix.last, prefix.length)
            if first < lo or last > hi:
                raise ValueError(f"保留范围必须在 {netmath.ip_text(lo)} – {netmath.ip_text(hi)} 之内")
            overlap = index.reservations_in(first, last)
            if any((r.first, r.last) == (first, last) for r in overlap):
                return self.result("保留范围", number, label, "无变化")
            if overlap:
                raise ValueError(f"与已有保留范围重叠（{overlap[0].purpose}）")
            busy = index.addresses_in(first, last)
            if busy:
                raise ValueError(f"范围内已有 {len(busy)} 个已分配地址，例如 {busy[0].ip}")
        except ValueError as exc:
            return self.result("保留范围", number, label, "错误", str(exc))
        reservation = Reservation(id=new_id(), space_id=space.id, first=first, last=last,
                                  purpose=row["用途"][:200], created_at=now(), created_by=self.actor.username)
        self.session.add(reservation)
        index.reservations.append(reservation)
        self.log("保留范围", reservation.id, f"保留 {label}", "导入",
                 [{"field": "用途", "before": None, "after": reservation.purpose}])
        self.result("保留范围", number, label, "新增")

    def address_row(self, number, row):
        label = row.get("IP") or "（未填 IP）"
        try:
            space = self.space(row.get("路由域"))
            value = netmath.ip_value(row["IP"])
            label = netmath.ip_text(value)
            index, prefix = self.leaf(space, value, value)
            values = {"kind": choice(row, "用途", ADDRESS_KINDS, "主机"), "hostname": row.get("主机名") or None,
                      "dns_name": row.get("DNS 名") or None, "note": row.get("备注") or None}
            existing = next(iter(index.addresses_in(value, value)), None)
            if not existing:
                service.check_free(index, prefix, value)
        except ValueError as exc:
            return self.result("IP 地址", number, label, "错误", str(exc))
        except HTTPException as exc:
            return self.result("IP 地址", number, label, "错误", exc.detail)
        if existing:
            if existing.host_id:
                values["hostname"] = existing.hostname
            changes = diff(existing, values, service.ADDRESS_NAMES)
            if not changes:
                return self.result("IP 地址", number, label, "无变化")
            self.log("IP 地址", existing.id, f"IP {label}", "导入", changes)
            return self.result("IP 地址", number, label, "更新", changes=changes)
        address = Address(id=new_id(), space_id=space.id, value=value, ip=label,
                          assigned_at=now(), assigned_by=self.actor.username, **values)
        self.session.add(address)
        index.add(address)
        self.log("IP 地址", address.id, f"IP {label}", "导入", [{"field": "网段", "before": None, "after": prefix.network}])
        self.result("IP 地址", number, label, "新增")

    def host_row(self, number, row, lookup, domains):
        label = row["主机名"]
        try:
            system = self.first(select(System).where(System.name == row["业务系统"]))
            if not system:
                raise ValueError(f"业务系统“{row['业务系统']}”不存在，请在“业务系统”表里填写")
            subsystem = None
            if row.get("子系统"):
                subsystem = self.first(select(Subsystem).where(
                    Subsystem.system_id == system.id, Subsystem.name == row["子系统"]))
                if not subsystem:
                    raise ValueError(f"{system.name} 没有子系统“{row['子系统']}”，请在“业务系统”表里填写")
            for column in ("上线日期", "下线日期"):
                if row.get(column) and not re.match(r"^\d{4}-\d{2}-\d{2}$", row[column]):
                    raise ValueError(f"{column}“{row[column]}”无效，请写成 2026-03-01")
            values = {
                "hostname": label[:128], "system_id": system.id, "subsystem_id": subsystem.id if subsystem else None,
                "environment": choice(row, "环境", ENVIRONMENTS, "生产"),
                "status": choice(row, "状态", HOST_STATES, "已投产"),
                "kind": choice(row, "类型", HOST_KINDS, "虚拟机"), "role": choice(row, "角色", HOST_ROLES, "应用"),
                "os": row.get("操作系统") or None, "owner": row.get("负责人") or None,
                "online_date": row.get("上线日期") or None, "offline_date": row.get("下线日期") or None,
                "tags": row.get("标签") or None, "note": row.get("备注") or None,
            }
            wanted = wanted_ips(split(row.get("IP 地址")))
            ports = wanted_ports(services_of(row.get("服务端口")))
        except ValueError as exc:
            return self.result("主机", number, label, "错误", str(exc))
        except HTTPException as exc:
            return self.result("主机", number, label, "错误", exc.detail)
        host = self.first(select(Host).where(Host.hostname == label, Host.system_id == system.id))
        if host:
            changes, outcome = diff(host, values, HOST_NAMES), "更新"
        else:
            host = Host(id=new_id(), updated_at=now(), updated_by=self.actor.username, **values)
            self.session.add(host)
            changes, outcome = [], "新增"
        self.session.flush()
        try:
            changes += sync_host(self.session, self.actor, host, wanted, ports, lookup, domains)
        except HTTPException as exc:
            return self.result("主机", number, label, "错误", exc.detail)
        if outcome == "更新" and not changes:
            return self.result("主机", number, label, "无变化")
        host.updated_at, host.updated_by = now(), self.actor.username
        changes = named(lookup, changes)
        self.log("主机", host.id, f"主机 {host.hostname}", "导入", changes or None)
        self.result("主机", number, label, outcome, changes=changes)


def run(actor, data, commit):
    if actor.role != "admin":
        raise HTTPException(403, "批量导入台账需要管理员账号")
    sheets = read(data)
    from sqlalchemy.orm import Session

    from src.platform.kernel.migrations import business_engine

    with Session(business_engine()) as session:
        transaction = session.begin()
        job = Run(session, actor)
        try:
            for number, row in sheets.get("站点", []):
                job.site_row(number, row)
            for number, row in sheets.get("网络区域", []):
                job.zone_row(number, row)
            for number, row in sheets.get("业务系统", []):
                job.system_row(number, row)
            for number, row in sheets.get("网段", []):
                job.prefix_row(number, row)
            job.finish_prefixes()
            for number, row in sheets.get("保留范围", []):
                job.reservation_row(number, row)
            for number, row in sheets.get("IP 地址", []):
                job.address_row(number, row)
            session.flush()
            if sheets.get("主机"):
                lookup = service.Lookup(session)
                domains = service.planners(session, lookup)
                for number, row in sheets["主机"]:
                    job.host_row(number, row, lookup, domains)
            session.flush()
            for sheet, handle in (("运营商", job.provider_row), ("运营商网络", job.network_row), ("局端", job.pop_row),
                                  ("专线类型", job.type_row), ("共享风险", job.risk_row), ("专线", job.circuit_row)):
                for number, row in sheets.get(sheet, []):
                    handle(number, row)
                session.flush()
            errors = sum(1 for r in job.rows if r["result"] == "错误")
            keep = commit and errors == 0
            if keep:
                transaction.commit()
            else:
                transaction.rollback()
        except Exception:
            transaction.rollback()
            raise
    summary = {}
    for row in job.rows:
        counts = summary.setdefault(row["sheet"], {"新增": 0, "更新": 0, "无变化": 0, "错误": 0})
        counts[row["result"]] += 1
    return {
        "committed": keep,
        "errors": errors,
        "summary": summary,
        "created": {k: v for k, v in job.created.items() if v},
        "rows": job.rows,
    }
