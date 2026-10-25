"""Human-facing patrol diff and report helpers."""

from __future__ import annotations

import io
from collections import Counter
from html import escape
from pathlib import Path
from typing import Any

from src.core.patrol.change_event_repository import ChangeEventRepository
from src.core.patrol.handlers import DetectChangesHandler
from src.core.patrol.observation_report import (
    ReportTable,
    build_observation_report_tables,
)
from src.core.patrol.repository import PatrolRepository
from src.core.patrol.unstructured_diff import summarize_unstructured_changes
from src.core.state.repository import StateRepository


def build_run_quality_summary(run_id: str) -> dict[str, Any]:
    repository = PatrolRepository()
    run = repository.get_run(run_id) or {}
    results = repository.list_collection_results(run_id)
    eligible = [
        item for item in results if item.get("parse_status") not in {"ignored", "not_applicable"}
    ]
    parser_eligible = [item for item in eligible if item.get("parse_status") != "raw_hashed"]
    parser_ready = [
        item for item in parser_eligible if item.get("parse_status") in {"parsed", "parse_empty"}
    ]
    fact_ready = [item for item in eligible if item.get("fact_status") in {"ready", "empty"}]
    normalized = [
        item for item in eligible if item.get("normalization_status") in {"normalized", "empty"}
    ]
    denominator = len(eligible)
    parser_denominator = len(parser_eligible)
    return {
        "collection_status": run.get("collection_status") or "unknown",
        "parser_readiness_status": run.get("parser_readiness_status") or "unknown",
        "fact_readiness_status": run.get("fact_readiness_status") or "unknown",
        "analysis_status": run.get("analysis_status") or "unknown",
        "command_count": len(results),
        "execution_success_count": sum(
            item.get("execution_status") == "complete" for item in results
        ),
        "parser_ready_rate": (
            len(parser_ready) / parser_denominator if parser_denominator else 1.0
        ),
        "fact_ready_rate": len(fact_ready) / denominator if denominator else 1.0,
        "normalization_ready_rate": len(normalized) / denominator if denominator else 1.0,
        "template_missing_count": sum(
            item.get("parse_status") == "template_missing" for item in results
        ),
        "parser_backlog_count": len(
            {item.get("parser_candidate_id") for item in results if item.get("parser_candidate_id")}
        ),
        "blocked_by_contract_count": sum(
            item.get("fact_status") == "blocked_by_contract" for item in results
        ),
        "blocked_by_parser_count": sum(
            item.get("fact_status") == "blocked_by_parser" for item in results
        ),
        "valid_empty_count": sum(item.get("output_status") == "valid_empty" for item in results),
        "empty_unverified_count": sum(
            item.get("output_status") == "empty_unverified" for item in results
        ),
        "unsupported_command_count": sum(
            item.get("output_status") == "unsupported_command" for item in results
        ),
        "command_error_count": sum(
            item.get("output_status") == "command_error" for item in results
        ),
        "raw_hash_diff_count": sum(item.get("parse_status") == "raw_hashed" for item in results),
    }


def build_human_diff(params: dict[str, Any]) -> dict[str, Any]:
    effective = dict(params)
    effective.setdefault("persist", False)
    result = DetectChangesHandler().handle(effective)
    return {
        **result,
        "report_type": "patrol_diff",
        "persisted_as_network_changes": bool(effective.get("persist")),
    }


def render_run_report(
    *,
    run_id: str,
    output_path: str | None = None,
    device_id: str | None = None,
    output_format: str = "markdown",
) -> dict[str, Any]:
    repository = ChangeEventRepository()
    changes = repository.list_changes(run_id=run_id, device_id=device_id)
    events = repository.list_run_event_facts(run_id=run_id, device_id=device_id)
    event_lifecycle = repository.summarize_run_events(
        run_id=run_id,
        device_id=device_id,
    )
    event_counter = Counter(event["event_type"] for event in events)
    severity_counter = Counter(event["severity"] for event in events)
    quality = build_run_quality_summary(run_id)
    unstructured = summarize_unstructured_changes(
        changes,
        collected_output_count=int(quality["raw_hash_diff_count"]),
    )
    lines = [
        f"# Patrol Report: {run_id}",
        "",
        "## Summary",
        "",
        f"- Changes: {len(changes)}",
        f"- Event lifecycle facts: {len(events)}",
        f"- New events: {event_lifecycle['created_event_count']}",
        f"- Event updates: {event_lifecycle['updated_event_count']}",
        f"- Recoveries: {event_lifecycle['recovered_event_count']}",
        f"- Open after this run: {event_lifecycle['open_after_run_count']}",
        f"- Still open now: {event_lifecycle['current_open_event_count']}",
        f"- Collection status: {quality['collection_status']}",
        f"- Parser ready: {quality['parser_ready_rate']:.1%}",
        f"- Fact ready: {quality['fact_ready_rate']:.1%}",
        f"- Normalization ready: {quality['normalization_ready_rate']:.1%}",
        f"- Missing templates: {quality['template_missing_count']}",
        f"- Parser backlog: {quality['parser_backlog_count']}",
        f"- Valid empty outputs: {quality['valid_empty_count']}",
        f"- Unverified empty outputs: {quality['empty_unverified_count']}",
        f"- Unsupported commands: {quality['unsupported_command_count']}",
        f"- Raw hash comparisons: {quality['raw_hash_diff_count']}",
        f"- Non-structured outputs changed: {unstructured['changed_output_count']}",
        (
            "- Non-structured changed lines: "
            f"+{unstructured['added_line_count']} / -{unstructured['removed_line_count']}"
        ),
    ]
    if event_counter:
        lines.extend(["", "## Events By Type", ""])
        lines.extend(
            f"- {event_type}: {count}" for event_type, count in sorted(event_counter.items())
        )
    if severity_counter:
        lines.extend(["", "## Events By Severity", ""])
        lines.extend(
            f"- {severity}: {count}" for severity, count in sorted(severity_counter.items())
        )
    if events:
        lines.extend(["", "## Event Details", ""])
        for event in events:
            lines.append(
                "- "
                f"{event['occurred_at']} "
                f"{event['severity']} "
                f"{event['event_type']} "
                f"{event.get('device_id') or '-'} "
                f"{event.get('entity_key') or ''}"
            )
    if unstructured["items"]:
        lines.extend(["", "## Non-structured Output Changes", ""])
        for item in unstructured["items"]:
            categories = (
                ", ".join(
                    str(category.get("label") or category.get("name") or "other")
                    for category in item.get("categories") or []
                )
                or "uncategorized"
            )
            lines.append(
                "- "
                f"{item.get('device_name') or '-'} · "
                f"`{item.get('command') or '-'}` · "
                f"+{item['added_line_count']} / -{item['removed_line_count']} · "
                f"{categories}"
            )
            for preview in item.get("preview") or []:
                marker = "+" if preview.get("operation") == "added" else "-"
                lines.append(
                    f"  - `{marker}` {preview.get('category_label') or 'other'}: "
                    f"`{preview.get('text') or ''}`"
                )
    markdown = "\n".join(lines) + "\n"
    output_format = str(output_format or "markdown").lower()
    rendered = markdown
    if output_format == "html":
        rendered = _render_html(
            run_id,
            changes,
            events,
            event_counter,
            severity_counter,
            quality,
            unstructured,
        )
    written_path = None
    observations: list[dict[str, Any]] = []
    if output_path:
        path = Path(output_path)
        path.parent.mkdir(parents=True, exist_ok=True)
        if output_format == "xlsx":
            observations = StateRepository().list_run_observations(
                run_id,
                device_ids=[device_id] if device_id else None,
            )
            _write_xlsx(
                path,
                run_id,
                changes,
                events,
                quality,
                unstructured,
                observations,
            )
            rendered = ""
        elif output_format == "html":
            path.write_text(rendered, encoding="utf-8")
        else:
            path.write_text(markdown, encoding="utf-8")
        written_path = str(path)
    return {
        "success": True,
        "run_id": run_id,
        "output_format": output_format,
        "change_count": len(changes),
        "event_count": len(events),
        "event_lifecycle": event_lifecycle,
        "event_types": dict(event_counter),
        "severities": dict(severity_counter),
        "quality": quality,
        "unstructured_comparison": unstructured,
        "observation_count": len(observations),
        "markdown": markdown,
        "content": rendered,
        "output_path": written_path,
    }


def _render_html(
    run_id: str,
    changes: list[dict[str, Any]],
    events: list[dict[str, Any]],
    event_counter: Counter,
    severity_counter: Counter,
    quality: dict[str, Any],
    unstructured: dict[str, Any],
) -> str:
    rows = "\n".join(
        "<tr>"
        f"<td>{escape(str(event.get('occurred_at') or ''))}</td>"
        f"<td>{escape(str(event.get('severity') or ''))}</td>"
        f"<td>{escape(str(event.get('event_type') or ''))}</td>"
        f"<td>{escape(str(event.get('device_id') or ''))}</td>"
        f"<td>{escape(str(event.get('entity_key') or ''))}</td>"
        "</tr>"
        for event in events
    )
    unstructured_rows = "\n".join(
        "<tr>"
        f"<td>{escape(str(item.get('device_name') or ''))}</td>"
        f"<td>{escape(str(item.get('command') or ''))}</td>"
        f"<td>{item.get('added_line_count') or 0}</td>"
        f"<td>{item.get('removed_line_count') or 0}</td>"
        f"<td>{escape(_unstructured_categories(item))}</td>"
        f"<td>{escape(_unstructured_preview(item))}</td>"
        "</tr>"
        for item in unstructured.get("items") or []
    )
    return f"""<!doctype html>
<html>
<head>
  <meta charset="utf-8">
  <title>Patrol Report {escape(run_id)}</title>
  <style>
    body {{ font-family: Arial, sans-serif; margin: 24px; color: #1f2937; }}
    table {{ border-collapse: collapse; width: 100%; }}
    th, td {{ border: 1px solid #d1d5db; padding: 6px 8px; text-align: left; }}
    th {{ background: #f3f4f6; }}
  </style>
</head>
<body>
  <h1>Patrol Report: {escape(run_id)}</h1>
  <h2>Summary</h2>
  <p>Changes: {len(changes)}<br>Events: {len(events)}<br>
  Collection: {escape(str(quality["collection_status"]))}<br>
  Parser ready: {quality["parser_ready_rate"]:.1%}<br>
  Fact ready: {quality["fact_ready_rate"]:.1%}<br>
  Normalization ready: {quality["normalization_ready_rate"]:.1%}<br>
  Missing templates: {quality["template_missing_count"]}<br>
  Parser backlog: {quality["parser_backlog_count"]}</p>
  <p>Valid empty outputs: {quality["valid_empty_count"]}<br>
  Unverified empty outputs: {quality["empty_unverified_count"]}<br>
  Unsupported commands: {quality["unsupported_command_count"]}<br>
  Raw hash comparisons: {quality["raw_hash_diff_count"]}<br>
  Non-structured outputs changed: {unstructured["changed_output_count"]}<br>
  Non-structured changed lines:
  +{unstructured["added_line_count"]} / -{unstructured["removed_line_count"]}</p>
  <h2>Events By Type</h2>
  <ul>{"".join(f"<li>{escape(k)}: {v}</li>" for k, v in sorted(event_counter.items()))}</ul>
  <h2>Events By Severity</h2>
  <ul>{"".join(f"<li>{escape(k)}: {v}</li>" for k, v in sorted(severity_counter.items()))}</ul>
  <h2>Event Details</h2>
  <table>
    <thead><tr><th>Time</th><th>Severity</th><th>Type</th><th>Device</th><th>Entity</th></tr></thead>
    <tbody>{rows}</tbody>
  </table>
  <h2>Non-structured Output Changes</h2>
  <table>
    <thead><tr><th>Device</th><th>Command</th><th>Added</th><th>Removed</th>
    <th>Categories</th><th>Safe Preview</th></tr></thead>
    <tbody>{unstructured_rows}</tbody>
  </table>
</body>
</html>
"""


def _write_xlsx(
    path: Path,
    run_id: str,
    changes: list[dict[str, Any]],
    events: list[dict[str, Any]],
    quality: dict[str, Any],
    unstructured: dict[str, Any],
    observations: list[dict[str, Any]],
) -> None:
    workbook = _build_xlsx_workbook(
        report_id=run_id,
        changes=changes,
        events=events,
        quality=quality,
        unstructured=unstructured,
        observations=observations,
    )
    workbook.save(path)


def render_patrol_batch_xlsx(
    patrol_batch_id: str,
    summary: dict[str, Any],
) -> bytes:
    """Render the batch technical appendix from canonical observations."""

    snapshot_run_ids = sorted(
        {
            str(row.get("snapshot_run_id") or "").strip()
            for row in summary.get("runs") or []
            if str(row.get("snapshot_run_id") or "").strip()
        }
    )
    observations = StateRepository().list_observations_for_runs(snapshot_run_ids)
    workbook = _build_xlsx_workbook(
        report_id=patrol_batch_id,
        changes=list(summary.get("changes") or []),
        events=list(summary.get("events") or []),
        quality=dict(summary.get("quality") or {}),
        unstructured=dict(summary.get("unstructured_comparison") or {}),
        observations=observations,
        current_findings=list(summary.get("current_findings") or []),
        assessment=dict(summary.get("assessment") or {}),
        device_count=int(summary.get("device_count") or 0),
        current_state_rule_coverage=dict(summary.get("current_state_rule_coverage") or {}),
    )
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def _build_xlsx_workbook(
    *,
    report_id: str,
    changes: list[dict[str, Any]],
    events: list[dict[str, Any]],
    quality: dict[str, Any],
    unstructured: dict[str, Any],
    observations: list[dict[str, Any]],
    current_findings: list[dict[str, Any]] | None = None,
    assessment: dict[str, Any] | None = None,
    device_count: int | None = None,
    current_state_rule_coverage: dict[str, Any] | None = None,
):
    from openpyxl import Workbook
    from openpyxl.formatting.rule import ColorScaleRule
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    current_findings = list(current_findings or [])
    assessment = dict(assessment or {})
    current_state_rule_coverage = dict(current_state_rule_coverage or {})
    parser_ready_rate = float(quality.get("parser_ready_rate") or 0.0)
    fact_ready_rate = float(quality.get("fact_ready_rate", parser_ready_rate) or 0.0)
    normalization_ready_rate = float(
        quality.get("normalization_ready_rate", parser_ready_rate) or 0.0
    )
    workbook = Workbook()
    overview = workbook.active
    overview.title = "总览"
    title_fill = PatternFill("solid", fgColor="263238")
    header_fill = PatternFill("solid", fgColor="365F4B")
    soft_fill = PatternFill("solid", fgColor="E8F0EB")
    white_font = Font(color="FFFFFF", bold=True)
    overview["A1"] = "巡检报告摘要"
    overview["A1"].font = Font(color="FFFFFF", bold=True, size=16)
    overview["A1"].fill = title_fill
    overview["B1"] = report_id
    overview["B1"].font = Font(color="FFFFFF")
    overview["B1"].fill = title_fill
    overview.column_dimensions["A"].width = 30
    overview.column_dimensions["B"].width = 28
    metrics: list[tuple[str, Any]] = [
        ("巡检批次/运行ID", report_id),
    ]
    if assessment:
        metrics.extend(
            [
                ("总体结论", assessment.get("status") or "未形成结论"),
                ("结论可信度", assessment.get("trust") or "未知"),
                ("判断依据", assessment.get("reason") or ""),
            ]
        )
    if device_count is not None:
        metrics.append(("设备数", device_count))
    metrics.extend(
        [
            ("变化数", len(changes)),
            ("事件数", len(events)),
            ("当前异常项", len(current_findings)),
            ("结构化观测数", len(observations)),
            (
                "采集/质量状态",
                quality.get("collection_status") or quality.get("status") or "unknown",
            ),
            ("Parser就绪率", parser_ready_rate),
            ("事实就绪率", fact_ready_rate),
            ("结构化就绪率", normalization_ready_rate),
            ("当前状态规则评估率", current_state_rule_coverage.get("evaluation_rate") or 0.0),
            ("缺少模板", quality.get("template_missing_count") or 0),
            ("Parser待治理", quality.get("parser_backlog_count") or 0),
            ("合法空输出", quality.get("valid_empty_count") or 0),
            ("未确认空输出", quality.get("empty_unverified_count") or 0),
            ("不支持命令", quality.get("unsupported_command_count") or 0),
            ("非结构化变更", unstructured.get("changed_output_count") or 0),
            ("非结构化新增行", unstructured.get("added_line_count") or 0),
            ("非结构化删除行", unstructured.get("removed_line_count") or 0),
        ]
    )
    for row_index, (label, value) in enumerate(metrics, start=3):
        overview.cell(row_index, 1, label)
        overview.cell(row_index, 2, _safe_excel_value(value))
        overview.cell(row_index, 1).font = Font(bold=True)
        if row_index % 2:
            overview.cell(row_index, 1).fill = soft_fill
            overview.cell(row_index, 2).fill = soft_fill
        if label.endswith("率"):
            overview.cell(row_index, 2).number_format = "0.0%"
    overview.freeze_panes = "A3"

    tables: list[ReportTable] = [
        ReportTable(
            key="current_findings",
            sheet_name="当前风险",
            headers=(
                "严重度",
                "风险类型",
                "设备ID",
                "设备",
                "实体类型",
                "实体",
                "字段",
                "当前值（B）",
                "单位",
                "规则ID",
                "证据ID",
            ),
            rows=tuple(
                (
                    finding.get("severity"),
                    finding.get("event_type"),
                    finding.get("device_id"),
                    finding.get("device_name"),
                    finding.get("entity_type"),
                    finding.get("entity_key"),
                    finding.get("field"),
                    finding.get("value", finding.get("current_value")),
                    finding.get("unit"),
                    finding.get("rule_id"),
                    finding.get("evidence_id"),
                )
                for finding in current_findings
            ),
        ),
        ReportTable(
            key="events",
            sheet_name="事件",
            headers=("时间", "严重度", "事件类型", "设备ID", "实体", "状态", "证据ID"),
            rows=tuple(
                (
                    event.get("occurred_at"),
                    event.get("severity"),
                    event.get("event_type"),
                    event.get("device_id"),
                    event.get("entity_key"),
                    event.get("lifecycle_action") or event.get("status"),
                    event.get("evidence_id"),
                )
                for event in events
            ),
        ),
        ReportTable(
            key="changes",
            sheet_name="变化与基线对比",
            headers=(
                "时间",
                "设备ID",
                "实体类型",
                "实体",
                "标准观测类型",
                "字段",
                "基线值",
                "当前值",
                "比较结果",
            ),
            rows=tuple(
                (
                    change.get("detected_at") or change.get("observed_at"),
                    change.get("device_id"),
                    change.get("entity_type"),
                    change.get("entity_key"),
                    change.get("observation_type"),
                    change.get("field_name") or change.get("field"),
                    (
                        change.get("previous_value")
                        if "previous_value" in change
                        else change.get("old_value")
                    ),
                    (
                        change.get("current_value")
                        if "current_value" in change
                        else change.get("new_value")
                    ),
                    change.get("decision") or change.get("change_type"),
                )
                for change in changes
            ),
        ),
        *[
            _enrich_domain_table(
                table,
                events=events,
                changes=changes,
                current_findings=current_findings,
                rule_coverage_available=bool(current_state_rule_coverage.get("available")),
            )
            for table in build_observation_report_tables(observations)
        ],
        ReportTable(
            key="unstructured",
            sheet_name="非结构化变更",
            headers=("设备", "命令", "新增行", "删除行", "分类", "安全预览"),
            rows=tuple(
                (
                    item.get("device_name"),
                    item.get("command"),
                    item.get("added_line_count") or 0,
                    item.get("removed_line_count") or 0,
                    _unstructured_categories(item),
                    _unstructured_preview(item),
                )
                for item in unstructured.get("items") or []
            ),
        ),
        ReportTable(
            key="data_quality",
            sheet_name="数据质量与待确认",
            headers=("项目", "值", "说明"),
            rows=_data_quality_rows(
                quality=quality,
                rule_coverage=current_state_rule_coverage,
            ),
        ),
    ]
    for table in tables:
        _append_report_table(
            workbook,
            table,
            header_fill=header_fill,
            white_font=white_font,
            alignment=Alignment(vertical="top", wrap_text=True),
            get_column_letter=get_column_letter,
            color_scale_rule=ColorScaleRule,
        )
    return workbook


def _data_quality_rows(
    *,
    quality: dict[str, Any],
    rule_coverage: dict[str, Any],
) -> tuple[tuple[Any, ...], ...]:
    """Keep incomplete evidence visible instead of implying a healthy result."""

    rows = [
        (
            "数据质量状态",
            quality.get("collection_status") or quality.get("status") or "unknown",
            "采集、解析与标准化结果的总体状态。",
        ),
        (
            "缺少 Parser 模板",
            int(quality.get("template_missing_count") or 0),
            "缺少模板的命令不能形成结构化 B 观测。",
        ),
        (
            "Parser 解析失败",
            int(quality.get("parse_failed_count") or 0),
            "解析失败的命令仅保留原始证据，不能用于结构化健康判断。",
        ),
        (
            "未确认空输出",
            int(quality.get("empty_unverified_count") or 0),
            "空输出尚未确认是否为不适用，不能据此判断设备健康。",
        ),
        (
            "阈值待确认观测",
            int(rule_coverage.get("threshold_unknown_observation_count") or 0),
            "未取得设备或模块阈值的观测不使用统一固定阈值。",
        ),
        (
            "基线/健康策略待确认观测",
            int(rule_coverage.get("assessment_policy_missing_observation_count") or 0),
            "缺少可用基线或健康策略时，结果保留为待确认。",
        ),
        (
            "规则评估观测数",
            int(rule_coverage.get("evaluated_observation_count") or 0),
            "已按本次冻结或当前活动规则完成健康评估的标准观测数。",
        ),
    ]
    return tuple(rows)


def _safe_excel_value(value: Any) -> Any:
    if value is None or isinstance(value, (bool, int, float)):
        return value
    text = str(value)
    if len(text) > 32_767:
        text = text[:32_764] + "..."
    if text.startswith(("=", "+", "-", "@")):
        return "'" + text
    return text


def _enrich_domain_table(
    table: ReportTable,
    *,
    events: list[dict[str, Any]],
    changes: list[dict[str, Any]],
    current_findings: list[dict[str, Any]],
    rule_coverage_available: bool,
) -> ReportTable:
    if table.key in {
        "evidence_index",
        "other_structured",
        "canonical_field_dictionary",
    }:
        return table

    def entity_key(value: Any) -> str:
        text = str(value or "").strip()
        return text or "__singleton__"

    event_index: dict[tuple[str, str], list[dict[str, Any]]] = {}
    for event in events:
        event_index.setdefault(
            (
                str(event.get("device_id") or ""),
                entity_key(event.get("entity_key")),
            ),
            [],
        ).append(event)
    change_index: Counter[tuple[str, str]] = Counter(
        (
            str(change.get("device_id") or ""),
            entity_key(change.get("entity_key")),
        )
        for change in changes
    )
    finding_index: dict[tuple[str, str], list[dict[str, Any]]] = {}
    for finding in current_findings:
        finding_index.setdefault(
            (
                str(finding.get("device_id") or ""),
                entity_key(finding.get("entity_key")),
            ),
            [],
        ).append(finding)

    enriched_rows: list[tuple[Any, ...]] = []
    for row in table.rows:
        device_id = str(row[0] or "")
        entity = entity_key(row[4])
        related = event_index.get((device_id, entity), [])
        related_findings = finding_index.get((device_id, entity), [])
        if related_findings:
            current_summary = "；".join(
                f"{finding.get('severity') or '异常'}:{finding.get('event_type') or 'Unknown'}"
                for finding in related_findings[:3]
            )
            if len(related_findings) > 3:
                current_summary += f"；另{len(related_findings) - 3}项"
        elif rule_coverage_available:
            current_summary = "覆盖范围内未发现当前异常项"
        else:
            current_summary = "未判定（缺少活动规则覆盖）"
        if related:
            labels = []
            for event in related[:3]:
                action = str(
                    event.get("lifecycle_action") or event.get("event_class") or ""
                ).casefold()
                prefix = (
                    "已恢复"
                    if action in {"recovered", "recovery"}
                    else str(event.get("severity") or "事件")
                )
                labels.append(f"{prefix}:{event.get('event_type') or 'Unknown'}")
            event_summary = "；".join(labels)
            if len(related) > 3:
                event_summary += f"；另{len(related) - 3}项"
        else:
            event_summary = "未发现关联事件"
        enriched_rows.append(
            (
                *row[:-1],
                current_summary,
                event_summary,
                change_index[(device_id, entity)],
                row[-1],
            )
        )
    return ReportTable(
        key=table.key,
        sheet_name=table.sheet_name,
        headers=(
            *table.headers[:-1],
            "当前判定",
            "本次事件",
            "变化数",
            table.headers[-1],
        ),
        rows=tuple(enriched_rows),
    )


def _append_report_table(
    workbook: Any,
    table: ReportTable,
    *,
    header_fill: Any,
    white_font: Any,
    alignment: Any,
    get_column_letter: Any,
    color_scale_rule: Any,
) -> None:
    worksheet = workbook.create_sheet(table.sheet_name[:31])
    worksheet.append(list(table.headers))
    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = white_font
        cell.alignment = alignment
    for row in table.rows:
        worksheet.append([_safe_excel_value(value) for value in row])
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    worksheet.sheet_view.showGridLines = False
    for column_index, header in enumerate(table.headers, start=1):
        values = [
            str(worksheet.cell(row_index, column_index).value or "")
            for row_index in range(1, min(worksheet.max_row, 300) + 1)
        ]
        width = min(42, max(10, len(str(header)) + 2, *(len(value) + 2 for value in values)))
        worksheet.column_dimensions[get_column_letter(column_index)].width = width
        for row_index in range(2, worksheet.max_row + 1):
            worksheet.cell(row_index, column_index).alignment = alignment
    if table.key in {"cpu", "memory"} and worksheet.max_row > 1:
        for column_index, header in enumerate(table.headers, start=1):
            if "(%)" not in str(header):
                continue
            letter = get_column_letter(column_index)
            worksheet.conditional_formatting.add(
                f"{letter}2:{letter}{worksheet.max_row}",
                color_scale_rule(
                    start_type="num",
                    start_value=0,
                    start_color="DDEBDD",
                    mid_type="num",
                    mid_value=70,
                    mid_color="F4E6A1",
                    end_type="num",
                    end_value=100,
                    end_color="E6A19A",
                ),
            )


def _unstructured_categories(item: dict[str, Any]) -> str:
    return (
        ", ".join(
            str(category.get("label") or category.get("name") or "other")
            for category in item.get("categories") or []
        )
        or "uncategorized"
    )


def _unstructured_preview(item: dict[str, Any]) -> str:
    return " | ".join(
        f"{'+' if row.get('operation') == 'added' else '-'} {row.get('text') or ''}"
        for row in item.get("preview") or []
    )
