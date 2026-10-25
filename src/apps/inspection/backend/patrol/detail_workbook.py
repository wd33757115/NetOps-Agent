"""Export saved parser records and CLI text, without rerunning rules or parsers."""

import json
import math
import re
from collections import defaultdict
from difflib import SequenceMatcher
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

MAX_ROWS = 1_048_576
EDGE = Side(style="thin", color="000000")
BORDER = Border(left=EDGE, right=EDGE, top=EDGE, bottom=EDGE)


def record_sets(value):
    if isinstance(value, list):
        return {"default": value}
    if isinstance(value, dict) and isinstance(value.get("record_sets"), dict):
        return value["record_sets"]
    return {}


def _literal(value):
    if isinstance(value, (dict, list)):
        value = json.dumps(value, ensure_ascii=False)
    if isinstance(value, float) and not math.isfinite(value):
        return str(value)
    if isinstance(value, int) and abs(value) >= 10**15:
        return str(value)
    if isinstance(value, str):
        value = re.sub(
            r"[\x00-\x08\x0b\x0c\x0e-\x1f\ufffe\uffff]", lambda m: f"\\u{ord(m[0]):04x}", value
        )
        if len(value) > 32767:
            raise ValueError("结构化单元格超过 Excel 长度限制，请下载原始输出")
    return value


def _sheet_name(name, used):
    clean = re.sub(r"[\x00-\x1f\\/*?:\[\]]", "_", name).strip(" '") or "检查项"
    clean = clean[:31]
    candidate, n = clean, 2
    while candidate.casefold() in used:
        suffix = f" ({n})"
        candidate = clean[: 31 - len(suffix)] + suffix
        n += 1
    used.add(candidate.casefold())
    return candidate


def _write(ws, values, *, header=False, changed=False):
    if ws.max_row >= MAX_ROWS:
        raise ValueError("详细数据超过 Excel 单页行数限制，请缩小巡检范围")
    row = 1 if header else ws.max_row + 1
    for index, value in enumerate(values, 1):
        cell = ws.cell(row, index, _literal(value))
        if isinstance(cell.value, str):
            cell.data_type = "s"
        cell.font = Font(
            name="Microsoft YaHei" if header else "Consolas", size=10, bold=header, color="000000"
        )
        cell.border = BORDER
        cell.alignment = Alignment(
            horizontal="center" if header else "left", vertical="center", wrap_text=True
        )
        if header:
            cell.fill = PatternFill("solid", fgColor="BDD7EE")
        elif changed and index > 2:
            cell.fill = PatternFill("solid", fgColor="FFC7CE")


def _width(value):
    return max(
        (sum(2 if ord(c) > 127 else 1 for c in line) for line in value.splitlines()), default=0
    )


def _format_sheet(ws, text=False):
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = ws.dimensions
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 100
    ws.print_title_rows = "1:1"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A3
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    for col in ws.columns:
        index = col[0].column
        length = max(_width(str(c.value or "")) for c in col)
        header_width = _width(str(col[0].value or "")) * 1.3
        width = (
            85
            if text and index > 2
            else min(80, max(14, math.ceil(max(header_width, length * 1.05)) + 3))
        )
        ws.column_dimensions[get_column_letter(index)].width = width
    for row in ws:
        lines = max(
            sum(
                max(
                    1,
                    math.ceil(
                        _width(line) / max(ws.column_dimensions[c.column_letter].width - 2, 1)
                    ),
                )
                for line in str(c.value or "").split("\n")
            )
            for c in row
        )
        ws.row_dimensions[row[0].row].height = min(409, max(20, lines * 14 + 4))


def _records(snapshot, version, definition, multiple_items):
    sets = record_sets(snapshot.get("structured_json"))
    selected = (snapshot.get("detail_record_sets") or {}).get(version)
    if selected is None:
        entity = (definition.get("contract") or {}).get("entity_type")
        selected = [entity] if entity in sets else ([] if multiple_items else list(sets))
    return [
        record for name in selected for record in sets.get(name, []) if isinstance(record, dict)
    ], bool(selected)


def _text_rows(ws, base, now):
    old, current = (base or {}).get("raw_text"), (now or {}).get("raw_text")
    # Retain line endings: concatenating either side reconstructs its source.
    left, right = (old or "").splitlines(keepends=True), (current or "").splitlines(keepends=True)
    operations = SequenceMatcher(None, left, right, autojunk=False).get_opcodes()
    snapshot = now or base
    for op, a, b, c, d in operations or [("equal", 0, 0, 0, 0)]:
        for index in range(max(b - a, d - c, 1)):
            left_line = left[a + index] if a + index < b else ""
            right_line = right[c + index] if c + index < d else ""
            for offset in range(0, max(len(left_line), len(right_line), 1), 800):
                _write(
                    ws,
                    [
                        snapshot["device_id"],
                        snapshot["command"],
                        left_line[offset : offset + 800],
                        right_line[offset : offset + 800],
                    ],
                    changed=op != "equal" and old is not None and current is not None,
                )


def render_detail_workbook(report, source, data):
    """One sheet per check item, original Value names and complete current rows.

    Conclusions never gate this export. Missing bindings are visible; unbound
    CLI is retained under its command rather than discarded or falsely linked.
    """
    wb = Workbook()
    wb.remove(wb.active)
    groups, notes = {}, []
    for snapshot in data["snapshots"]:
        meta = snapshot.get("metadata_json") or {}
        versions = list(
            dict.fromkeys(
                [v for v in meta.get("check_item_version_ids", []) if v]
                + [
                    b["check_item_version_id"]
                    for b in meta.get("check_item_bindings", [])
                    if b.get("check_item_version_id")
                ]
            )
        )
        definitions = [(v, data["definitions"][v]) for v in versions if v in data["definitions"]]
        if not definitions:
            command = snapshot.get("canonical_command") or snapshot["command"]
            definitions = [
                (
                    None,
                    {
                        "id": "unbound:" + command,
                        "name": command,
                        "mode": "structured"
                        if record_sets(snapshot.get("structured_json"))
                        or meta.get("handling") == "structured"
                        else "text",
                    },
                )
            ]
            notes.append(
                [
                    snapshot["device_id"],
                    snapshot["command"],
                    command,
                    "采集时未保存可用检查项关联，按命令命名，保留原始数据。",
                ]
            )
        for version, definition in definitions:
            group = groups.setdefault(definition["id"], {"name": definition["name"], "entries": []})
            structured = definition.get("mode") == "structured"
            records, selected = _records(snapshot, version, definition, len(definitions) > 1)
            group["entries"].append((snapshot, structured, records))
            if structured and snapshot["period"] == "current" and not records:
                reason = (
                    "本期无可导出的解析记录（"
                    + str(snapshot.get("parse_status") or "未知状态")
                    + "）。"
                    if selected
                    else "缺少记录集关联或解析记录，未混入其他检查项数据。"
                )
                notes.append(
                    [snapshot["device_id"], snapshot["command"], definition["name"], reason]
                )
    used = set()
    for group in groups.values():
        entries = group["entries"]
        # Mode changes between historical versions must not mix incompatible layouts.
        for structured in dict.fromkeys(s for _, s, _ in entries):
            selected = [(s, rows) for s, mode, rows in entries if mode == structured]
            name = group["name"]
            if len({mode for _, mode, _ in entries}) > 1:
                name += "（解析数据）" if structured else "（原文）"
            ws = wb.create_sheet(_sheet_name(name, used))
            if structured:
                current = [(s, rows) for s, rows in selected if s["period"] == "current"]
                fields = list(dict.fromkeys(k for _, rows in current for row in rows for k in row))
                _write(ws, ["设备ID", "命令"] + fields, header=True)
                for snapshot, rows in current:
                    for row in rows:
                        _write(
                            ws,
                            [snapshot["device_id"], snapshot["command"]]
                            + [row.get(f) for f in fields],
                        )
                if not current:
                    for s, _ in selected:
                        notes.append(
                            [s["device_id"], s["command"], name, "本期未采集，未填入基线数据。"]
                        )
            else:
                _write(ws, ["设备ID", "命令", "基线数据", "当前数据"], header=True)
                pairs = defaultdict(lambda: defaultdict(list))
                for s, _ in selected:
                    pairs[(s["device_id"], s.get("canonical_command") or s["command"])][
                        s["period"]
                    ].append(s)
                for sides in pairs.values():
                    baseline, current = sides["baseline"], sides["current"]
                    if len(baseline) <= 1 and len(current) <= 1:
                        _text_rows(
                            ws, baseline[0] if baseline else None, current[0] if current else None
                        )
                    else:
                        for s in baseline:
                            _text_rows(ws, s, None)
                        for s in current:
                            _text_rows(ws, None, s)
                        s = (current or baseline)[0]
                        notes.append(
                            [
                                s["device_id"],
                                s["command"],
                                name,
                                "同一期存在多个快照，分别保留，未自动配对。",
                            ]
                        )
                    if not baseline or not current:
                        s = (current or baseline)[0]
                        notes.append(
                            [s["device_id"], s["command"], name, "缺少一期原文，空白不代表无变化。"]
                        )
            _format_sheet(ws, text=not structured)
    if notes or not wb.worksheets:
        ws = wb.create_sheet(_sheet_name("导出说明", used))
        _write(ws, ["设备ID", "命令", "检查项", "数据说明"], header=True)
        for row in dict.fromkeys(tuple(n) for n in notes):
            _write(ws, row)
        if not notes:
            _write(ws, [None, None, None, "本次任务没有命令快照。"])
        _format_sheet(ws)
    output = BytesIO()
    wb.save(output)
    return output.getvalue()
