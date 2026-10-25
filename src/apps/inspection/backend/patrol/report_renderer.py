"""Fill approved DOCX controls and append local CLI evidence, without any AI calls."""

import csv
import io
import json
from copy import deepcopy
from functools import lru_cache
from hashlib import sha256
from io import BytesIO
from zipfile import ZIP_DEFLATED, ZipFile

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

from src.core.patrol.report_evidence import build_excerpts
from src.core.patrol.report_findings import date_text, field_label, project_findings
from src.core.patrol.report_templates import binding_schema
from src.core.patrol.report_view import display, identity, project_report


def _replace_control(control, value):
    content = control.find(qn("w:sdtContent"))
    parent = control.getparent()
    run = OxmlElement("w:r")
    existing = next(iter(content.iter(qn("w:rPr"))), None) if content is not None else None
    if existing is not None:
        run.append(deepcopy(existing))
    for i, line in enumerate(display(value).splitlines()):
        if i:
            run.append(OxmlElement("w:br"))
        text = OxmlElement("w:t")
        text.set(qn("xml:space"), "preserve")
        text.text = line
        run.append(text)
    if parent.tag in [qn("w:body"), qn("w:tc")]:
        paragraph = OxmlElement("w:p")
        paragraph.append(run)
        parent.replace(control, paragraph)
    else:
        parent.replace(control, run)


def _tag(control):
    tag = control.find("./" + qn("w:sdtPr") + "/" + qn("w:tag"))
    return tag.get(qn("w:val"), "") if tag is not None else ""


@lru_cache(maxsize=1)
def binding_columns():
    """Freeze the user's approved column meanings; imported columns may move."""
    return binding_schema()["columns"]


def _clip_template_appendices(doc):
    # Evidence layout is rendered dynamically; remove its prototype from this point.
    for child in list(doc.element.body):
        tags = [e.get(qn("w:val"), "") for e in child.iter(qn("w:tag"))]
        text = "".join(t.text or "" for t in child.iter(qn("w:t")))
        if text.startswith("F 非结构化文本变化对照") or any(t.startswith("cli.") for t in tags):
            previous = child.getprevious()
            if previous is not None and list(previous.iter(qn("w:br"))):
                previous.getparent().remove(previous)
            while child is not None and child.tag != qn("w:sectPr"):
                nxt = child.getnext()
                child.getparent().remove(child)
                child = nxt
            break


def fill_template(
    content, report, source=None, metadata=None, drafts=None, *, bundle=False, include_details=False
):
    doc = Document(BytesIO(content))
    _clip_template_appendices(doc)
    view = project_findings(report, source, metadata, project_report(report, source, metadata))
    view["scalars"]["trace.template_version"] = "问题与变化模板"
    view["scalars"]["trace.revision"] = (metadata or {}).get("revision") or "未填写"
    view["tables"]["deliverables"] = [
        ["DOC-01", "patrol-report.docx", "本报告及 CLI 摘录", "随本次导出生成"]
    ]
    if bundle:
        view["tables"]["deliverables"] += [
            ["DATA-01", "observations.csv", "本次全部字段观测与来源索引", "随证据包提供"],
            ["CLI", "evidence/*.txt", "附录所引用的完整命令；缺失侧无文件", "随证据包提供"],
        ]
    if include_details:
        view["tables"]["deliverables"].append(
            [
                "DATA-02",
                "patrol-details.xlsx",
                "按检查项分表的完整详细数据与 CLI 对比",
                "随证据包提供",
            ]
        )
    drafts = drafts or {}
    for key in ["narrative.summary", "changes.summary"]:
        if key in drafts:
            view["scalars"][key] = drafts[key]
    for table in list(doc.tables):
        for row in list(table.rows):
            controls = list(row._tr.iter(qn("w:sdt")))
            numeric = [_tag(c) for c in controls if _tag(c).rsplit(".", 1)[-1].isdigit()]
            if not numeric:
                continue
            key = numeric[0].rsplit(".", 1)[0]
            rows = view["tables"].get(key) or [
                ["未取得对应数据或尚未确认"] + [""] * (len(row.cells) - 1)
            ]
            for values in rows:
                clone = deepcopy(row._tr)
                if key == "findings":
                    cell = clone.find(qn("w:tc"))
                    props = cell.get_or_add_tcPr()
                    for sh in list(props.findall(qn("w:shd"))):
                        props.remove(sh)
                    shade = OxmlElement("w:shd")
                    shade.set(
                        qn("w:fill"),
                        {"当前异常": "FCE4D6", "已恢复": "E2EFDA", "数据缺口": "EEEEEE"}.get(
                            values[0], "FFF2CC"
                        ),
                    )
                    props.append(shade)
                for cell in clone.findall(qn("w:tc")):
                    for c in list(cell.iter(qn("w:sdt"))):
                        index = binding_columns().get(_tag(c))
                        value = (
                            values[index]
                            if index is not None and index < len(values)
                            else view["scalars"].get(_tag(c))
                        )
                        _replace_control(c, value)
                row._tr.addprevious(clone)
            row._tr.getparent().remove(row._tr)
    for c in list(doc.element.iter(qn("w:sdt"))):
        _replace_control(c, view["scalars"].get(_tag(c), "未提供或待人工确认"))
    return doc, view


def _line(cell, value, change, side):
    p = cell.paragraphs[0]
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_after = Pt(2)
    if not value:
        p.add_run("—（无对应原文行）")
        return
    p.add_run(f"{value['line']:>4}  ")
    text = value["text"]
    spans = value.get("anomaly_spans", []) + (
        [value["anomaly_span"]] if value.get("anomaly_span") else []
    )
    bounds = sorted({0, len(text), *[x for s in spans for x in s]})
    for a, b in zip(bounds, bounds[1:]):
        run = p.add_run(text[a:b])
        run.font.name = "Consolas"
        run.font.size = Pt(8.5)
        if any(a >= s and b <= e for s, e in spans):
            run.font.color.rgb = RGBColor(192, 0, 0)
            run.bold = True
        elif change == "delete" and side == "baseline":
            run.font.color.rgb = RGBColor(100, 100, 100)
            run.font.strike = True
        elif change == "insert" and side == "current":
            run.font.color.rgb = RGBColor(31, 78, 121)
        elif change == "replace":
            run.font.color.rgb = RGBColor(156, 101, 0)


def append_evidence(doc, excerpts):
    previous_groups = None
    for kind, heading in [
        ("text", "D 非结构化文本变化对照"),
        ("event", "E 结构化异常与事件原文核对"),
    ]:
        if previous_groups is None or previous_groups:
            doc.add_page_break()
        doc.add_heading(heading, level=2)
        groups = [e for e in excerpts if e["kind"] == kind]
        previous_groups = groups
        if not groups:
            doc.add_paragraph("本次未形成此类可展示证据；无事件或未比较不等同于健康。")
        for entry in groups:
            doc.add_heading(f"{entry['id']}  {entry['device']}", level=3)
            description = doc.add_paragraph(
                f"命令：{entry['command']}\n对象：{identity(entry['entity'])}；"
                f"检查字段：{'、'.join(field_label(f) for f in entry['fields'])}"
            )
            description.paragraph_format.keep_with_next = True
            note = doc.add_paragraph(entry["note"])
            note.paragraph_format.keep_with_next = True
            table = doc.add_table(rows=1, cols=2)
            table.style = "Table Grid"
            table.autofit = False
            section = doc.sections[-1]
            column_width = int(
                (section.page_width - section.left_margin - section.right_margin) / 2
            )
            for col in table.columns:
                col.width = column_width
            for side, c in zip(["baseline", "current"], table.rows[0].cells):
                c.width = column_width
                c.text = (
                    f"{'基线' if side == 'baseline' else '当前'} CLI 原文\n"
                    f"采集时间：{date_text(entry[side].get('observed_at'))}"
                )
                sh = OxmlElement("w:shd")
                sh.set(qn("w:fill"), "D9E2F3")
                c._tc.get_or_add_tcPr().append(sh)
            table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
            rows = entry["rows"]
            last = {"baseline": 0, "current": 0}
            for row in rows:
                if any(row.get(s) and last[s] and row[s]["line"] > last[s] + 1 for s in last):
                    gap = table.add_row()
                    for c in gap.cells:
                        c.text = "……省略其他原文行……"
                cells = table.add_row().cells
                for side, cell in zip(["baseline", "current"], cells):
                    cell.width = column_width
                    _line(cell, row.get(side), row["change"], side)
                    if row.get(side):
                        last[side] = row[side]["line"]
            if not rows:
                for c in table.add_row().cells:
                    c.text = "未能生成可靠摘录，请核对完整命令证据。"
            # Keep one original CLI line intact when a comparison spans pages.
            for evidence_row in table.rows:
                evidence_row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
                if len(table.rows) <= 20:
                    for cell in evidence_row.cells:
                        for p in cell.paragraphs:
                            p.paragraph_format.keep_with_next = True
            doc.add_paragraph(
                f"完整原文：导出证据包内 evidence/{entry['id']}-baseline.txt "
                f"和 {entry['id']}-current.txt（缺失侧不提供文件）。"
            )
    return doc


def render_report(
    template, report, source=None, metadata=None, drafts=None, *, bundle=False, detail_workbook=None
):
    doc, view = fill_template(
        template,
        report,
        source,
        metadata,
        drafts,
        bundle=bundle,
        include_details=bundle and detail_workbook is not None,
    )
    excerpts = build_excerpts(report)
    # Only claim bundled attachments when this actual export contains them.
    append_evidence(doc, excerpts)
    if not bundle:
        for p in doc.paragraphs:
            if p.text.startswith("完整原文：导出证据包"):
                p.text = (
                    "完整原文请在本任务下载原始输出，或使用“导出报告与证据包”；"
                    "本 Word 未内嵌全文附件。"
                )
    output = BytesIO()
    doc.save(output)
    if not bundle:
        return output.getvalue()
    result = BytesIO()
    with ZipFile(result, "w", ZIP_DEFLATED) as z:
        z.writestr("patrol-report.docx", output.getvalue())
        if detail_workbook is not None:
            z.writestr("patrol-details.xlsx", detail_workbook)
        manifest = {
            "schema_version": "patrol-report-package/v1",
            "template_sha256": sha256(template).hexdigest(),
            "analysis": (report.get("evidence") or {}).get("analysis"),
            "observations_digest": report.get("source_observations_digest"),
            "ai_sections_adopted": sorted((drafts or {}).keys()),
            "evidence": [],
        }
        for entry in excerpts:
            references = {
                "id": entry["id"],
                "references": entry["references"],
                "device": entry["device"],
                "command": entry["command"],
                "entity": entry["entity"],
                "files": {},
            }
            for side, text in entry["full_source"].items():
                if text is not None:
                    name = f"evidence/{entry['id']}-{side}.txt"
                    z.writestr(name, text)
                    references["files"][side] = {
                        "path": name,
                        "sha256": sha256(text.encode()).hexdigest(),
                        "snapshot_id": entry[side].get("snapshot_id"),
                    }
            manifest["evidence"].append(references)
        csv_output = io.StringIO()
        writer = csv.writer(csv_output)
        headers = [
            "device_id",
            "entity_type",
            "entity_key",
            "field",
            "value",
            "unit",
            "command",
            "snapshot_id",
            "evidence_id",
            "observed_at",
        ]
        writer.writerow(headers)
        for obs in view["observations"]:
            # Prevent spreadsheet formula execution in user-controlled CLI fields.
            writer.writerow(
                [
                    ("'" + str(obs.get(k)))
                    if str(obs.get(k) or "").startswith(("=", "+", "-", "@", "\t", "\r"))
                    else display(obs.get(k))
                    for k in headers
                ]
            )
        z.writestr("observations.csv", csv_output.getvalue().encode("utf-8-sig"))
        z.writestr("manifest.json", json.dumps(manifest, ensure_ascii=False, indent=2, default=str))
    return result.getvalue()
