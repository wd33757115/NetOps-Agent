"""Deterministic patrol deliverables rendered from one immutable report manifest."""

from __future__ import annotations

import csv
import hashlib
import io
import json
import re
import unicodedata
import zipfile
from collections import Counter
from collections.abc import Iterable, Mapping, Sequence
from dataclasses import dataclass
from datetime import datetime, timezone
from typing import Any

from src.core.check_items.presentation import check_item_display_name
from src.core.check_items.report_claims import (
    validate_report_claim_ledger,
)
from src.core.check_items.report_contract import (
    CURRENT_STATE_AUTHORITATIVE_PURPOSES,
    OFFLINE_HISTORICAL,
    PARSER_VALIDATION,
    PRODUCTION_CURRENT,
    RULE_ASSET_REPLAY,
    SHADOW_VALIDATION,
    build_report_conclusion_permission_from_source_facts,
    normalize_report_purpose,
    raw_evidence_set_digest,
    report_authority_context_is_valid,
    report_conclusion_permission_is_valid,
)
from src.core.check_items.report_projection import (
    project_report_tables,
    validate_report_projection_snapshot,
)
from src.core.check_items.risk_issue_projection import (
    build_risk_issue_projection,
)

_ARCHIVE_TIMESTAMP = (1980, 1, 1, 0, 0, 0)
_SEVERITY_ORDER = {
    "critical": 5,
    "major": 4,
    "warning": 3,
    "minor": 2,
    "info": 1,
    "unknown": 0,
}
_STATUS_LABELS = {
    "healthy": "正常",
    "at_risk": "存在风险",
    "unable_to_determine": "无法判断",
    "not_applicable": "不适用",
    "informational": "信息性结果",
}
_COMPLETENESS_LABELS = {
    "complete": "完整",
    "partial": "部分完整",
    "insufficient": "证据不足",
}
_COMPARISON_LABELS = {
    "comparable": "可比较",
    "first_observation": "首次观测",
    "baseline_missing": "缺少基线",
    "baseline_stale": "基线过期",
    "schema_incompatible": "口径不兼容",
}
_SEVERITY_LABELS = {
    "critical": "严重",
    "major": "重大",
    "warning": "告警",
    "minor": "提示",
    "info": "信息",
    "unknown": "未知",
}
_EXECUTION_STATUS_LABELS = {
    "planned": "已计划",
    "pending": "待执行",
    "queued": "排队中",
    "running": "执行中",
    "processing": "处理中",
    "completed": "已完成",
    "failed": "失败",
    "collection_failed": "采集失败",
    "parse_failed": "Parser 解析失败",
    "normalization_failed": "状态标准化失败",
    "comparison_failed": "变化比较失败",
    "assessment_failed": "检查项评估失败",
    "blocked_asset_missing": "规则资产缺失",
    "unsupported": "设备不支持",
    "not_applicable": "不适用",
    "cancelled": "已取消",
    "skipped": "已跳过",
}
_REPORT_STATUS_LABELS = {
    "partial": "事实汇集中",
    "final": "事实已冻结",
    "revised": "事实已修订",
}
_SOURCE_IDENTITY_LABELS = {
    "confirmed": "身份已确认",
    "ambiguous": "身份有歧义",
    "conflict": "身份冲突",
    "conflicted": "身份冲突",
    "external_unbound": "外部设备未绑定",
    "unknown": "身份未知",
}
_OUTPUT_QUALITY_LABELS = {
    "complete": "输出完整",
    "partial": "输出不完整",
    "truncated": "输出被截断",
    "paged": "输出含分页",
    "device_error": "设备返回错误",
    "empty": "输出为空",
    "unknown": "质量未知",
}
_COLLECTION_OUTPUT_STATUS_LABELS = {
    "success": "采集成功",
    "empty": "采集成功但输出为空",
    "failed": "采集失败",
    "timeout": "采集超时",
    "command_error": "设备命令错误",
    "permission_denied": "权限不足",
    "skipped": "已跳过",
    "unknown": "采集状态未知",
}
_FALLBACK_RESILIENCE_LABELS = {
    "full": "完整冗余",
    "reduced": "韧性降低",
    "none": "无可用路径",
    "not_applicable": "不适用",
    "unknown": "状态未知",
}
_PARSE_STATUS_LABELS = {
    "parsed": "解析成功",
    "template_missing": "缺少 Parser 模板",
    "parse_failed": "解析失败",
    "ignored": "无需结构化解析",
    "empty": "无可解析记录",
    "unknown": "解析状态未知",
}
_FRESHNESS_LABELS = {
    "fresh": "时效合格",
    "stale": "数据已过期",
    "unknown": "时效未知",
}
_COVERAGE_STATUS_LABELS = {
    "complete": "覆盖完整",
    "partial": "部分覆盖",
    "insufficient": "覆盖证据不足",
    "not_evaluable": "无法评估覆盖",
    "unknown": "覆盖未知",
}
_COVERAGE_BASIS_LABELS = {
    "fixed": "固定分母",
    "inventory": "资产清单分母",
    "inventory_entity_set": "资产清单实体分母",
    "derived_entity_set": "历史实体分母",
    "command_reported_total": "命令报告总数分母",
    "non_exhaustive": "非穷举结果",
    "previous_qualified_observation": "历史合格观测分母",
    "unbounded": "无确定分母",
    "unknown": "分母未知",
}
_BLOCKING_REASON_LABELS = {
    "blocked_asset_missing": "运行规则资产缺失",
    "collection_failed": "命令采集失败",
    "template_missing": "缺少 Parser 模板",
    "parser_template_missing": "缺少 Parser 模板",
    "parser_compile_failed": "Parser 模板编译失败",
    "parse_failed": "Parser 解析失败",
    "parser_failed": "Parser 解析失败",
    "normalization_failed": "状态标准化失败",
    "comparison_failed": "变化比较失败",
    "assessment_failed": "检查项评估失败",
    "coverage_not_evaluable": "无法评估实体覆盖",
    "coverage_empty": "未观测到业务实体",
    "coverage_partial": "仅覆盖部分业务实体",
    "coverage_unknown": "实体覆盖范围未知",
    "no_observations": "未形成结构化观测",
    "no_observed_entities": "未识别到业务实体",
    "no_entities_evaluated": "没有实体完成规则评估",
    "current_state_rules_not_ready": "当前状态规则未就绪",
    "evidence_qualification_missing": "缺少原始证据资质判定",
    "output_quality_empty": "命令输出为空",
    "output_quality_truncated": "命令输出被截断",
    "output_quality_paged": "命令输出含未完成分页",
    "output_quality_device_error": "设备返回命令错误",
    "output_quality_ambiguous": "命令输出完整性无法确认",
    "output_quality_unknown": "命令输出质量未知",
    "source_identity_ambiguous": "设备身份有歧义",
    "source_identity_conflict": "设备身份冲突",
    "source_identity_external_unbound": "外部设备未绑定资产",
    "source_identity_unknown": "设备身份未知",
    "freshness_stale": "巡检事实已过期",
    "freshness_unknown": "巡检事实时效未知",
    "entity_coverage_incomplete": "业务实体覆盖不完整",
    "entity_keys_missing": "存在缺失业务实体",
    "entity_keys_unexpected": "存在计划外业务实体",
    "entity_evaluation_incomplete": "部分业务实体未完成规则评估",
    "threshold_unknown": "缺少可审计阈值",
    "threshold_coverage_partial": "仅部分实体具备可审计阈值",
    "applicability_rule_no_match": "检查项不适用于该设备",
    "asset_unavailable_at_plan_time": "冻结运行资产不支持该命令",
    "assessment_policy_missing": "缺少可审计的评估策略来源",
    "assessment_policy_conflict": "同优先级评估策略相互冲突",
    "assessment_policy_source_invalid": "评估策略来源未通过审批或完整性校验",
    "assessment_policy_snapshot_invalid": "冻结评估策略快照无效",
    "assessment_policy_snapshot_mismatch": "评估使用的策略快照与执行快照不一致",
    "assessment_policy_unit_mismatch": "观测单位与评估策略单位不一致",
    "expected_state_missing": "缺少该设备或实体的期望状态",
    "assessment_basis_missing": "缺少可追溯的判断依据",
    "assessment_basis_invalid": "判断依据完整性校验失败",
    "assessment_basis_policy_snapshot_mismatch": "判断依据与冻结策略快照不一致",
}
_ASSESSMENT_SOURCE_LABELS = {
    "device_output": "设备输出阈值",
    "inventory_profile": "设备或模块档案",
    "scope_override": "本次范围策略",
    "catalog_default": "检查项目录缺省值",
    "approved_rule_constant": "旧版规则常量",
}
_ASSESSMENT_BASIS_STATUS_LABELS = {
    "resolved": "已确定",
    "missing": "缺失",
    "conflict": "冲突",
    "unknown": "未知",
}
_RISK_STATUS_LABELS = {
    "open": "待处理",
    "acknowledged": "已确认",
    "recovering": "恢复观察中",
    "resolved": "技术已恢复",
    "suppressed": "已抑制",
}
_REMEDIATION_STATUS_LABELS = {
    "not_started": "未开始",
    "planned": "已计划",
    "in_progress": "处理中",
    "completed": "已完成",
    "accepted": "验收通过",
    "verified": "已验证",
    "cancelled": "已取消",
}
_RISK_CLOSURE_LABELS = {
    "open": "风险未关闭",
    "technical_recovered": "技术已恢复，待关闭",
    "pending_remediation": "技术已恢复，整改未完成",
    "pending_acceptance": "整改完成，待验收",
    "ready_to_close": "满足正式关闭条件",
    "closed": "风险已正式关闭",
    "suppressed": "风险已抑制",
}
_EMPTY_RESULT_LABELS = {
    "not_empty": "有有效数据",
    "expected_empty": "预期为空",
    "not_applicable": "不适用",
    "invalid": "空结果无效",
    "not_evaluated": "未评估",
    "unknown": "未知",
}
_RISK_VERIFICATION_LABELS = {
    "confirmed_current": "本批次已确认",
    "no_current_risk": "本批次未确认风险",
    "not_reverified": "历史风险待复核",
    "recovery_pending": "恢复待确认",
    "resolved": "已恢复",
    "not_applicable": "不适用",
    "unknown": "未知",
}
_REFERENCE_SOURCE_LABELS = {
    "device_output": "设备输出阈值",
    "event_rule_static": "事件规则静态阈值",
    "policy_static": "变化策略静态阈值",
    "policy_fallback": "变化策略兜底阈值",
}
_REPORT_PURPOSE_LABELS = {
    PRODUCTION_CURRENT: "当前周期生产巡检",
    OFFLINE_HISTORICAL: "历史离线巡检",
    SHADOW_VALIDATION: "影子验证",
    PARSER_VALIDATION: "Parser 验证",
    RULE_ASSET_REPLAY: "规则资产回放",
}
_REPORT_PURPOSE_DISCLOSURES = {
    PRODUCTION_CURRENT: (
        "本报告仅在身份、时效、完整性、实体覆盖和规则证据均满足时，"
        "对对应设备与检查项形成当前状态结论。"
    ),
    OFFLINE_HISTORICAL: (
        "本报告反映原始数据采集时点的历史网络状态，不代表报告生成时或当前时刻的网络状态。"
    ),
    SHADOW_VALIDATION: (
        "本报告用于影子链路和资产验证，不作为生产巡检结论，也不得用于声明当前网络健康。"
    ),
    PARSER_VALIDATION: ("本报告仅用于 Parser 解析能力验证，不形成网络健康、风险或整改结论。"),
    RULE_ASSET_REPLAY: ("本报告仅用于规则资产历史回放，不形成当前网络状态结论。"),
}


def _build_conclusion_dimensions(
    *,
    formal: bool,
    report_purpose: str,
    current_state_authoritative: bool,
    matrix_count: int,
    health_assessments: Sequence[Mapping[str, Any]],
    status_counts: Counter[str],
    completeness_counts: Counter[str],
    comparison_counts: Counter[str],
) -> dict[str, dict[str, Any]]:
    if status_counts["at_risk"]:
        risk_code = "confirmed_risk"
        risk_label = "已确认风险"
        risk_basis = f"本次事实中有 {status_counts['at_risk']} 项检查形成风险结论。"
    elif (
        formal
        and health_assessments
        and not status_counts["unable_to_determine"]
        and any(item.get("assessment_status") == "healthy" for item in health_assessments)
        and all(
            item.get("assessment_status") in {"healthy", "not_applicable"}
            for item in health_assessments
        )
    ):
        risk_code = "no_confirmed_risk"
        risk_label = "未确认风险"
        risk_basis = "已完成且可判定的健康检查未形成当前风险；该结论不外推到未纳入或不适用范围。"
    else:
        risk_code = "unknown"
        risk_label = "风险未知"
        risk_basis = "当前事实不足以形成无风险结论，需结合无法判断项和覆盖范围复核。"

    if (
        formal
        and matrix_count > 0
        and completeness_counts["complete"] == matrix_count
        and not status_counts["unable_to_determine"]
    ):
        completeness_code = "complete"
        completeness_label = "评估完整"
    elif matrix_count > 0 and (completeness_counts["complete"] or completeness_counts["partial"]):
        completeness_code = "partial"
        completeness_label = "部分完整"
    else:
        completeness_code = "insufficient"
        completeness_label = "证据不足"
    completeness_basis = (
        f"完整 {completeness_counts['complete']} 项，"
        f"部分完整 {completeness_counts['partial']} 项，"
        f"证据不足 {completeness_counts['insufficient']} 项。"
    )

    comparable_count = comparison_counts["comparable"]
    blocked_comparison_count = (
        comparison_counts["baseline_stale"] + comparison_counts["schema_incompatible"]
    )
    if matrix_count > 0 and comparable_count == matrix_count:
        comparison_code = "comparable"
        comparison_label = "全部可比较"
    elif comparable_count > 0:
        comparison_code = "partial"
        comparison_label = "部分可比较"
    elif blocked_comparison_count > 0:
        comparison_code = "blocked"
        comparison_label = "比较受阻"
    else:
        comparison_code = "no_baseline"
        comparison_label = "暂无可用基线"
    no_baseline_count = (
        comparison_counts["first_observation"]
        + comparison_counts["baseline_missing"]
        + comparison_counts["no_baseline"]
    )
    comparison_basis = (
        f"可比较 {comparable_count} 项；"
        f"首次观测或缺少基线 {no_baseline_count} 项；"
        f"基线过期或口径不兼容 {blocked_comparison_count} 项。"
    )

    if not formal:
        authority_code = "draft"
        authority_label = "过程数据"
    elif report_purpose == PRODUCTION_CURRENT and current_state_authoritative:
        authority_code = "production_current"
        authority_label = "当前生产结论"
    elif report_purpose == PRODUCTION_CURRENT:
        authority_code = "blocked"
        authority_label = "当前结论权限未成立"
    elif report_purpose == OFFLINE_HISTORICAL:
        authority_code = "historical"
        authority_label = "历史采集时点结论"
    elif report_purpose == SHADOW_VALIDATION:
        authority_code = "shadow"
        authority_label = "影子验证"
    else:
        authority_code = "validation"
        authority_label = "验证用途"
    authority_basis = _REPORT_PURPOSE_DISCLOSURES[report_purpose]
    if authority_code == "blocked":
        authority_basis = (
            "报告用途虽为当前生产巡检，但来源证明、冻结摘要或终态"
            "门禁未通过，不能形成当前网络状态结论。"
        )

    return {
        "risk_conclusion": {
            "code": risk_code,
            "label": risk_label,
            "basis": risk_basis,
        },
        "assessment_completeness": {
            "code": completeness_code,
            "label": completeness_label,
            "basis": completeness_basis,
        },
        "comparison_readiness": {
            "code": comparison_code,
            "label": comparison_label,
            "basis": comparison_basis,
        },
        "authority": {
            "code": authority_code,
            "label": authority_label,
            "basis": authority_basis,
        },
    }


def _conclusion_from_permission(
    permission: Mapping[str, Any],
) -> tuple[str, str, str, str]:
    counts = _as_dict(permission.get("counts"))
    at_risk = int(counts.get("at_risk") or 0)
    unable = int(counts.get("unable_to_determine") or 0)
    healthy = int(counts.get("healthy") or 0)
    health_total = int(counts.get("health_target_pair_count") or 0)
    not_applicable = int(counts.get("not_applicable") or 0)
    unresolved = int(counts.get("unresolved_in_scope_risk_case_count") or 0)
    blockers = ", ".join(str(code) for code in permission.get("blocker_reason_codes") or [])
    if not permission.get("current_conclusion_ready"):
        return (
            "current_authority_blocked",
            "当前网络状态无法形成正式结论",
            "冻结事实未取得当前结论权限。阻断原因："
            f"{blockers or 'conclusion_permission_unavailable'}。",
            "受限结论",
        )
    if at_risk and unable:
        return (
            "at_risk_with_unknown",
            "本次巡检发现当前风险，且部分检查项无法判断",
            f"已确认 {at_risk} 项当前风险；另有 {unable} 项"
            "无法判断。风险与未知范围均已保留在正式分母中。",
            "受限结论",
        )
    if at_risk:
        return (
            "at_risk",
            "本次巡检发现当前风险",
            f"已确认 {at_risk} 项当前风险，结论来自冻结的 Assessment、Finding 和 RiskCase 事实。",
            "正式结论",
        )
    if unresolved and unable:
        return (
            "unresolved_risk_with_unknown",
            "存在未关闭风险，且部分检查项无法判断",
            f"目标范围内仍有 {unresolved} 项风险未取得合格恢复证据，另有 {unable} 项检查无法判断。",
            "受限结论",
        )
    if unresolved:
        return (
            "unresolved_risk",
            "目标范围内仍存在未关闭风险",
            f"有 {unresolved} 项 RiskCase 未取得合格恢复证据；"
            "当前数值未再次触发规则不能替代风险关闭。",
            "正式结论",
        )
    if unable:
        if permission.get("evaluated_scope_no_abnormality_allowed"):
            return (
                "evaluated_scope_healthy_with_unknown",
                "已评估范围内未发现异常，以下范围无法判断",
                f"已评估且正常 {healthy} 项；无法判断 {unable} 项。"
                "未知范围未从分母删除，不能表述网络整体健康。",
                "受限结论",
            )
        return (
            "partially_unknown",
            "部分检查项无法判断",
            f"有 {unable} 项检查因采集、解析、规则、阈值、覆盖或时效证据不足而无法判断。",
            "受限结论",
        )
    if health_total > 0 and not_applicable == health_total:
        return (
            "no_applicable_health_checks",
            "本次无适用健康检查项",
            "全部健康类目标均由冻结适用性证据判定为不适用，因此不能使用“未发现异常”或“整体健康”。",
            "信息性结果",
        )
    if permission.get("overall_healthy_allowed"):
        return (
            "healthy",
            "本次巡检范围整体健康",
            f"{health_total} 项健康类目标均已终态对账，不存在无法判断、当前风险或未关闭历史风险。",
            "正式结论",
        )
    if permission.get("evaluated_scope_no_abnormality_allowed"):
        return (
            "evaluated_scope_no_abnormality",
            "已评估范围内未发现异常",
            f"已评估且正常 {healthy} 项；该表述不外推到未纳入或不适用范围。",
            "正式结论",
        )
    if health_total == 0:
        return (
            "informational_only",
            "本次仅形成信息性结果，未判定网络健康",
            "冻结目标范围内没有承担健康结论的检查项。",
            "信息性结果",
        )
    return (
        "mixed",
        "本次检查结果需工程复核",
        "冻结事实未满足任何更高等级的结论表达条件。",
        "受限结论",
    )


@dataclass(frozen=True, slots=True)
class ReportArtifact:
    """One immutable report artifact ready for object storage."""

    name: str
    filename: str
    content_type: str
    content: bytes

    @property
    def sha256(self) -> str:
        return hashlib.sha256(self.content).hexdigest()


class CheckItemReportMaterializer:
    """Atomically materialize every formal artifact from one manifest."""

    def __init__(
        self,
        *,
        manifest_service: Any | None = None,
        object_storage: Any | None = None,
    ) -> None:
        if manifest_service is None:
            from src.core.check_items.report_manifest import (
                ReportManifestService,
            )

            manifest_service = ReportManifestService()
        if object_storage is None:
            from src.infrastructure.storage.minio_client import (
                get_minio_storage,
            )

            object_storage = get_minio_storage()
        self._manifest_service = manifest_service
        self._object_storage = object_storage

    def materialize(
        self,
        patrol_batch_id: str,
        *,
        generated_by: str,
    ) -> dict[str, Any]:
        manifest = self._manifest_service.build(
            patrol_batch_id,
            generated_by=generated_by,
        )
        if str(manifest.get("status") or "") != "final":
            reasons = (
                _as_dict(_as_dict(manifest.get("source_facts_json")).get("readiness")).get(
                    "reason_codes"
                )
                or []
            )
            raise ValueError(
                "formal report is blocked by non-final facts: "
                + ",".join(str(item) for item in reasons)
            )
        existing_keys = _as_dict(manifest.get("artifact_object_keys_json"))
        existing_digests = _as_dict(manifest.get("artifact_sha256_json"))
        if existing_keys or existing_digests:
            expected_names = set(_ARTIFACT_CONTENT_TYPES)
            if set(existing_keys) != expected_names or set(existing_digests) != expected_names:
                raise RuntimeError("attached report artifact set is incomplete")
            artifact_bytes: dict[str, int] = {}
            for name in sorted(expected_names):
                content = self._object_storage.download_file(str(existing_keys[name]))
                if content is None:
                    raise RuntimeError(f"attached report artifact is unavailable: {name}")
                actual_digest = hashlib.sha256(content).hexdigest()
                if actual_digest != str(existing_digests[name]).lower():
                    raise RuntimeError(f"attached report artifact digest mismatch: {name}")
                artifact_bytes[name] = len(content)
            return {
                "success": True,
                "reused": True,
                "patrol_batch_id": patrol_batch_id,
                "report_manifest_id": str(manifest["id"]),
                "source_digest": str(manifest["source_digest"]),
                "artifact_object_keys": existing_keys,
                "artifact_sha256": existing_digests,
                "artifact_bytes": artifact_bytes,
                "manifest": manifest,
            }
        artifacts = render_check_item_artifacts(manifest)
        source_digest = str(manifest["source_digest"])
        prefix = f"reports/patrol/{patrol_batch_id}/{source_digest}"
        object_keys: dict[str, str] = {}
        sha256_by_artifact: dict[str, str] = {}
        for name, artifact in artifacts.items():
            object_key = f"{prefix}/{artifact.filename}"
            uploaded = self._object_storage.upload_file(
                object_key,
                artifact.content,
                content_type=artifact.content_type,
            )
            if not uploaded:
                raise RuntimeError(f"failed to upload report artifact: {name}")
            object_keys[name] = object_key
            sha256_by_artifact[name] = artifact.sha256
        attached = self._manifest_service.attach_artifacts(
            str(manifest["id"]),
            object_keys=object_keys,
            sha256_by_artifact=sha256_by_artifact,
        )
        return {
            "success": True,
            "reused": False,
            "patrol_batch_id": patrol_batch_id,
            "report_manifest_id": str(manifest["id"]),
            "source_digest": source_digest,
            "artifact_object_keys": object_keys,
            "artifact_sha256": sha256_by_artifact,
            "artifact_bytes": {name: len(artifact.content) for name, artifact in artifacts.items()},
            "manifest": attached,
        }


_ARTIFACT_CONTENT_TYPES = {
    "markdown": "text/markdown; charset=utf-8",
    "docx": ("application/vnd.openxmlformats-officedocument.wordprocessingml.document"),
    "excel": ("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
    "json": "application/json; charset=utf-8",
    "evidence": "application/zip",
}


def load_attached_report_artifact(
    manifest: Mapping[str, Any],
    name: str,
    *,
    object_storage: Any | None = None,
) -> ReportArtifact:
    """Read one immutable artifact and verify it against the manifest."""

    artifact_name = str(name)
    if artifact_name not in _ARTIFACT_CONTENT_TYPES:
        raise ValueError(f"unsupported report artifact: {artifact_name}")
    object_keys = _as_dict(manifest.get("artifact_object_keys_json"))
    artifact_digests = _as_dict(manifest.get("artifact_sha256_json"))
    object_key = str(object_keys.get(artifact_name) or "").strip()
    expected_digest = str(artifact_digests.get(artifact_name) or "").strip().lower()
    if not object_key or not expected_digest:
        raise ValueError(f"report artifact is not attached: {artifact_name}")
    if object_storage is None:
        from src.infrastructure.storage.minio_client import (
            get_minio_storage,
        )

        object_storage = get_minio_storage()
    content = object_storage.download_file(object_key)
    if content is None:
        raise RuntimeError(f"attached report artifact is unavailable: {artifact_name}")
    actual_digest = hashlib.sha256(content).hexdigest()
    if actual_digest != expected_digest:
        raise RuntimeError(f"attached report artifact digest mismatch: {artifact_name}")
    return ReportArtifact(
        name=artifact_name,
        filename=object_key.rsplit("/", 1)[-1],
        content_type=_ARTIFACT_CONTENT_TYPES[artifact_name],
        content=content,
    )


def verify_attached_report_artifacts(
    manifest: Mapping[str, Any],
    *,
    object_storage: Any | None = None,
) -> dict[str, Any]:
    """Verify every formal artifact against the immutable manifest digest."""

    object_keys = _as_dict(manifest.get("artifact_object_keys_json"))
    expected_digests = {
        str(name): str(value).strip().lower()
        for name, value in _as_dict(manifest.get("artifact_sha256_json")).items()
    }
    expected_names = set(_ARTIFACT_CONTENT_TYPES)
    metadata_complete = (
        set(object_keys) == expected_names
        and set(expected_digests) == expected_names
        and all(
            len(value) == 64 and all(char in "0123456789abcdef" for char in value)
            for value in expected_digests.values()
        )
    )
    if object_storage is None:
        from src.infrastructure.storage.minio_client import (
            get_minio_storage,
        )

        object_storage = get_minio_storage()

    results: dict[str, dict[str, Any]] = {}
    for name in sorted(expected_names):
        object_key = str(object_keys.get(name) or "").strip()
        expected_digest = str(expected_digests.get(name) or "")
        if not object_key or not expected_digest:
            results[name] = {
                "status": "metadata_missing",
                "size": None,
            }
            continue
        try:
            if callable(getattr(object_storage, "sha256_object", None)):
                observed = object_storage.sha256_object(object_key)
                actual_digest = str((observed or {}).get("sha256") or "").lower()
                size = (observed or {}).get("size")
            else:
                content = object_storage.download_file(object_key)
                actual_digest = hashlib.sha256(content).hexdigest() if content is not None else ""
                size = len(content) if content is not None else None
        except Exception as exc:  # storage adapters must fail closed
            results[name] = {
                "status": "storage_error",
                "size": None,
                "error": type(exc).__name__,
            }
            continue
        if not actual_digest:
            status = "object_missing"
        elif actual_digest != expected_digest:
            status = "digest_mismatch"
        else:
            status = "verified"
        results[name] = {
            "status": status,
            "size": size,
            "expected_sha256": expected_digest,
            "actual_sha256": actual_digest or None,
        }

    passed = metadata_complete and all(
        item.get("status") == "verified" for item in results.values()
    )
    result = {
        "schema_version": "report-artifact-integrity/v1",
        "passed": passed,
        "metadata_complete": metadata_complete,
        "artifacts": results,
        "failed_artifact_names": [
            name for name, item in results.items() if item.get("status") != "verified"
        ],
    }
    result["digest"] = hashlib.sha256(
        json.dumps(
            result,
            ensure_ascii=False,
            sort_keys=True,
            separators=(",", ":"),
            default=str,
        ).encode("utf-8")
    ).hexdigest()
    return result


def _digest(value: Any) -> str:
    return hashlib.sha256(
        json.dumps(
            value,
            ensure_ascii=False,
            sort_keys=True,
            separators=(",", ":"),
            default=str,
        ).encode("utf-8")
    ).hexdigest()


def _iso(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.isoformat()
    return str(value)


def _jsonable(value: Any) -> Any:
    if isinstance(value, datetime):
        return value.isoformat()
    if isinstance(value, Mapping):
        return {str(key): _jsonable(item) for key, item in value.items()}
    if isinstance(value, (list, tuple, set)):
        return [_jsonable(item) for item in value]
    return value


def _as_dict(value: Any) -> dict[str, Any]:
    return dict(value) if isinstance(value, Mapping) else {}


def _as_rows(value: Any) -> list[dict[str, Any]]:
    return [dict(item) for item in value or [] if isinstance(item, Mapping)]


def _display_value(row: Mapping[str, Any]) -> Any:
    if row.get("value_number") is not None:
        return row.get("value_number")
    if row.get("value_text") not in (None, ""):
        return row.get("value_text")
    return row.get("value")


def _value_with_unit(value: Any, unit: Any) -> str:
    if value in (None, ""):
        return "-"
    suffix = str(unit or "").strip()
    return f"{value} {suffix}".strip()


def _reference_source_label(value: Any) -> str:
    source = str(value or "").strip()
    if not source:
        return "-"
    if source.startswith("context:"):
        return f"观测上下文字段 {source.split(':', 1)[1]}"
    return _REFERENCE_SOURCE_LABELS.get(source, source)


def _event_criterion_label(row: Mapping[str, Any]) -> str:
    kind = str(row.get("condition_kind") or "").strip()
    condition = _as_dict(row.get("condition"))
    reference = row.get("reference_value")
    unit = row.get("unit")
    operator = {
        "threshold": ">=",
        "threshold_recovery": "<",
        "threshold_lte": "<=",
        "threshold_low_recovery": ">=",
        "threshold_field_gte": ">=",
        "threshold_field_high_recovery": "<=",
        "threshold_field_lte": "<=",
        "threshold_field_low_recovery": ">=",
    }.get(kind)
    boundary = row.get("recovery_boundary")
    if "recovery" in kind and boundary is not None:
        reference = boundary
    if operator and reference is not None:
        source = _reference_source_label(row.get("reference_source"))
        return (
            f"{row.get('field') or '-'} {operator} {_value_with_unit(reference, unit)}（{source}）"
        )
    if kind == "transition":
        return f"{condition.get('from_group') or '*'} -> {condition.get('to_group') or '*'}"
    return kind or "-"


def _risk_criterion_label(row: Mapping[str, Any]) -> str:
    reference = (
        row.get("reference_value_number")
        if row.get("reference_value_number") is not None
        else row.get("reference_value_text")
        if row.get("reference_value_text") not in (None, "")
        else row.get("reference_value")
    )
    if reference in (None, ""):
        return "-"
    return (
        f"{_value_with_unit(reference, row.get('unit'))}"
        f"（{_reference_source_label(row.get('reference_source'))}）"
    )


def _risk_actual_value(row: Mapping[str, Any]) -> Any:
    return (
        row.get("actual_value_number")
        if row.get("actual_value_number") is not None
        else row.get("actual_value_text")
        if row.get("actual_value_text") not in (None, "")
        else row.get("actual_value")
    )


def _risk_actual_value_label(row: Mapping[str, Any]) -> str:
    return _value_with_unit(_risk_actual_value(row), row.get("unit"))


def _risk_remediation_owner_label(row: Mapping[str, Any]) -> str:
    status = _localized_label(
        row.get("remediation_status"),
        _REMEDIATION_STATUS_LABELS,
    )
    return f"{status} / {row.get('owner') or '未分配'}"


def _risk_closure_status(row: Mapping[str, Any]) -> str:
    explicit = str(row.get("closure_status") or "").strip()
    if explicit:
        return explicit
    if str(row.get("status") or "") == "suppressed":
        return "suppressed"
    if row.get("closed_at"):
        return "closed"
    if str(row.get("status") or "") != "resolved":
        return "open"
    remediation = str(row.get("remediation_status") or "not_started")
    if remediation == "accepted":
        return "ready_to_close"
    if remediation == "completed":
        return "pending_acceptance"
    if remediation in {"planned", "in_progress"}:
        return "pending_remediation"
    return "technical_recovered"


def _device_label(
    device_id: str,
    target_by_device: Mapping[str, Mapping[str, Any]],
) -> str:
    target = target_by_device.get(device_id, {})
    return str(target.get("device_name") or device_id or "-")


def _check_item_label(check_item: Mapping[str, Any]) -> str:
    check_item_id = str(check_item.get("check_item_id") or "").strip()
    display_name = _check_item_display_name(check_item)
    if display_name and check_item_id and display_name != check_item_id:
        return f"{display_name}（{check_item_id}）"
    return display_name or check_item_id or "-"


def _check_item_display_name(check_item: Mapping[str, Any]) -> str:
    name = str(check_item.get("name") or "").strip()
    check_item_id = str(check_item.get("check_item_id") or "").strip()
    return (
        check_item_display_name(
            check_item_id,
            name if name and name != check_item_id else None,
        )
        or "-"
    )


def _localized_label(
    value: Any,
    labels: Mapping[str, str],
) -> str:
    normalized = str(value or "").strip().lower()
    return labels.get(normalized, str(value or "-"))


def _reason_code_label(
    value: Any,
    *,
    include_code: bool = True,
) -> str:
    code = str(value or "").strip()
    if not code:
        return "-"
    label = _BLOCKING_REASON_LABELS.get(code)
    if label is None and code.startswith("execution_"):
        status = code.removeprefix("execution_")
        localized = _EXECUTION_STATUS_LABELS.get(status)
        if localized:
            label = f"执行终态：{localized}"
    if label is None:
        return code
    return f"{label}（{code}）" if include_code else label


def _reason_codes_label(
    value: Any,
    *,
    include_codes: bool = True,
) -> str:
    values = _as_rows_or_values(value)
    if not values:
        return "-"
    return "；".join(
        _reason_code_label(
            item,
            include_code=include_codes,
        )
        for item in values
    )


def _count_map_label(value: Any, labels: Mapping[str, str]) -> str:
    counts = _as_dict(value)
    if not counts:
        return "-"
    return "；".join(
        f"{labels.get(str(key), str(key))} {int(count or 0)}"
        for key, count in sorted(counts.items())
    )


def _assessment_basis_label(item: Mapping[str, Any]) -> str:
    digest = str(item.get("assessment_basis_digest") or "").strip()
    if not digest:
        return "历史数据未冻结判断依据"
    source_label = _count_map_label(
        item.get("assessment_basis_source_counts"),
        _ASSESSMENT_SOURCE_LABELS,
    )
    status_label = _count_map_label(
        item.get("assessment_basis_status_counts"),
        _ASSESSMENT_BASIS_STATUS_LABELS,
    )
    entry_count = int(item.get("assessment_basis_entry_count") or 0)
    if entry_count <= 0:
        return "本项未形成实体级规则判定"
    return f"来源：{source_label}；状态：{status_label}"


def _collection_route_count_label(item: Mapping[str, Any]) -> str:
    return (
        f"{int(item.get('executable_command_count') or 0)} / "
        f"{int(item.get('declared_command_count') or 0)}"
    )


def _unavailable_command_label(item: Mapping[str, Any]) -> str:
    routes = _as_rows(item.get("unavailable_commands"))
    if not routes:
        return "-"
    return "；".join(
        (
            f"{route.get('command') or route.get('canonical_command') or '-'}："
            f"{_reason_code_label(route.get('reason_code'), include_code=False)}"
        )
        for route in routes
    )


def _compact_identifier(value: Any) -> str:
    text = str(value or "-")
    if len(text) <= 30:
        return text
    return f"{text[:18]}…{text[-8:]}"


def _compact_digest(value: Any) -> str:
    text = str(value or "-")
    if len(text) <= 24:
        return text
    return f"{text[:12]}…{text[-8:]}"


def _evidence_quality_label(item: Mapping[str, Any]) -> str:
    return " / ".join(
        (
            _localized_label(
                item.get("source_identity_status"),
                _SOURCE_IDENTITY_LABELS,
            ),
            _localized_label(
                item.get("output_quality_status"),
                _OUTPUT_QUALITY_LABELS,
            ),
            _localized_label(
                item.get("freshness_status"),
                _FRESHNESS_LABELS,
            ),
        )
    )


def _coverage_label(item: Mapping[str, Any]) -> str:
    status = _localized_label(
        item.get("coverage_status"),
        _COVERAGE_STATUS_LABELS,
    )
    basis = _localized_label(
        item.get("coverage_basis"),
        _COVERAGE_BASIS_LABELS,
    )
    details: list[str] = [basis]
    expected = item.get("expected_entity_count")
    observed = item.get("observed_entity_count")
    if expected is not None:
        details.append(f"已观测 {observed or 0} / 期望 {expected}")
    missing = _as_rows_or_values(item.get("missing_entity_sample"))
    if missing:
        details.append(
            "缺失样例 " + "、".join(_compact_entity_sample(value) for value in missing[:5])
        )
    return f"{status}（{'；'.join(details)}）"


def _as_rows_or_values(value: Any) -> list[Any]:
    if isinstance(value, (list, tuple, set)):
        return list(value)
    return [] if value in (None, "") else [value]


def _compact_entity_sample(value: Any) -> str:
    if isinstance(value, (dict, list, tuple, set)):
        text = json.dumps(
            _jsonable(value),
            ensure_ascii=False,
            sort_keys=True,
            separators=(",", ":"),
        )
    else:
        text = str(value)
    return text if len(text) <= 80 else f"{text[:77]}..."


def validate_report_manifest(manifest: Mapping[str, Any]) -> None:
    """Reject mutable, incomplete or digest-inconsistent report inputs."""

    source_facts = _as_dict(manifest.get("source_facts_json"))
    if not source_facts:
        raise ValueError("report manifest has no frozen source facts")
    source_schema = str(source_facts.get("schema_version") or "")
    if source_schema not in {
        "patrol-report-source/v2",
        "patrol-report-source/v3",
        "patrol-report-source/v4",
        "patrol-report-source/v5",
    }:
        raise ValueError("unsupported patrol report source schema")
    expected_digest = str(manifest.get("source_digest") or "")
    if len(expected_digest) != 64 or _digest(source_facts) != expected_digest:
        raise ValueError("report manifest source digest is invalid")
    if str(source_facts.get("patrol_batch_id") or "") != str(manifest.get("patrol_batch_id") or ""):
        raise ValueError("report manifest batch identity mismatch")
    if source_schema in {
        "patrol-report-source/v3",
        "patrol-report-source/v4",
        "patrol-report-source/v5",
    }:
        expected_manifest_schema = {
            "patrol-report-source/v3": "patrol-report-manifest/v3",
            "patrol-report-source/v4": "patrol-report-manifest/v4",
            "patrol-report-source/v5": "patrol-report-manifest/v5",
        }[source_schema]
        if str(manifest.get("schema_version") or "") not in {
            "",
            expected_manifest_schema,
        }:
            raise ValueError("report manifest schema identity mismatch")
        manifest_purpose = normalize_report_purpose(manifest.get("report_purpose"))
        source_purpose = normalize_report_purpose(source_facts.get("report_purpose"))
        if manifest_purpose != source_purpose:
            raise ValueError("report manifest purpose mismatch")
        expected_raw_digest = str(manifest.get("raw_evidence_set_digest") or "")
        if (
            len(expected_raw_digest) != 64
            or str(source_facts.get("raw_evidence_set_digest") or "") != expected_raw_digest
            or raw_evidence_set_digest(_as_rows(source_facts.get("evidence_index")))
            != expected_raw_digest
        ):
            raise ValueError("report manifest raw evidence set digest is invalid")
    if source_schema in {
        "patrol-report-source/v4",
        "patrol-report-source/v5",
    }:
        frozen_permission = _as_dict(source_facts.get("conclusion_permission"))
        if not report_conclusion_permission_is_valid(frozen_permission):
            raise ValueError("report conclusion permission digest is invalid")
        expected_permission = build_report_conclusion_permission_from_source_facts(source_facts)
        if frozen_permission != expected_permission:
            raise ValueError("report conclusion permission does not match the frozen source facts")
    if source_schema == "patrol-report-source/v5":
        validate_report_claim_ledger(source_facts)
        ledger = _as_dict(source_facts.get("claim_ledger"))
        if str(manifest.get("claim_ledger_digest") or "") != str(
            ledger.get("claim_ledger_digest") or ""
        ):
            raise ValueError("report manifest claim ledger digest is invalid")
        if int(manifest.get("claim_count") or 0) != len(_as_rows(ledger.get("claims"))):
            raise ValueError("report manifest claim count is invalid")
    projection_snapshot = source_facts.get("report_projection_snapshot")
    if projection_snapshot:
        validate_report_projection_snapshot(_as_dict(projection_snapshot))
        if str(source_facts.get("report_projection_digest") or "") != str(
            _as_dict(projection_snapshot).get("digest") or ""
        ):
            raise ValueError("report projection snapshot digest is invalid")


def build_check_item_report_view(
    manifest: Mapping[str, Any],
) -> dict[str, Any]:
    """Build a human-oriented view without consulting mutable runtime tables."""

    validate_report_manifest(manifest)
    facts = _as_dict(manifest.get("source_facts_json"))
    source_schema = str(facts.get("schema_version") or "")
    report_purpose = (
        normalize_report_purpose(facts.get("report_purpose"))
        if source_schema
        in {
            "patrol-report-source/v3",
            "patrol-report-source/v4",
            "patrol-report-source/v5",
        }
        else SHADOW_VALIDATION
    )
    purpose_label = _REPORT_PURPOSE_LABELS[report_purpose]
    purpose_disclosure = _REPORT_PURPOSE_DISCLOSURES[report_purpose]
    readiness = _as_dict(facts.get("readiness"))
    targets = _as_rows(facts.get("target_facts"))
    executions = _as_rows(facts.get("executions"))
    assessments = _as_rows(facts.get("assessments"))
    coverage_decisions = _as_rows(facts.get("coverage_decisions"))
    claim_ledger = _as_dict(facts.get("claim_ledger"))
    report_claims = _as_rows(claim_ledger.get("claims"))
    claim_by_code = {
        str(item.get("claim_code") or ""): item
        for item in report_claims
        if str(item.get("claim_code") or "")
    }
    claim_by_fact_ref = {
        str(reference): item
        for item in report_claims
        for reference in item.get("fact_refs") or []
        if str(reference)
    }

    def claim_id_for(*references: str) -> str | None:
        for reference in references:
            claim = claim_by_fact_ref.get(str(reference))
            if claim:
                return str(claim.get("claim_id") or "") or None
        return None

    check_items = {
        str(key): _as_dict(value) for key, value in _as_dict(facts.get("check_items")).items()
    }
    target_by_device = {str(item.get("device_id") or ""): item for item in targets}
    execution_by_id = {str(item.get("id") or ""): item for item in executions}
    assessment_by_execution = {str(item.get("execution_id") or ""): item for item in assessments}
    coverage_by_id = {
        str(item.get("id") or ""): item for item in coverage_decisions if str(item.get("id") or "")
    }
    resilience_rows = _as_rows(facts.get("collection_resilience"))
    resilience_by_pair = {
        (
            str(item.get("device_id") or ""),
            str(item.get("check_item_version_id") or ""),
        ): item
        for item in resilience_rows
    }

    matrix: list[dict[str, Any]] = []
    for execution in executions:
        execution_id = str(execution.get("id") or "")
        device_id = str(execution.get("device_id") or "")
        version_id = str(execution.get("check_item_version_id") or "")
        check_item = check_items.get(version_id, {})
        assessment = assessment_by_execution.get(execution_id, {})
        coverage_id = str(
            assessment.get("coverage_decision_id") or execution.get("coverage_decision_id") or ""
        )
        coverage = coverage_by_id.get(coverage_id, {})
        resilience = resilience_by_pair.get(
            (device_id, version_id),
            {},
        )
        assessment_basis = _as_dict(assessment.get("assessment_basis_json"))
        assessment_basis_summary = _as_dict(assessment_basis.get("summary"))
        matrix.append(
            {
                "execution_id": execution_id,
                "assessment_id": assessment.get("id"),
                "claim_id": claim_id_for(f"CheckAssessment:{assessment.get('id')}"),
                "device_id": device_id,
                "device_name": _device_label(
                    device_id,
                    target_by_device,
                ),
                "management_ip": target_by_device.get(
                    device_id,
                    {},
                ).get("management_ip"),
                "check_item_version_id": version_id,
                "check_item_id": check_item.get("check_item_id"),
                "check_item_name": check_item.get("name"),
                "check_item_display_name": _check_item_display_name(check_item),
                "check_item_label": _check_item_label(check_item),
                "category": check_item.get("category"),
                "result_mode": check_item.get("result_mode"),
                "assessment_mode": check_item.get("assessment_mode"),
                "report_projection": check_item.get("report_projection"),
                "execution_status": execution.get("execution_status"),
                "assessment_status": assessment.get(
                    "assessment_status",
                    "unable_to_determine",
                ),
                "completeness_status": assessment.get(
                    "completeness_status",
                    "insufficient",
                ),
                "comparison_status": assessment.get(
                    "comparison_status",
                    "baseline_missing",
                ),
                "blocking_reason_codes": list(
                    assessment.get("blocking_reason_codes_json")
                    or execution.get("reason_codes_json")
                    or []
                ),
                "expected_entity_count": assessment.get(
                    "expected_entity_count",
                    coverage.get("expected_entity_count"),
                ),
                "observed_entity_count": assessment.get(
                    "observed_entity_count",
                    execution.get("observed_entity_count", 0),
                ),
                "evaluated_entity_count": assessment.get(
                    "evaluated_entity_count",
                    0,
                ),
                "threshold_unknown_count": assessment.get(
                    "threshold_unknown_count",
                    0,
                ),
                "observation_count": execution.get(
                    "observation_count",
                    0,
                ),
                "change_count": execution.get("change_count", 0),
                "event_count": execution.get("event_count", 0),
                "current_finding_count": execution.get(
                    "current_finding_count",
                    0,
                ),
                "evidence_qualification_id": assessment.get(
                    "evidence_qualification_id",
                    execution.get("evidence_qualification_id"),
                ),
                "coverage_decision_id": assessment.get(
                    "coverage_decision_id",
                    execution.get("coverage_decision_id"),
                ),
                "source_identity_status": assessment.get(
                    "source_identity_status",
                    execution.get("source_identity_status"),
                ),
                "output_quality_status": assessment.get(
                    "output_quality_status",
                    execution.get("output_quality_status"),
                ),
                "freshness_status": assessment.get(
                    "freshness_status",
                    execution.get("freshness_status"),
                ),
                "empty_result_classification": assessment.get(
                    "empty_result_classification",
                    execution.get("empty_result_classification"),
                ),
                "coverage_status": assessment.get(
                    "coverage_status",
                    execution.get("coverage_status"),
                ),
                "coverage_basis": assessment.get(
                    "coverage_basis",
                    execution.get("coverage_basis"),
                ),
                "coverage_denominator_ref": dict(coverage.get("denominator_ref_json") or {}),
                "coverage_reason_codes": list(coverage.get("reason_codes_json") or []),
                "missing_entity_sample": list(coverage.get("missing_entity_sample_json") or []),
                "unexpected_entity_sample": list(
                    coverage.get("unexpected_entity_sample_json") or []
                ),
                "risk_verification_status": assessment.get(
                    "risk_verification_status",
                    execution.get("risk_verification_status"),
                ),
                "evidence_digest": assessment.get("evidence_digest"),
                "runtime_bundle_digest": execution.get("runtime_bundle_digest"),
                "evidence_refs": list(execution.get("evidence_refs_json") or []),
                "assessment_policy_snapshot_digest": assessment.get(
                    "assessment_policy_snapshot_digest",
                    execution.get("assessment_policy_snapshot_digest"),
                ),
                "assessment_basis_digest": assessment.get("assessment_basis_digest"),
                "assessment_basis_entry_count": int(
                    assessment_basis_summary.get("entry_count") or 0
                ),
                "assessment_basis_source_counts": _as_dict(
                    assessment_basis_summary.get("source_counts")
                ),
                "assessment_basis_status_counts": _as_dict(
                    assessment_basis_summary.get("status_counts")
                ),
                "fallback_resilience": resilience.get(
                    "fallback_resilience",
                    "unknown",
                ),
                "declared_command_count": int(resilience.get("declared_command_count") or 0),
                "executable_command_count": int(resilience.get("executable_command_count") or 0),
                "selected_commands": list(resilience.get("selected_commands") or []),
                "unavailable_commands": list(resilience.get("unavailable_commands") or []),
            }
        )
    matrix.sort(
        key=lambda item: (
            str(item.get("device_name") or ""),
            str(item.get("check_item_id") or ""),
        )
    )

    health_assessments = [row for row in matrix if row.get("assessment_mode") == "health"]
    status_counts = Counter(
        str(row.get("assessment_status") or "unable_to_determine") for row in matrix
    )
    completeness_counts = Counter(
        str(row.get("completeness_status") or "insufficient") for row in matrix
    )
    comparison_counts = Counter(
        str(row.get("comparison_status") or "baseline_missing") for row in matrix
    )
    blocker_counts = Counter(
        str(code) for row in matrix for code in row.get("blocking_reason_codes") or []
    )
    resilience_counts = Counter(str(row.get("fallback_resilience") or "unknown") for row in matrix)
    conclusion_permission = (
        _as_dict(facts.get("conclusion_permission"))
        if source_schema in {"patrol-report-source/v4", "patrol-report-source/v5"}
        else {}
    )

    manifest_status = str(manifest.get("status") or "")
    # Supersession only changes which immutable report is the latest version.
    # It must not downgrade a previously final report's frozen conclusion.
    formal = manifest_status in {"final", "superseded"} and bool(readiness.get("ready_for_final"))
    if not formal:
        conclusion_code = "not_final"
        conclusion = "正式巡检结论尚未形成"
        conclusion_basis = (
            "目标设备、检查项执行或评估尚未全部完成，当前结果仅供过程排错，不能作为正式健康结论。"
        )
        trust = "过程数据"
    elif status_counts["at_risk"] and status_counts["unable_to_determine"]:
        conclusion_code = "at_risk_with_unknown"
        conclusion = "本次巡检发现当前风险，且部分检查项无法判断"
        conclusion_basis = (
            f"已确认 {status_counts['at_risk']} 项当前风险；另有 "
            f"{status_counts['unable_to_determine']} 项检查因证据不足"
            "无法判断，已确认风险和未知范围均需跟进。"
        )
        trust = "受限结论"
    elif status_counts["at_risk"]:
        conclusion_code = "at_risk"
        conclusion = "本次巡检发现当前风险"
        conclusion_basis = (
            f"已完成 {len(matrix)} 项检查，其中 {status_counts['at_risk']} 项形成可追溯风险结论。"
        )
        trust = "正式结论"
    elif status_counts["unable_to_determine"]:
        conclusion_code = "partially_unknown"
        conclusion = "部分检查项无法判断"
        conclusion_basis = (
            f"已有 {status_counts['unable_to_determine']} 项检查"
            "因采集、解析、规则、阈值或基线证据不足而无法判断；"
            "其余检查结果不得代表未覆盖范围。"
        )
        trust = "受限结论"
    elif not health_assessments or not any(
        row.get("assessment_status") == "healthy" for row in health_assessments
    ):
        conclusion_code = "informational_only"
        conclusion = "本次仅形成信息性结果，未判定网络健康"
        conclusion_basis = (
            "本次范围未包含可形成健康判定的检查项，因此不能使用“未发现异常”作为结论。"
        )
        trust = "信息性结果"
    elif all(
        row.get("assessment_status") in {"healthy", "not_applicable"} for row in health_assessments
    ):
        conclusion_code = "healthy"
        conclusion = "本次覆盖范围内未发现异常"
        conclusion_basis = (
            f"{len(health_assessments)} 项健康类检查均已形成完整评估，"
            "且未产生当前风险；该结论仅适用于本次目标范围和已发布规则版本。"
        )
        trust = "正式结论"
    else:
        conclusion_code = "mixed"
        conclusion = "本次检查结果需工程复核"
        conclusion_basis = "存在尚未纳入明确结论类别的评估状态。"
        trust = "受限结论"

    assessment_pattern_code = conclusion_code
    current_state_authoritative = bool(
        formal and report_purpose in CURRENT_STATE_AUTHORITATIVE_PURPOSES
    )
    authority_context = _as_dict(facts.get("report_authority_context"))
    authority_context_integrity_valid = report_authority_context_is_valid(
        authority_context,
        expected_purpose=report_purpose,
    )
    authority_digest_consistent = all(
        str(authority_context.get(key) or "") == str(expected or "")
        for key, expected in {
            "target_manifest_digest": (manifest.get("target_manifest_digest")),
            "runtime_manifest_digest": (manifest.get("runtime_manifest_digest")),
            "inventory_snapshot_digest": (facts.get("inventory_snapshot_digest")),
        }.items()
    )
    authority_context_valid = bool(
        authority_context_integrity_valid
        and authority_digest_consistent
        and authority_context.get("current_state_eligible")
    )
    current_state_authoritative = bool(current_state_authoritative and authority_context_valid)
    if source_schema in {
        "patrol-report-source/v4",
        "patrol-report-source/v5",
    }:
        current_state_authoritative = bool(
            formal and conclusion_permission.get("current_conclusion_ready")
        )
    if formal and not current_state_authoritative:
        if report_purpose == PRODUCTION_CURRENT:
            conclusion_code = "current_authority_blocked"
            conclusion = "当前网络状态无法形成正式结论"
            if source_schema in {
                "patrol-report-source/v4",
                "patrol-report-source/v5",
            }:
                reason_codes = [
                    str(item)
                    for item in conclusion_permission.get(
                        "blocker_reason_codes",
                        [],
                    )
                    if str(item)
                ]
            else:
                reason_codes = [
                    str(item)
                    for item in authority_context.get(
                        "reason_codes",
                        [],
                    )
                    if str(item)
                ]
                if not authority_context:
                    reason_codes.append("report_authority_context_missing")
                elif not authority_context_integrity_valid:
                    reason_codes.append("report_authority_context_invalid")
                elif not authority_digest_consistent:
                    reason_codes.append("report_authority_digest_mismatch")
            if not reason_codes:
                reason_codes.append("report_authority_context_not_eligible")
            conclusion_basis = (
                "报告缺少完整的当前状态来源证明，不能使用“网络健康”"
                "或“未发现异常”作为结论。阻断原因："
                f"{', '.join(reason_codes)}。"
            )
            trust = "受限结论"
        else:
            purpose_conclusions = {
                OFFLINE_HISTORICAL: (
                    "historical_result",
                    "历史采集时点巡检结果已形成",
                ),
                SHADOW_VALIDATION: (
                    "shadow_validation_result",
                    "影子验证结果已形成",
                ),
                PARSER_VALIDATION: (
                    "parser_validation_result",
                    "Parser 验证结果已形成",
                ),
                RULE_ASSET_REPLAY: (
                    "rule_asset_replay_result",
                    "规则资产回放结果已形成",
                ),
            }
            conclusion_code, conclusion = purpose_conclusions[report_purpose]
            conclusion_basis = (
                f"{purpose_disclosure}"
                f"本次事实中风险判定 {status_counts['at_risk']} 项，"
                f"无法判断 {status_counts['unable_to_determine']} 项，"
                f"不适用 {status_counts['not_applicable']} 项。"
            )
            trust = "非生产结论"
    elif source_schema in {
        "patrol-report-source/v4",
        "patrol-report-source/v5",
    }:
        (
            conclusion_code,
            conclusion,
            conclusion_basis,
            trust,
        ) = _conclusion_from_permission(conclusion_permission)

    conclusion_dimensions = _build_conclusion_dimensions(
        formal=formal,
        report_purpose=report_purpose,
        current_state_authoritative=current_state_authoritative,
        matrix_count=len(matrix),
        health_assessments=health_assessments,
        status_counts=status_counts,
        completeness_counts=completeness_counts,
        comparison_counts=comparison_counts,
    )
    unresolved_risk_count = int(
        _as_dict(conclusion_permission.get("counts")).get("unresolved_in_scope_risk_case_count")
        or 0
    )
    if unresolved_risk_count:
        conclusion_dimensions["risk_conclusion"] = {
            "code": "unresolved_risk",
            "label": "存在未关闭风险",
            "basis": (f"目标范围内有 {unresolved_risk_count} 项 RiskCase 尚未取得合格恢复证据。"),
        }
    assessment_pattern_code = conclusion_code

    findings = _as_rows(facts.get("findings"))
    for finding in findings:
        finding["claim_id"] = claim_id_for(f"CurrentStateFinding:{finding.get('id')}")
    changes = _as_rows(facts.get("changes"))
    aggregate_events = _as_rows(facts.get("events"))
    event_occurrences = _as_rows(facts.get("event_occurrences"))
    if event_occurrences:
        occurrence_aggregate_ids = {
            str(item.get("aggregate_event_id") or "")
            for item in event_occurrences
            if item.get("aggregate_event_id")
        }
        events = [
            item
            for item in aggregate_events
            if str(item.get("id") or "") not in occurrence_aggregate_ids
        ] + event_occurrences
    else:
        events = aggregate_events
    risks = _as_rows(facts.get("risks"))
    occurrences = _as_rows(facts.get("risk_occurrences"))
    for occurrence in occurrences:
        occurrence["claim_id"] = claim_id_for(f"RiskOccurrence:{occurrence.get('id')}")
    remediations = _as_rows(facts.get("risk_remediations"))
    observations = _as_rows(facts.get("observations"))
    evidence_index = _as_rows(facts.get("evidence_index"))
    enriched_observations: list[dict[str, Any]] = []
    check_item_labels_by_snapshot: dict[str, set[str]] = {}
    for observation in observations:
        row = dict(observation)
        device_id = str(row.get("device_id") or "")
        linked_executions = [
            execution_by_id[execution_id]
            for execution_id in (str(item) for item in row.get("execution_ids") or [])
            if execution_id in execution_by_id
        ]
        linked_check_items = [
            check_items.get(
                str(execution.get("check_item_version_id") or ""),
                {},
            )
            for execution in linked_executions
        ]
        check_item_ids = sorted(
            {
                str(item.get("check_item_id") or "")
                for item in linked_check_items
                if item.get("check_item_id")
            }
        )
        check_item_labels = sorted({_check_item_label(item) for item in linked_check_items if item})
        row["device_name"] = _device_label(device_id, target_by_device)
        row["management_ip"] = target_by_device.get(device_id, {}).get("management_ip")
        row["check_item_ids"] = check_item_ids
        row["check_item_labels"] = check_item_labels
        snapshot_id = str(row.get("snapshot_id") or "")
        if snapshot_id:
            check_item_labels_by_snapshot.setdefault(snapshot_id, set()).update(check_item_labels)
        enriched_observations.append(row)
    enriched_observations.sort(
        key=lambda item: (
            str(item.get("device_name") or ""),
            str((item.get("check_item_labels") or [""])[0]),
            str(item.get("observation_type") or ""),
            str(item.get("entity_key") or ""),
            str(item.get("field") or ""),
        )
    )
    enriched_evidence_index: list[dict[str, Any]] = []
    for evidence in evidence_index:
        row = dict(evidence)
        row["check_item_labels"] = sorted(
            check_item_labels_by_snapshot.get(
                str(row.get("snapshot_id") or ""),
                set(),
            )
        )
        enriched_evidence_index.append(row)
    risk_by_id = {str(item.get("id") or ""): item for item in risks}
    finding_by_id = {str(item.get("id") or ""): item for item in findings if item.get("id")}
    current_execution_ids = {str(item.get("id") or "") for item in executions if item.get("id")}
    current_finding_ids = {str(item.get("id") or "") for item in findings if item.get("id")}
    current_event_ids = {
        str(identifier)
        for item in events
        for identifier in (
            item.get("id"),
            item.get("aggregate_event_id"),
        )
        if identifier
    }
    check_item_name_by_id = {
        str(item.get("check_item_id") or ""): _check_item_label(item)
        for item in check_items.values()
    }
    enriched_risks: list[dict[str, Any]] = []
    for risk in risks:
        row = dict(risk)
        row["claim_id"] = claim_id_for(f"RiskCase:{row.get('id')}")
        row["closure_status"] = _risk_closure_status(row)
        device_id = str(row.get("device_id") or "")
        row["device_name"] = _device_label(
            device_id,
            target_by_device,
        )
        row["check_item_label"] = check_item_name_by_id.get(
            str(row.get("check_item_id") or ""),
            str(row.get("check_item_id") or "-"),
        )
        finding_refs = {str(item) for item in row.get("source_finding_ids_json") or [] if item}
        event_refs = {str(item) for item in row.get("source_event_ids_json") or [] if item}
        source_findings = [
            finding_by_id[finding_id]
            for finding_id in sorted(finding_refs)
            if finding_id in finding_by_id
        ]
        if source_findings:
            latest_finding = max(
                source_findings,
                key=lambda item: (
                    str(item.get("observed_at") or ""),
                    str(item.get("id") or ""),
                ),
            )
            row["actual_value"] = latest_finding.get("actual_value")
            row["actual_value_text"] = latest_finding.get("actual_value_text")
            row["actual_value_number"] = latest_finding.get("actual_value_number")
            row["reference_value"] = latest_finding.get("reference_value")
            row["reference_value_text"] = latest_finding.get("reference_value_text")
            row["reference_value_number"] = latest_finding.get("reference_value_number")
            row["reference_source"] = latest_finding.get("reference_source")
            row["unit"] = latest_finding.get("unit")
            row["event_rule_version_id"] = latest_finding.get("event_rule_version_id")
            row["rule_digest"] = latest_finding.get("rule_digest")
        verification_execution_id = str(row.get("last_verification_execution_id") or "")
        row["verified_by_current_batch"] = bool(
            verification_execution_id in current_execution_ids
            and (
                finding_refs.intersection(current_finding_ids)
                or event_refs.intersection(current_event_ids)
            )
        )
        enriched_risks.append(row)
    enriched_risks.sort(
        key=lambda item: (
            -_SEVERITY_ORDER.get(
                str(item.get("severity") or "unknown").lower(),
                0,
            ),
            str(item.get("device_name") or ""),
            str(item.get("entity_key") or ""),
        )
    )
    current_risks = [
        row
        for row in enriched_risks
        if str(row.get("verification_status") or "") == "confirmed_current"
        and bool(row.get("verified_by_current_batch"))
        and str(row.get("status") or "") not in {"resolved", "suppressed"}
    ]
    risk_issue_summary = build_risk_issue_projection(
        risks=current_risks,
        associations=_as_rows(facts.get("risk_episode_associations")),
        event_occurrences=event_occurrences,
    )
    for issue in _as_rows(risk_issue_summary.get("issues")):
        issue["claim_id"] = claim_id_for(f"RiskIssue:{issue.get('issue_id')}") or issue.get(
            "claim_id"
        )
    risk_issue_summary["issues"] = _as_rows(risk_issue_summary.get("issues"))
    recovery_pending_risks = [
        row
        for row in enriched_risks
        if str(row.get("verification_status") or "") == "recovery_pending"
        and str(row.get("status") or "") != "resolved"
    ]
    resolved_risks = [
        row
        for row in enriched_risks
        if str(row.get("verification_status") or "") == "resolved"
        or str(row.get("status") or "") == "resolved"
    ]
    formally_closed_risks = [
        row for row in resolved_risks if str(row.get("closure_status") or "") == "closed"
    ]
    technically_recovered_risks = [
        row for row in resolved_risks if str(row.get("closure_status") or "") != "closed"
    ]
    current_risk_ids = {str(row.get("id") or "") for row in current_risks}
    recovery_pending_ids = {str(row.get("id") or "") for row in recovery_pending_risks}
    resolved_risk_ids = {str(row.get("id") or "") for row in resolved_risks}
    historical_not_reverified_risks = [
        row
        for row in enriched_risks
        if str(row.get("id") or "") not in current_risk_ids
        and str(row.get("id") or "") not in recovery_pending_ids
        and str(row.get("id") or "") not in resolved_risk_ids
    ]
    risk_verification_counts = {
        "confirmed_current": len(current_risks),
        "recovery_pending": len(recovery_pending_risks),
        "not_reverified": len(historical_not_reverified_risks),
        "resolved": len(resolved_risks),
        "technical_recovered": len(technically_recovered_risks),
        "closed": len(formally_closed_risks),
    }
    enriched_remediations: list[dict[str, Any]] = []
    for remediation in remediations:
        row = dict(remediation)
        row["claim_id"] = claim_id_for(f"RiskRemediation:{row.get('id')}")
        risk = risk_by_id.get(str(row.get("risk_case_id") or ""), {})
        row["device_id"] = risk.get("device_id")
        row["device_name"] = _device_label(
            str(risk.get("device_id") or ""),
            target_by_device,
        )
        row["check_item_id"] = risk.get("check_item_id")
        row["risk_type"] = risk.get("risk_type")
        row["severity"] = risk.get("severity")
        enriched_remediations.append(row)
    enriched_changes: list[dict[str, Any]] = []
    for change in changes:
        row = dict(change)
        row["claim_id"] = claim_id_for(f"NetworkChange:{row.get('id')}")
        row["device_name"] = _device_label(
            str(row.get("device_id") or ""),
            target_by_device,
        )
        row["check_item_label"] = check_item_name_by_id.get(
            str(row.get("check_item_id") or ""),
            str(row.get("check_item_id") or "-"),
        )
        enriched_changes.append(row)
    enriched_changes.sort(
        key=lambda item: (
            str(item.get("observed_at") or ""),
            str(item.get("device_name") or ""),
            str(item.get("entity_key") or ""),
            str(item.get("field") or ""),
        ),
        reverse=True,
    )
    enriched_events: list[dict[str, Any]] = []
    for event in events:
        row = dict(event)
        row["claim_id"] = claim_id_for(
            f"NetworkEventOccurrence:{row.get('id')}",
            f"NetworkEvent:{row.get('id')}",
        )
        row["device_name"] = _device_label(
            str(row.get("device_id") or ""),
            target_by_device,
        )
        row["check_item_label"] = check_item_name_by_id.get(
            str(row.get("check_item_id") or ""),
            str(row.get("check_item_id") or "-"),
        )
        row["criterion_label"] = _event_criterion_label(row)
        enriched_events.append(row)
    enriched_events.sort(
        key=lambda item: (
            str(item.get("occurred_at") or ""),
            _SEVERITY_ORDER.get(
                str(item.get("severity") or "unknown").lower(),
                0,
            ),
            str(item.get("device_name") or ""),
            str(item.get("entity_key") or ""),
        ),
        reverse=True,
    )

    comparison_contexts = _as_rows(facts.get("comparison_contexts"))
    comparison_by_id = {str(item.get("id") or ""): item for item in comparison_contexts}
    for row in matrix:
        execution = execution_by_id.get(
            str(row.get("execution_id") or ""),
            {},
        )
        context = comparison_by_id.get(
            str(execution.get("comparison_context_id") or ""),
            {},
        )
        row["baseline_run_id"] = context.get("baseline_run_id")
        row["baseline_snapshot_ids"] = list(context.get("baseline_snapshot_ids") or [])

    source_digest = str(manifest.get("source_digest") or "")
    report_view = {
        "schema_version": (
            "check-item-patrol-report/v2"
            if source_schema == "patrol-report-source/v5"
            else "check-item-patrol-report/v1"
        ),
        "report_manifest_id": str(manifest.get("id") or ""),
        "report_manifest_version": int(manifest.get("version") or 0),
        "report_status": str(manifest.get("status") or ""),
        "patrol_batch_id": str(manifest.get("patrol_batch_id") or ""),
        "report_purpose": report_purpose,
        "report_purpose_label": purpose_label,
        "report_purpose_disclosure": purpose_disclosure,
        "report_authority_context": authority_context,
        "conclusion_permission": conclusion_permission,
        "current_state_authoritative": (current_state_authoritative),
        "generated_at": _iso(manifest.get("generated_at")),
        "report_cutoff_at": _iso(manifest.get("report_cutoff_at") or facts.get("report_cutoff_at")),
        "fact_ingestion_high_watermark": _iso(
            manifest.get("fact_ingestion_high_watermark")
            or facts.get("fact_ingestion_high_watermark")
        ),
        "late_fact_count": int(facts.get("late_fact_count") or 0),
        "generator_version": str(manifest.get("generator_version") or ""),
        "source_digest": source_digest,
        "raw_evidence_set_digest": str(
            manifest.get("raw_evidence_set_digest") or facts.get("raw_evidence_set_digest") or ""
        ),
        "qualification_fact_digest": str(manifest.get("qualification_fact_digest") or ""),
        "coverage_fact_digest": str(manifest.get("coverage_fact_digest") or ""),
        "target_manifest_digest": str(manifest.get("target_manifest_digest") or ""),
        "runtime_manifest_digest": str(manifest.get("runtime_manifest_digest") or ""),
        "report_projection_digest": str(facts.get("report_projection_digest") or ""),
        "report_projection_snapshot": _as_dict(facts.get("report_projection_snapshot")),
        "claim_basis_digest": str(facts.get("claim_basis_digest") or ""),
        "claim_ledger_digest": str(facts.get("claim_ledger_digest") or ""),
        "claim_ledger": claim_ledger,
        "claims": report_claims,
        "conclusion": {
            "claim_id": str(claim_by_code.get("conclusion.overall", {}).get("claim_id") or "")
            or None,
            "code": conclusion_code,
            "text": conclusion,
            "basis": conclusion_basis,
            "trust": trust,
            "formal": formal,
            "current_state_authoritative": (current_state_authoritative),
            "assessment_pattern_code": (assessment_pattern_code),
        },
        "conclusion_dimensions": conclusion_dimensions,
        "coverage": {
            "claim_id": str(claim_by_code.get("scope.summary", {}).get("claim_id") or "") or None,
            "target_device_count": int(readiness.get("target_count") or len(targets)),
            "target_record_count": int(readiness.get("target_record_count") or len(targets)),
            "expected_check_count": int(readiness.get("expected_execution_count") or len(matrix)),
            "executed_check_count": len(executions),
            "assessed_check_count": len(assessments),
            "complete_check_count": completeness_counts["complete"],
            "partial_check_count": completeness_counts["partial"],
            "insufficient_check_count": completeness_counts["insufficient"],
        },
        "assessment_counts": dict(sorted(status_counts.items())),
        "completeness_counts": dict(sorted(completeness_counts.items())),
        "comparison_counts": dict(sorted(comparison_counts.items())),
        "blocking_reason_counts": dict(sorted(blocker_counts.items())),
        "collection_resilience": {
            "counts": dict(sorted(resilience_counts.items())),
            "reduced_count": int(resilience_counts.get("reduced") or 0),
            "none_count": int(resilience_counts.get("none") or 0),
        },
        "readiness": readiness,
        "targets": targets,
        "check_items": check_items,
        "matrix": matrix,
        # Keep the legacy key truthful: it now means risks confirmed by this
        # report's frozen current-batch evidence, not every open historical case.
        "risks": current_risks,
        "current_risks": current_risks,
        "historical_not_reverified_risks": (historical_not_reverified_risks),
        "recovery_pending_risks": recovery_pending_risks,
        "resolved_risks": resolved_risks,
        "technically_recovered_risks": technically_recovered_risks,
        "formally_closed_risks": formally_closed_risks,
        "all_risk_cases": enriched_risks,
        "risk_verification_counts": risk_verification_counts,
        "risk_issue_summary": risk_issue_summary,
        "risk_issues": _as_rows(risk_issue_summary.get("issues")),
        "risk_episode_associations": _as_rows(facts.get("risk_episode_associations")),
        "risk_occurrences": occurrences,
        "remediations": enriched_remediations,
        "findings": findings,
        "changes": enriched_changes,
        "events": enriched_events,
        "observations": enriched_observations,
        "evidence_index": enriched_evidence_index,
    }
    return _jsonable(report_view)


def _markdown_text(value: Any) -> str:
    text = str(value if value not in (None, "") else "-")
    return text.replace("|", "\\|").replace("\r", " ").replace("\n", " ")


def _markdown_table(
    headers: Sequence[str],
    rows: Iterable[Sequence[Any]],
) -> list[str]:
    output = [
        "| " + " | ".join(headers) + " |",
        "| " + " | ".join("---" for _ in headers) + " |",
    ]
    materialized = list(rows)
    if not materialized:
        materialized = [["-" for _ in headers]]
    output.extend(
        "| " + " | ".join(_markdown_text(value) for value in row) + " |" for row in materialized
    )
    return output


def _comparison_change_disclosure(
    report: Mapping[str, Any],
    *,
    change_count: int,
) -> str:
    readiness = _as_dict(_as_dict(report.get("conclusion_dimensions")).get("comparison_readiness"))
    code = str(readiness.get("code") or "no_baseline")
    basis = str(readiness.get("basis") or "").strip()
    if code == "comparable":
        if change_count:
            return f"全部目标具备合格比较基线，本次形成 {change_count} 条标准化变化。"
        return "全部目标具备合格比较基线；本次可比较范围内未形成标准化变化。"
    if code == "partial":
        return (
            f"本次形成 {change_count} 条标准化变化，仅代表具备合格基线的范围；"
            f"其余范围不能判断是否发生变化。{basis}"
        )
    return (
        f"本次形成 {change_count} 条标准化变化；当前比较条件未成立，不能据此判断未发生变化。{basis}"
    )


def render_check_item_markdown(report: Mapping[str, Any]) -> str:
    """Render the leadership report with explicit confidence limitations."""

    coverage = _as_dict(report.get("coverage"))
    conclusion = _as_dict(report.get("conclusion"))
    conclusion_dimensions = _as_dict(report.get("conclusion_dimensions"))
    matrix = _as_rows(report.get("matrix"))
    risks = _as_rows(report.get("risks"))
    risk_issue_summary = _as_dict(report.get("risk_issue_summary"))
    risk_issues = _as_rows(report.get("risk_issues"))
    historical_risks = _as_rows(report.get("historical_not_reverified_risks"))
    recovery_pending_risks = _as_rows(report.get("recovery_pending_risks"))
    resolved_risks = _as_rows(report.get("resolved_risks"))
    remediations = _as_rows(report.get("remediations"))
    changes = _as_rows(report.get("changes"))
    events = _as_rows(report.get("events"))
    blockers = _as_dict(report.get("blocking_reason_counts"))
    purpose_label = str(report.get("report_purpose_label") or "未声明用途")
    purpose_disclosure = str(report.get("report_purpose_disclosure") or "")
    risk_heading = (
        "当前风险" if report.get("current_state_authoritative") else "风险结果（非当前状态）"
    )
    basic_operational_tables = _basic_operational_report_tables(report)
    device_inventory_rows = _report_device_inventory_rows(
        report,
        basic_operational_tables,
    )
    basic_sections: list[str] = []
    for index, operational_table in enumerate(basic_operational_tables, start=1):
        basic_sections.extend(
            [
                "",
                f"##### 3.1.1.{index} {operational_table['title']}",
                "",
                *_markdown_table(
                    tuple(operational_table["headers"]),
                    operational_table["rows"],
                ),
            ]
        )
    lines = [
        f"# 网络设备巡检报告（{purpose_label}）",
        "",
        f"- 巡检批次：`{report.get('patrol_batch_id')}`",
        (
            f"- 报告清单：`{report.get('report_manifest_id')}`"
            f"（版本 {report.get('report_manifest_version')}）"
        ),
        f"- 报告用途：**{purpose_label}**",
        f"- 生成时间：{report.get('generated_at') or '-'}",
        f"- 数据指纹：`{report.get('source_digest')}`",
        (f"- 原始证据集合指纹：`{report.get('raw_evidence_set_digest') or '-'}`"),
        f"- 用途声明：{purpose_disclosure}",
        "",
        "## 第一章 概述",
        "",
        "### 1.1 巡检概述",
        "",
        f"**{conclusion.get('text')}。**",
        f"声明编号：`{conclusion.get('claim_id') or '-'}`",
        "",
        str(conclusion.get("basis") or ""),
        "",
        f"结论可信度：**{conclusion.get('trust')}**。",
        "",
        *_markdown_table(
            ("结论维度", "判定", "依据"),
            [
                (
                    label,
                    _as_dict(conclusion_dimensions.get(key)).get("label"),
                    _as_dict(conclusion_dimensions.get(key)).get("basis"),
                )
                for key, label in (
                    ("risk_conclusion", "风险结论"),
                    ("assessment_completeness", "评估完整性"),
                    ("comparison_readiness", "周期可比性"),
                    ("authority", "结论权威性"),
                )
            ],
        ),
        "",
        "### 1.2 巡检说明",
        "",
        "本次巡检仅采集和核验设备运行信息，不对设备配置作变更。健康结论以检查项"
        "绑定的标准语义 B、设备或模块上报阈值、有效基线和已发布规则为依据；"
        "数据不足时明确标记为无法判断。",
        "",
        "### 1.3 巡检范围",
        "",
        *_markdown_table(
            ("目标设备", "计划检查", "已执行", "已评估", "完整", "部分", "证据不足"),
            [
                (
                    coverage.get("target_device_count"),
                    coverage.get("expected_check_count"),
                    coverage.get("executed_check_count"),
                    coverage.get("assessed_check_count"),
                    coverage.get("complete_check_count"),
                    coverage.get("partial_check_count"),
                    coverage.get("insufficient_check_count"),
                )
            ],
        ),
        "",
        "#### 1.3.1 设备清单",
        "",
        *_markdown_table(
            ("序号", "设备名称", "设备型号", "管理 IP"),
            device_inventory_rows,
        ),
        "",
        "## 第二章 巡检内容",
        "",
        "### 2.1 网络设备巡检项",
        "",
        "本次巡检按已启用检查项和实际采集命令执行，涵盖设备基础信息、硬件环境、"
        "性能、接口状态、配置与日志等内容。具体健康判断以检查项绑定的标准语义 B 和"
        "规则资产为准。",
        "",
        "#### 2.1.1 采集路径韧性",
        "",
        (
            "可用路径少于检查项声明路径时，巡检仍可执行，但设备兼容性或命令故障时的"
            "自动回退能力已经降低；该状态不等同于检查项健康。"
        ),
        "",
        *_markdown_table(
            ("设备", "检查项", "路径状态", "可用 / 声明", "不可用路径及原因"),
            [
                (
                    item.get("device_name"),
                    item.get("check_item_label"),
                    _localized_label(
                        item.get("fallback_resilience"),
                        _FALLBACK_RESILIENCE_LABELS,
                    ),
                    _collection_route_count_label(item),
                    _unavailable_command_label(item),
                )
                for item in matrix
            ],
        ),
        "",
        "### 2.2 巡检设备配置备份",
        "",
        "由于设备数量较多，配置备份文件较大，不适宜粘贴到 Word 文档内，因此将以"
        "电子版的形式单独提交，特此说明。",
        "",
        "## 第三章 巡检结果",
        "",
        "### 3.1 网络设备",
        "",
        f"**{conclusion.get('text')}。** {conclusion.get('basis') or ''}",
        "",
        "#### 3.1.1 设备基本运行状况",
        "",
        "以下表格展示本次已采集并映射为标准字段的设备基础运行数据；表中不展示采集时间。"
        if basic_sections
        else "本批次未形成可用于报告的设备版本、CPU、内存、电源、风扇或温度标准观测数据。",
        *basic_sections,
        "",
        f"#### 3.1.2 {risk_heading}",
        "",
        (
            "巡检主问题数按经批准关联后的 RiskEpisode 主问题统计；技术风险、关联症状、"
            "影响实体和事件发生次数分别列示。未审批关联不会改变问题数。"
        ),
        "",
        *_markdown_table(
            (
                "巡检主问题",
                "风险Episode",
                "技术风险",
                "关联症状",
                "影响设备",
                "影响实体",
                "事件发生",
                "待复核关联",
            ),
            [
                (
                    risk_issue_summary.get("problem_count", len(risks)),
                    risk_issue_summary.get("risk_episode_count", len(risks)),
                    risk_issue_summary.get("technical_risk_count", len(risks)),
                    risk_issue_summary.get("associated_symptom_count", 0),
                    risk_issue_summary.get("affected_device_count", 0),
                    risk_issue_summary.get("affected_entity_count", 0),
                    risk_issue_summary.get("event_occurrence_count", len(events)),
                    risk_issue_summary.get("pending_review_count", 0),
                )
            ],
        ),
        "",
        "### 巡检主问题清单",
        "",
        *_markdown_table(
            (
                "等级",
                "主问题",
                "检查项",
                "技术风险",
                "影响设备",
                "影响实体",
                "事件发生",
                "关联状态",
                "声明编号",
            ),
            [
                (
                    _SEVERITY_LABELS.get(
                        str(item.get("severity") or "unknown").lower(),
                        item.get("severity"),
                    ),
                    item.get("risk_type"),
                    "、".join(item.get("check_item_ids") or []),
                    item.get("technical_risk_count", item.get("symptom_count")),
                    item.get("affected_device_count"),
                    item.get("affected_entity_count"),
                    item.get("event_occurrence_count"),
                    "已批准关联" if item.get("correlation_applied") else "独立问题",
                    item.get("claim_id"),
                )
                for item in risk_issues
            ],
        ),
        "",
        "### 技术风险明细",
        "",
        *_markdown_table(
            (
                "等级",
                "设备",
                "检查项",
                "风险",
                "对象",
                "当前值",
                "判据",
                "状态",
                "整改状态",
                "声明编号",
            ),
            [
                (
                    _SEVERITY_LABELS.get(
                        str(item.get("severity") or "unknown").lower(),
                        item.get("severity"),
                    ),
                    item.get("device_name"),
                    item.get("check_item_label"),
                    item.get("risk_type"),
                    item.get("entity_key"),
                    _risk_actual_value_label(item),
                    _risk_criterion_label(item),
                    _localized_label(
                        item.get("status"),
                        _RISK_STATUS_LABELS,
                    ),
                    _localized_label(
                        item.get("remediation_status"),
                        _REMEDIATION_STATUS_LABELS,
                    ),
                    item.get("claim_id"),
                )
                for item in risks
            ],
        ),
        "",
        "#### 3.1.3 历史风险待复核",
        "",
        (
            "以下风险来自历史批次，但本批次没有形成足以确认其"
            "当前状态的合格证据；不得计入“本次确认风险”，也不得"
            "直接视为已恢复。"
        ),
        "",
        *_markdown_table(
            ("等级", "设备", "检查项", "风险", "对象", "历史状态", "声明编号"),
            [
                (
                    _SEVERITY_LABELS.get(
                        str(item.get("severity") or "unknown").lower(),
                        item.get("severity"),
                    ),
                    item.get("device_name"),
                    item.get("check_item_label"),
                    item.get("risk_type"),
                    item.get("entity_key"),
                    _localized_label(
                        item.get("status"),
                        _RISK_STATUS_LABELS,
                    ),
                    item.get("claim_id"),
                )
                for item in historical_risks
            ],
        ),
        "",
        "#### 3.1.4 恢复待确认",
        "",
        ("以下风险在本批次未再次出现，但仍需恢复规则或后续合格巡检证据完成闭环。"),
        "",
        *_markdown_table(
            ("等级", "设备", "检查项", "风险", "对象", "当前状态", "声明编号"),
            [
                (
                    _SEVERITY_LABELS.get(
                        str(item.get("severity") or "unknown").lower(),
                        item.get("severity"),
                    ),
                    item.get("device_name"),
                    item.get("check_item_label"),
                    item.get("risk_type"),
                    item.get("entity_key"),
                    _localized_label(
                        item.get("status"),
                        _RISK_STATUS_LABELS,
                    ),
                    item.get("claim_id"),
                )
                for item in recovery_pending_risks
            ],
        ),
        "",
        "##### 3.1.4.1 技术恢复与正式关闭",
        "",
        (
            "技术恢复、整改验收和正式关闭分别列示；技术恢复不代表整改已验收，"
            "也不代表风险已经正式关闭。"
        ),
        "",
        *_markdown_table(
            (
                "设备",
                "检查项",
                "风险",
                "对象",
                "技术状态",
                "整改状态",
                "关闭状态",
                "关闭人",
                "声明编号",
            ),
            [
                (
                    item.get("device_name"),
                    item.get("check_item_label"),
                    item.get("risk_type"),
                    item.get("entity_key"),
                    _localized_label(
                        item.get("status"),
                        _RISK_STATUS_LABELS,
                    ),
                    _localized_label(
                        item.get("remediation_status"),
                        _REMEDIATION_STATUS_LABELS,
                    ),
                    _localized_label(
                        item.get("closure_status"),
                        _RISK_CLOSURE_LABELS,
                    ),
                    item.get("closed_by"),
                    item.get("claim_id"),
                )
                for item in resolved_risks
            ],
        ),
        "",
        "#### 3.1.5 检查项结论",
        "",
        *_markdown_table(
            (
                "设备",
                "检查项",
                "执行",
                "结论",
                "完整性",
                "证据资质",
                "覆盖",
                "判断依据",
                "基线",
                "阻断原因",
                "声明编号",
            ),
            [
                (
                    item.get("device_name"),
                    item.get("check_item_label"),
                    _localized_label(
                        item.get("execution_status"),
                        _EXECUTION_STATUS_LABELS,
                    ),
                    _STATUS_LABELS.get(
                        str(item.get("assessment_status")),
                        item.get("assessment_status"),
                    ),
                    _COMPLETENESS_LABELS.get(
                        str(item.get("completeness_status")),
                        item.get("completeness_status"),
                    ),
                    _evidence_quality_label(item),
                    _coverage_label(item),
                    _assessment_basis_label(item),
                    _COMPARISON_LABELS.get(
                        str(item.get("comparison_status")),
                        item.get("comparison_status"),
                    ),
                    _reason_codes_label(
                        item.get("blocking_reason_codes"),
                    ),
                    item.get("claim_id"),
                )
                for item in matrix
            ],
        ),
        "",
        "#### 3.1.6 变化与事件",
        "",
        f"- 标准化变化：{len(changes)} 条。",
        f"- 网络事件：{len(events)} 条。",
        f"- {_comparison_change_disclosure(report, change_count=len(changes))}",
        "- 未产生事件只表示没有命中已发布事件规则，不能单独证明设备健康。",
        "",
        *_markdown_table(
            ("设备", "检查项", "对象", "字段", "原值", "新值", "判定", "声明编号"),
            [
                (
                    item.get("device_name"),
                    item.get("check_item_label"),
                    item.get("entity_key"),
                    item.get("field"),
                    item.get("old"),
                    item.get("new"),
                    item.get("decision_reason"),
                    item.get("claim_id"),
                )
                for item in changes[:20]
            ],
        ),
        "",
        *_markdown_table(
            (
                "等级",
                "设备",
                "检查项",
                "事件",
                "对象",
                "当前值",
                "触发判据",
                "规则版本",
                "状态",
                "发生时间",
                "声明编号",
            ),
            [
                (
                    _SEVERITY_LABELS.get(
                        str(item.get("severity") or "unknown").lower(),
                        item.get("severity"),
                    ),
                    item.get("device_name"),
                    item.get("check_item_label"),
                    item.get("event_type"),
                    item.get("entity_key"),
                    _value_with_unit(
                        item.get("current_value"),
                        item.get("unit"),
                    ),
                    item.get("criterion_label"),
                    (f"{item.get('event_rule_id') or '-'}@{item.get('event_rule_version') or '-'}"),
                    _localized_label(
                        item.get("status"),
                        _RISK_STATUS_LABELS,
                    ),
                    item.get("occurred_at"),
                    item.get("claim_id"),
                )
                for item in events[:20]
            ],
        ),
        "",
        ("- 巡检报告最多展示 20 条变化和 20 条事件；完整明细见技术附表。"),
        "",
        "#### 3.1.7 已登记整改情况",
        "",
        *_markdown_table(
            ("设备", "风险", "整改状态", "责任人", "期限", "整改动作", "声明编号"),
            [
                (
                    item.get("device_name"),
                    item.get("risk_type"),
                    _localized_label(
                        item.get("status"),
                        _REMEDIATION_STATUS_LABELS,
                    ),
                    item.get("owner"),
                    item.get("due_at"),
                    item.get("action_text"),
                    item.get("claim_id"),
                )
                for item in remediations
            ],
        ),
        "",
        "## 第四章 巡检总结",
        "",
        "### 4.1 数据可信度与结论边界",
        "",
    ]
    if blockers:
        lines.extend(
            f"- {_reason_code_label(code)}：{count} 项。"
            for code, count in sorted(blockers.items())
        )
    else:
        lines.append(
            "- 未发现影响本次当前结论的数据阻断项。"
            if report.get("current_state_authoritative")
            else ("- 未发现影响本报告用途内结论的数据阻断项；这不提升其当前网络状态结论权限。")
        )
    lines.extend(
        [
            "",
            "## 第五章 后续工作计划",
            "",
            "### 5.1 后续工作计划",
            "",
        ]
    )
    if remediations:
        lines.extend(
            f"- {item.get('device_name') or '-'}：{item.get('action_text') or '按整改台账持续跟踪'}"
            f"（责任：{item.get('owner') or '-'}；期限：{item.get('due_at') or '-'}）。"
            for item in remediations
        )
    else:
        lines.append("本批次暂无已登记整改事项；后续应按第三章风险和数据质量结论持续复核。")

    lines.extend(
        [
            "",
            "## 附录 A：证据与复核说明",
            "",
            f"- 目标清单指纹：`{report.get('target_manifest_digest')}`",
            f"- 运行资产清单指纹：`{report.get('runtime_manifest_digest')}`",
            f"- 报告事实指纹：`{report.get('source_digest')}`",
            f"- 事实声明账本指纹：`{report.get('claim_ledger_digest') or '-'}`",
            f"- 报告投影合同指纹：`{report.get('report_projection_digest') or '-'}`",
            f"- 事实声明数量：{len(_as_rows(report.get('claims')))} 条。",
            (
                "- 结论权限证明指纹："
                f"`{_as_dict(report.get('report_authority_context')).get('digest') or '-'}`"
            ),
            f"- 证据资质事实指纹：`{report.get('qualification_fact_digest') or '-'}`",
            f"- 覆盖决策事实指纹：`{report.get('coverage_fact_digest') or '-'}`",
            f"- 事实截止时间：{report.get('report_cutoff_at') or '-'}",
            (f"- 事实摄入高水位：{report.get('fact_ingestion_high_watermark') or '-'}"),
            f"- 截止后事实修订：{report.get('late_fact_count') or 0} 条。",
            f"- 证据索引：{len(_as_rows(report.get('evidence_index')))} 条命令快照。",
            "- 每项结论均可由设备、检查项、命令快照、Parser/Contract 版本和规则资产摘要复核。",
            "- 原始 CLI 不嵌入本报告；原始输出通过独立、受权限控制的交付物提供。",
            "",
        ]
    )
    return "\n".join(lines)


def _normalize_zip_archive(content: bytes) -> bytes:
    source_stream = io.BytesIO(content)
    output_stream = io.BytesIO()
    with (
        zipfile.ZipFile(source_stream, "r") as source,
        zipfile.ZipFile(
            output_stream,
            "w",
            compression=zipfile.ZIP_DEFLATED,
            compresslevel=9,
        ) as target,
    ):
        for source_info in sorted(
            source.infolist(),
            key=lambda item: item.filename,
        ):
            entry_content = source.read(source_info.filename)
            if source_info.filename == "docProps/core.xml":
                created_match = re.search(
                    rb"(<dcterms:created\b[^>]*>)([^<]*)(</dcterms:created>)",
                    entry_content,
                )
                if created_match is not None:
                    created_value = created_match.group(2)
                    entry_content = re.sub(
                        rb"(<dcterms:modified\b[^>]*>)[^<]*(</dcterms:modified>)",
                        lambda match: match.group(1) + created_value + match.group(2),
                        entry_content,
                    )
            target_info = zipfile.ZipInfo(
                filename=source_info.filename,
                date_time=_ARCHIVE_TIMESTAMP,
            )
            target_info.compress_type = zipfile.ZIP_DEFLATED
            target_info.create_system = source_info.create_system
            target_info.external_attr = source_info.external_attr
            target_info.internal_attr = source_info.internal_attr
            target_info.comment = source_info.comment
            target.writestr(
                target_info,
                entry_content,
            )
    return output_stream.getvalue()


def _report_datetime(value: Any) -> datetime:
    try:
        parsed = datetime.fromisoformat(str(value or "").replace("Z", "+00:00"))
    except ValueError:
        parsed = datetime(2000, 1, 1, tzinfo=timezone.utc)
    if parsed.tzinfo is not None:
        return parsed.astimezone(timezone.utc).replace(tzinfo=None)
    return parsed


_BASIC_OPERATIONAL_REPORT_TABLES = (
    ("device.version", "设备启动信息与版本状态", False),
    ("environment.power", "设备电源状态", True),
    ("environment.fan", "设备风扇状态", True),
    ("environment.temperature", "设备环境温度状态", True),
    ("system.cpu_usage", "设备 CPU 利用率状态", False),
    ("system.memory_usage", "设备内存利用率状态", False),
)
_REPORT_PROJECTION_BASE_COLUMNS = {
    "device_name",
    "management_ip",
    "entity_type",
    "entity_key",
}
_REPORT_PROJECTION_AUXILIARY_COLUMNS = {
    "execution_status",
    "assessment_status",
    "completeness_status",
    "comparison_status",
    "threshold_unknown_count",
    "risk_verification_status",
    "blocking_reason_codes",
    "execution_id",
    "snapshot_id",
    "previous_snapshot_id",
    "current_snapshot_id",
    "parser_template_id",
    "runtime_bundle_digest",
    "observed_at",
}


def _basic_operational_report_tables(
    report: Mapping[str, Any],
) -> list[dict[str, Any]]:
    """Return customer-readable B tables for the main Word report.

    The report projection snapshot is already frozen with the report facts.  This
    only removes IDs, timestamps and runtime bookkeeping columns; it does not
    reinterpret vendor A fields or apply any new health logic.
    """

    projected = {table.check_item_id: table for table in project_report_tables(report)}
    result: list[dict[str, Any]] = []
    for check_item_id, title, include_component in _BASIC_OPERATIONAL_REPORT_TABLES:
        source = projected.get(check_item_id)
        if source is None or not source.rows:
            continue
        positions = {column_id: index for index, column_id in enumerate(source.column_ids)}
        canonical_ids = [
            column_id
            for column_id in source.column_ids
            if column_id not in _REPORT_PROJECTION_BASE_COLUMNS
            and column_id not in _REPORT_PROJECTION_AUXILIARY_COLUMNS
        ]
        selected_ids = ["device_name"]
        if include_component:
            selected_ids.append("entity_key")
        selected_ids.extend(canonical_ids)
        selected = [column_id for column_id in selected_ids if column_id in positions]
        if not selected:
            continue

        headers: list[str] = []
        for column_id in selected:
            if column_id == "device_name":
                headers.append("设备")
            elif column_id == "entity_key":
                headers.append("部件")
            else:
                headers.append(f"{source.headers[positions[column_id]]}（{column_id}）")
        rows = [
            [
                row[positions[column_id]] if row[positions[column_id]] not in (None, "") else "-"
                for column_id in selected
            ]
            for row in source.rows
        ]
        result.append(
            {
                "check_item_id": check_item_id,
                "title": title,
                "headers": headers,
                "rows": rows,
            }
        )
    return result


def _report_device_inventory_rows(
    report: Mapping[str, Any],
    basic_tables: Sequence[Mapping[str, Any]],
) -> list[tuple[str, str, str, str]]:
    """Build the overview device list without depending on mutable inventory."""

    models: dict[str, str] = {}
    for table in basic_tables:
        if table.get("check_item_id") != "device.version":
            continue
        headers = list(table.get("headers") or [])
        model_index = next(
            (index for index, value in enumerate(headers) if "hardware_model" in str(value)),
            None,
        )
        for row in list(table.get("rows") or []):
            if row and model_index is not None and len(row) > model_index:
                models[str(row[0])] = str(row[model_index] or "-")

    devices: dict[str, tuple[str, str, str]] = {}
    for target in _as_rows(report.get("targets")):
        name = str(target.get("device_name") or target.get("device_id") or "-")
        devices[name] = (
            name,
            models.get(name, "-"),
            str(target.get("management_ip") or "-"),
        )
    for name, model in models.items():
        devices.setdefault(name, (name, model, "-"))
    return [
        (str(index), name, model, management_ip)
        for index, (name, model, management_ip) in enumerate(
            sorted(devices.values(), key=lambda item: item[0]),
            start=1,
        )
    ]


def render_check_item_docx(report: Mapping[str, Any]) -> bytes:
    """Render a formal customer patrol report from the frozen report view."""

    from docx import Document
    from docx.enum.table import (
        WD_CELL_VERTICAL_ALIGNMENT,
        WD_TABLE_ALIGNMENT,
    )
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor

    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(2.4)
    section.bottom_margin = Cm(2.2)
    section.left_margin = Cm(2.6)
    section.right_margin = Cm(2.6)
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    generated_at = _report_datetime(report.get("generated_at"))
    purpose_label = str(report.get("report_purpose_label") or "未声明用途")
    purpose_disclosure = str(report.get("report_purpose_disclosure") or "")
    risk_heading = (
        "当前风险" if report.get("current_state_authoritative") else "风险结果（非当前状态）"
    )
    doc.core_properties.created = generated_at
    doc.core_properties.modified = generated_at
    doc.core_properties.author = "NetOps Agent"
    doc.core_properties.title = f"网络设备巡检报告（{purpose_label}）"

    styles = doc.styles
    styles["Normal"].font.name = "宋体"
    styles["Normal"].font.size = Pt(10.5)
    styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_run = footer.add_run("第 ")
    footer_run.font.name = "宋体"
    footer_run.font.size = Pt(9)

    def field(paragraph: Any, instruction: str) -> None:
        begin = OxmlElement("w:fldChar")
        begin.set(qn("w:fldCharType"), "begin")
        text = OxmlElement("w:instrText")
        text.set(qn("xml:space"), "preserve")
        text.text = instruction
        separate = OxmlElement("w:fldChar")
        separate.set(qn("w:fldCharType"), "separate")
        display = OxmlElement("w:t")
        display.text = "1"
        end = OxmlElement("w:fldChar")
        end.set(qn("w:fldCharType"), "end")
        run = paragraph.add_run()
        run._r.extend((begin, text, separate, display, end))

    field(footer, " PAGE ")
    footer.add_run(" 页 / 共 ")
    field(footer, " NUMPAGES ")
    footer.add_run(" 页")

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(16)
    title_run = title.add_run(f"网络设备巡检报告（{purpose_label}）")
    title_run.bold = True
    title_run.font.name = "方正小标宋简体"
    title_run.font.size = Pt(22)
    title_run._element.rPr.rFonts.set(
        qn("w:eastAsia"),
        "方正小标宋简体",
    )

    def heading(text: str, *, level: int = 1) -> None:
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.space_before = Pt({1: 16, 2: 12, 3: 8}.get(level, 5))
        paragraph.paragraph_format.space_after = Pt(4)
        paragraph.paragraph_format.keep_with_next = True
        run = paragraph.add_run(text)
        run.bold = True
        run.font.name = "黑体"
        run.font.size = Pt({1: 16, 2: 15, 3: 14}.get(level, 12))
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")

    def paragraph(text: str, *, bold: bool = False) -> None:
        item = doc.add_paragraph()
        item.paragraph_format.first_line_indent = Cm(0.74)
        item.paragraph_format.line_spacing = 1.5
        run = item.add_run(text)
        run.bold = bold

    def table(
        headers: Sequence[str],
        rows: Iterable[Sequence[Any]],
        *,
        empty_text: str = "本节无可列示事项。",
    ) -> None:
        materialized = list(rows)
        if not materialized:
            paragraph(empty_text)
            return
        grid = doc.add_table(rows=1, cols=len(headers))
        grid.style = "Table Grid"
        grid.alignment = WD_TABLE_ALIGNMENT.CENTER
        header_properties = grid.rows[0]._tr.get_or_add_trPr()
        repeat_header = OxmlElement("w:tblHeader")
        repeat_header.set(qn("w:val"), "true")
        header_properties.append(repeat_header)
        for index, header in enumerate(headers):
            cell = grid.rows[0].cells[index]
            cell.text = str(header)
            shading = OxmlElement("w:shd")
            shading.set(qn("w:fill"), "365F4B")
            cell._tc.get_or_add_tcPr().append(shading)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for run in cell.paragraphs[0].runs:
                run.bold = True
                run.font.name = "黑体"
                run.font.size = Pt(9.5)
                run.font.color.rgb = RGBColor(255, 255, 255)
                run._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
        for row in materialized:
            table_row = grid.add_row()
            no_split = OxmlElement("w:cantSplit")
            table_row._tr.get_or_add_trPr().append(no_split)
            cells = table_row.cells
            for index, value in enumerate(row):
                cells[index].text = str(value if value not in (None, "") else "-")
                cells[index].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                for cell_paragraph in cells[index].paragraphs:
                    cell_paragraph.paragraph_format.space_after = Pt(0)
                    cell_paragraph.paragraph_format.line_spacing = 1.1
                    for run in cell_paragraph.runs:
                        run.font.name = "宋体"
                        run.font.size = Pt(9)
                        run._element.rPr.rFonts.set(
                            qn("w:eastAsia"),
                            "宋体",
                        )
        doc.add_paragraph()

    paragraph(f"报告用途：{purpose_label}。{purpose_disclosure}")

    table(
        ("批次编号", "目标设备", "计划检查", "事实状态"),
        [
            (
                _compact_identifier(report.get("patrol_batch_id")),
                _as_dict(report.get("coverage")).get("target_device_count"),
                _as_dict(report.get("coverage")).get("expected_check_count"),
                _localized_label(
                    report.get("report_status"),
                    _REPORT_STATUS_LABELS,
                ),
            )
        ],
    )
    conclusion = _as_dict(report.get("conclusion"))
    coverage = _as_dict(report.get("coverage"))
    basic_operational_tables = _basic_operational_report_tables(report)
    device_inventory_rows = _report_device_inventory_rows(
        report,
        basic_operational_tables,
    )

    heading("第一章 概述", level=1)
    heading("1.1 巡检概述", level=2)
    paragraph(
        f"{conclusion.get('text')}。{conclusion.get('basis')}",
        bold=True,
    )
    paragraph(f"事实声明编号：{_compact_identifier(conclusion.get('claim_id'))}。")
    paragraph(f"结论可信度：{conclusion.get('trust')}。")
    conclusion_dimensions = _as_dict(report.get("conclusion_dimensions"))
    table(
        ("结论维度", "判定", "依据"),
        [
            (
                label,
                _as_dict(conclusion_dimensions.get(key)).get("label"),
                _as_dict(conclusion_dimensions.get(key)).get("basis"),
            )
            for key, label in (
                ("risk_conclusion", "风险结论"),
                ("assessment_completeness", "评估完整性"),
                ("comparison_readiness", "周期可比性"),
                ("authority", "结论权威性"),
            )
        ],
    )

    heading("1.2 巡检说明", level=2)
    paragraph(
        "本次巡检仅采集和核验设备运行信息，不对设备配置作变更。健康结论以检查项"
        "绑定的标准语义 B、设备或模块上报阈值、有效基线和已发布规则为依据；"
        "数据不足时明确标记为无法判断。"
    )

    heading("1.3 巡检范围", level=2)
    table(
        ("目标设备", "计划检查", "已执行", "已评估", "完整", "部分", "证据不足"),
        [
            (
                coverage.get("target_device_count"),
                coverage.get("expected_check_count"),
                coverage.get("executed_check_count"),
                coverage.get("assessed_check_count"),
                coverage.get("complete_check_count"),
                coverage.get("partial_check_count"),
                coverage.get("insufficient_check_count"),
            )
        ],
    )

    heading("1.3.1 设备清单", level=3)
    table(
        ("序号", "设备名称", "设备型号", "管理 IP"),
        device_inventory_rows,
        empty_text="本批次未形成可列示的设备清单，设备范围以巡检任务清单为准。",
    )

    heading("第二章 巡检内容", level=1)
    heading("2.1 网络设备巡检项", level=2)
    paragraph(
        "本次巡检按已启用检查项和实际采集命令执行，涵盖设备基础信息、硬件环境、"
        "性能、接口状态、配置与日志等内容。具体健康判断以检查项绑定的标准语义 B 和"
        "规则资产为准。"
    )
    paragraph(
        "采集路径韧性表示冻结运行资产中可执行命令相对于检查项声明命令的覆盖情况；"
        "韧性降低不代表检查项健康。"
    )
    table(
        ("设备", "检查项", "路径状态", "可用 / 声明", "不可用路径及原因"),
        [
            (
                item.get("device_name"),
                item.get("check_item_label"),
                _localized_label(
                    item.get("fallback_resilience"),
                    _FALLBACK_RESILIENCE_LABELS,
                ),
                _collection_route_count_label(item),
                _unavailable_command_label(item),
            )
            for item in _as_rows(report.get("matrix"))
        ],
        empty_text="本批次没有可列示的采集路径信息。",
    )

    heading("2.2 巡检设备配置备份", level=2)
    paragraph(
        "由于设备数量较多，配置备份文件较大，不适宜粘贴到 Word 文档内，因此将以"
        "电子版的形式单独提交，特此说明。"
    )

    heading("第三章 巡检结果", level=1)
    heading("3.1 网络设备", level=2)
    paragraph(
        f"{conclusion.get('text')}。{conclusion.get('basis')}",
        bold=True,
    )
    heading("3.1.1 设备基本运行状况", level=3)
    if not basic_operational_tables:
        paragraph("本批次未形成可用于报告的设备版本、CPU、内存、电源、风扇或温度标准观测数据。")
    else:
        paragraph(
            "以下表格展示本次已采集并映射为标准字段的设备基础运行数据；"
            "表中不展示采集时间，缺失字段保留为空。"
        )
        for index, operational_table in enumerate(basic_operational_tables, start=1):
            heading(f"3.1.1.{index} {operational_table['title']}", level=4)
            table(
                tuple(operational_table["headers"]),
                operational_table["rows"],
            )

    heading(f"3.1.2 {risk_heading}", level=3)
    risk_issue_summary = _as_dict(report.get("risk_issue_summary"))
    risk_issues = _as_rows(report.get("risk_issues"))
    paragraph(
        "巡检主问题数按经批准关联后的 RiskEpisode 主问题统计；技术风险、关联症状、"
        "影响实体和事件发生次数分别列示，未审批关联不参与去重。"
    )
    table(
        (
            "巡检主问题",
            "风险Episode",
            "技术风险",
            "关联症状",
            "影响设备",
            "影响实体",
            "事件发生",
            "待复核关联",
        ),
        [
            (
                risk_issue_summary.get("problem_count", len(_as_rows(report.get("risks")))),
                risk_issue_summary.get("risk_episode_count", len(_as_rows(report.get("risks")))),
                risk_issue_summary.get("technical_risk_count", len(_as_rows(report.get("risks")))),
                risk_issue_summary.get("associated_symptom_count", 0),
                risk_issue_summary.get("affected_device_count", 0),
                risk_issue_summary.get("affected_entity_count", 0),
                risk_issue_summary.get(
                    "event_occurrence_count",
                    len(_as_rows(report.get("events"))),
                ),
                risk_issue_summary.get("pending_review_count", 0),
            )
        ],
    )
    table(
        ("等级", "主问题", "检查项", "技术风险", "影响设备 / 实体", "事件", "关联", "声明"),
        [
            (
                _SEVERITY_LABELS.get(
                    str(item.get("severity") or "unknown").lower(),
                    item.get("severity"),
                ),
                item.get("risk_type"),
                "、".join(item.get("check_item_ids") or []),
                item.get("technical_risk_count", item.get("symptom_count")),
                f"{item.get('affected_device_count')} / {item.get('affected_entity_count')}",
                item.get("event_occurrence_count"),
                "已批准" if item.get("correlation_applied") else "独立",
                _compact_identifier(item.get("claim_id")),
            )
            for item in risk_issues
        ],
        empty_text="本报告未形成可列示的巡检主问题。",
    )
    paragraph("以下为全部技术风险 Episode，关联后仍完整保留。")
    table(
        (
            "等级",
            "设备",
            "检查项 / 对象",
            "风险 / 状态",
            "当前值 / 判据",
            "整改 / 责任",
            "声明",
        ),
        [
            (
                _SEVERITY_LABELS.get(
                    str(item.get("severity") or "unknown").lower(),
                    item.get("severity"),
                ),
                item.get("device_name"),
                (f"{item.get('check_item_label') or '-'}\n{item.get('entity_key') or '-'}"),
                (
                    f"{item.get('risk_type') or '-'} / "
                    f"{_localized_label(item.get('status'), _RISK_STATUS_LABELS)}"
                ),
                (f"{_risk_actual_value_label(item)}\n{_risk_criterion_label(item)}"),
                _risk_remediation_owner_label(item),
                _compact_identifier(item.get("claim_id")),
            )
            for item in _as_rows(report.get("risks"))
        ],
        empty_text="本报告未形成可列示的风险结果。",
    )

    heading("3.1.3 历史风险待复核", level=3)
    paragraph(
        "以下风险来自历史批次，本批次未形成足以确认其当前状态的合格证据，不计入本次确认风险。"
    )
    table(
        ("等级", "设备", "检查项", "风险", "对象", "历史状态", "声明"),
        [
            (
                _SEVERITY_LABELS.get(
                    str(item.get("severity") or "unknown").lower(),
                    item.get("severity"),
                ),
                item.get("device_name"),
                item.get("check_item_label"),
                item.get("risk_type"),
                item.get("entity_key"),
                _localized_label(
                    item.get("status"),
                    _RISK_STATUS_LABELS,
                ),
                _compact_identifier(item.get("claim_id")),
            )
            for item in _as_rows(report.get("historical_not_reverified_risks"))
        ],
        empty_text="本批次没有需要列示的历史待复核风险。",
    )

    heading("3.1.4 恢复待确认", level=3)
    paragraph("以下风险在本批次未再次出现，但尚未满足恢复闭环条件。")
    table(
        ("等级", "设备", "检查项", "风险", "对象", "状态", "声明"),
        [
            (
                _SEVERITY_LABELS.get(
                    str(item.get("severity") or "unknown").lower(),
                    item.get("severity"),
                ),
                item.get("device_name"),
                item.get("check_item_label"),
                item.get("risk_type"),
                item.get("entity_key"),
                _localized_label(
                    item.get("status"),
                    _RISK_STATUS_LABELS,
                ),
                _compact_identifier(item.get("claim_id")),
            )
            for item in _as_rows(report.get("recovery_pending_risks"))
        ],
        empty_text="本批次没有需要列示的恢复待确认事项。",
    )
    paragraph(
        "技术恢复、整改验收和正式关闭分别列示；技术恢复不代表整改已验收，也不代表风险已经正式关闭。"
    )
    table(
        ("设备", "检查项", "风险", "对象", "技术状态", "整改", "关闭状态", "声明"),
        [
            (
                item.get("device_name"),
                item.get("check_item_label"),
                item.get("risk_type"),
                item.get("entity_key"),
                _localized_label(
                    item.get("status"),
                    _RISK_STATUS_LABELS,
                ),
                _localized_label(
                    item.get("remediation_status"),
                    _REMEDIATION_STATUS_LABELS,
                ),
                _localized_label(
                    item.get("closure_status"),
                    _RISK_CLOSURE_LABELS,
                ),
                _compact_identifier(item.get("claim_id")),
            )
            for item in _as_rows(report.get("resolved_risks"))
        ],
        empty_text="本批次没有需要列示的技术恢复或正式关闭事项。",
    )

    heading("3.1.5 检查项结论", level=3)
    table(
        (
            "设备",
            "检查项",
            "执行",
            "结论",
            "完整性",
            "证据",
            "覆盖",
            "判断依据",
            "失败层 / 原因",
            "声明",
        ),
        [
            (
                item.get("device_name"),
                item.get("check_item_label"),
                _localized_label(
                    item.get("execution_status"),
                    _EXECUTION_STATUS_LABELS,
                ),
                _STATUS_LABELS.get(
                    str(item.get("assessment_status")),
                    item.get("assessment_status"),
                ),
                _COMPLETENESS_LABELS.get(
                    str(item.get("completeness_status")),
                    item.get("completeness_status"),
                ),
                _evidence_quality_label(item),
                _coverage_label(item),
                _assessment_basis_label(item),
                _reason_codes_label(
                    item.get("blocking_reason_codes"),
                    include_codes=False,
                ),
                _compact_identifier(item.get("claim_id")),
            )
            for item in _as_rows(report.get("matrix"))
        ],
    )

    heading("3.1.6 变化与事件", level=3)
    changes = _as_rows(report.get("changes"))
    events = _as_rows(report.get("events"))
    paragraph(
        f"本次形成标准化变化 {len(changes)} 条、网络事件 "
        f"{len(events)} 条。事件为零不能单独证明设备健康。"
    )
    paragraph(
        _comparison_change_disclosure(
            report,
            change_count=len(changes),
        )
    )
    if changes:
        table(
            ("设备", "检查项 / 对象", "字段", "变化", "判定", "声明"),
            [
                (
                    item.get("device_name"),
                    (f"{item.get('check_item_label') or '-'}\n{item.get('entity_key') or '-'}"),
                    item.get("field"),
                    f"{item.get('old')} -> {item.get('new')}",
                    item.get("decision_reason"),
                    _compact_identifier(item.get("claim_id")),
                )
                for item in changes[:20]
            ],
        )
    if events:
        table(
            (
                "等级",
                "设备",
                "检查项 / 对象",
                "事件 / 当前值 / 判据",
                "发生时间",
                "声明",
            ),
            [
                (
                    _SEVERITY_LABELS.get(
                        str(item.get("severity") or "unknown").lower(),
                        item.get("severity"),
                    ),
                    item.get("device_name"),
                    (f"{item.get('check_item_label') or '-'}\n{item.get('entity_key') or '-'}"),
                    (
                        f"{item.get('event_type') or '-'} / "
                        f"{_localized_label(item.get('status'), _RISK_STATUS_LABELS)}\n"
                        f"{_value_with_unit(item.get('current_value'), item.get('unit'))}\n"
                        f"{item.get('criterion_label') or '-'}"
                    ),
                    item.get("occurred_at"),
                    _compact_identifier(item.get("claim_id")),
                )
                for item in events[:20]
            ],
        )
    paragraph("巡检报告最多展示 20 条变化和 20 条事件；完整明细见技术附表。")

    heading("第四章 巡检总结", level=1)
    heading("4.1 数据可信度与结论边界", level=2)
    blockers = _as_dict(report.get("blocking_reason_counts"))
    if blockers:
        for code, count in sorted(blockers.items()):
            paragraph(f"{_reason_code_label(code, include_code=False)}：{count} 项。")
    else:
        paragraph(
            "未发现影响本次当前结论的数据阻断项。"
            if report.get("current_state_authoritative")
            else ("未发现影响本报告用途内结论的数据阻断项；这不提升其当前网络状态结论权限。")
        )

    heading("第五章 后续工作计划", level=1)
    heading("5.1 后续工作计划", level=2)
    table(
        ("设备", "风险", "状态", "责任人", "期限", "整改动作", "声明"),
        [
            (
                item.get("device_name"),
                item.get("risk_type"),
                _localized_label(
                    item.get("status"),
                    _REMEDIATION_STATUS_LABELS,
                ),
                item.get("owner"),
                item.get("due_at"),
                item.get("action_text"),
                _compact_identifier(item.get("claim_id")),
            )
            for item in _as_rows(report.get("remediations"))
        ],
        empty_text="本批次暂无已登记的整改事项。",
    )

    heading("附录 A：证据与复核说明", level=1)
    table(
        ("复核对象", "标识摘要"),
        (
            (
                "报告清单",
                _compact_identifier(report.get("report_manifest_id")),
            ),
            ("报告事实", _compact_digest(report.get("source_digest"))),
            (
                "目标清单",
                _compact_digest(report.get("target_manifest_digest")),
            ),
            (
                "运行资产清单",
                _compact_digest(report.get("runtime_manifest_digest")),
            ),
            (
                "结论权限证明",
                _compact_digest(_as_dict(report.get("report_authority_context")).get("digest")),
            ),
            (
                "原始证据集合",
                _compact_digest(report.get("raw_evidence_set_digest")),
            ),
            (
                "事实声明账本",
                _compact_digest(report.get("claim_ledger_digest")),
            ),
            (
                "报告投影合同",
                _compact_digest(report.get("report_projection_digest")),
            ),
        ),
    )
    paragraph("完整 SHA-256、证据资质事实和覆盖决策事实见机器可读报告与证据索引包。")
    paragraph(
        "本报告不以“流程完成”或“事件为零”代替健康判断。"
        "数据不完整时明确标记无法判断；原始 CLI 通过独立、"
        "受权限控制的原始输出交付物提供。"
    )
    stream = io.BytesIO()
    doc.save(stream)
    return _normalize_zip_archive(stream.getvalue())


def _safe_excel_value(value: Any) -> Any:
    if value is None or isinstance(value, (bool, int, float)):
        return value
    if isinstance(value, (dict, list, tuple, set)):
        text = json.dumps(
            _jsonable(value),
            ensure_ascii=False,
            sort_keys=True,
            separators=(",", ":"),
        )
    else:
        text = str(value)
    if len(text) > 32_767:
        text = text[:32_764] + "..."
    return text


def _safe_csv_value(value: Any) -> Any:
    value = _safe_excel_value(value)
    if isinstance(value, str) and value.startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


def _sha256_label(value: Any) -> Any:
    text = str(value or "").strip()
    if len(text) == 64 and all(character in "0123456789abcdefABCDEF" for character in text):
        return f"sha256:{text.lower()}"
    return value


def _set_excel_cell_value(cell: Any, value: Any) -> None:
    value = _safe_excel_value(value)
    cell.value = value
    if isinstance(value, str):
        # Keep hashes, identifiers and CLI-like text literal without a visible
        # apostrophe or accidental formula coercion.
        cell.data_type = "s"
        cell.number_format = "@"


def _joined_display_values(value: Any) -> str:
    if isinstance(value, (list, tuple, set)):
        return "、".join(str(item) for item in value if item not in (None, ""))
    return str(value or "")


def _projection_display_value(
    column_id: str,
    value: Any,
) -> Any:
    label_maps = {
        "execution_status": _EXECUTION_STATUS_LABELS,
        "assessment_status": _STATUS_LABELS,
        "completeness_status": _COMPLETENESS_LABELS,
        "comparison_status": _COMPARISON_LABELS,
        "risk_verification_status": _RISK_VERIFICATION_LABELS,
    }
    if column_id in label_maps:
        return _localized_label(value, label_maps[column_id])
    if column_id == "blocking_reason_codes":
        return _reason_codes_label(value)
    if column_id == "runtime_bundle_digest":
        return _sha256_label(value)
    return value


def _excel_text_width(value: Any) -> int:
    text = str(value or "")
    return sum(
        2 if unicodedata.east_asian_width(character) in {"W", "F"} else 1 for character in text
    )


def _append_sheet(
    workbook: Any,
    *,
    name: str,
    headers: Sequence[str],
    rows: Iterable[Sequence[Any]],
    freeze_columns: int = 3,
) -> None:
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    sheet = workbook.create_sheet(name)
    sheet.sheet_view.showGridLines = False
    sheet.sheet_view.zoomScale = 85
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.print_title_rows = "1:1"
    sheet.append(list(headers))
    materialized = list(rows)
    for row_index, row in enumerate(materialized, start=2):
        for column_index, value in enumerate(row, start=1):
            _set_excel_cell_value(
                sheet.cell(row_index, column_index),
                value,
            )
    if not materialized:
        sheet.append(["暂无数据"] + [None] * (len(headers) - 1))
        sheet.merge_cells(
            start_row=2,
            start_column=1,
            end_row=2,
            end_column=len(headers),
        )
        sheet["A2"].font = Font(color="666666", italic=True)
        sheet["A2"].alignment = Alignment(
            horizontal="center",
            vertical="center",
        )
        sheet["A2"].fill = PatternFill("solid", fgColor="F3F5F4")
        sheet.row_dimensions[2].height = 24
    header_fill = PatternFill("solid", fgColor="365F4B")
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )
    sheet.row_dimensions[1].height = 38
    freeze_column = get_column_letter(min(max(freeze_columns, 0), len(headers)) + 1)
    sheet.freeze_panes = f"{freeze_column}2"
    if materialized:
        sheet.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{sheet.max_row}"
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(
                vertical="top",
                wrap_text=True,
            )
    minimum_widths = {
        "设备名称": 18,
        "检查项": 24,
        "检查项名称": 18,
        "风险类型": 22,
        "事件类型": 22,
        "触发判据": 30,
        "整改动作": 30,
    }
    for index, header in enumerate(headers, start=1):
        max_length = max(
            [_excel_text_width(header)]
            + [
                _excel_text_width(sheet.cell(row, index).value)
                for row in range(2, min(sheet.max_row, 200) + 1)
            ]
        )
        sheet.column_dimensions[get_column_letter(index)].width = min(
            max(
                max_length + 2,
                minimum_widths.get(str(header), 12),
            ),
            42,
        )


def _render_audit_check_item_xlsx(report: Mapping[str, Any]) -> bytes:
    """Render a scalable technical annex from the immutable report view."""

    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    workbook = Workbook()
    overview = workbook.active
    overview.title = "总览"
    generated_at = _report_datetime(report.get("generated_at"))
    workbook.properties.creator = "NetOps Agent"
    workbook.properties.created = generated_at
    workbook.properties.modified = generated_at
    overview.sheet_view.showGridLines = False
    overview.sheet_view.zoomScale = 90
    overview.sheet_properties.pageSetUpPr.fitToPage = True
    overview.page_setup.orientation = "portrait"
    overview.page_setup.fitToWidth = 1
    overview.page_setup.fitToHeight = 1
    purpose_label = str(report.get("report_purpose_label") or "未声明用途")
    purpose_disclosure = str(report.get("report_purpose_disclosure") or "")
    current_state_authoritative = bool(report.get("current_state_authoritative"))
    risk_label = "当前风险" if current_state_authoritative else "风险结果（非当前状态）"
    overview["A1"] = f"网络设备巡检技术附表（{purpose_label}）"
    overview["A1"].font = Font(
        color="FFFFFF",
        bold=True,
        size=16,
    )
    overview.merge_cells("A1:B1")
    overview["A1"].alignment = Alignment(
        horizontal="center",
        vertical="center",
    )
    overview["A1"].fill = PatternFill("solid", fgColor="263238")
    overview.row_dimensions[1].height = 28
    overview["A2"] = "巡检批次"
    overview["A2"].font = Font(bold=True)
    overview["B2"] = report.get("patrol_batch_id")
    conclusion_dimensions = _as_dict(report.get("conclusion_dimensions"))
    metrics = [
        ("总体结论", _as_dict(report.get("conclusion")).get("text")),
        ("结论可信度", _as_dict(report.get("conclusion")).get("trust")),
        (
            "风险结论",
            _as_dict(conclusion_dimensions.get("risk_conclusion")).get("label"),
        ),
        (
            "评估完整性",
            _as_dict(conclusion_dimensions.get("assessment_completeness")).get("label"),
        ),
        (
            "周期可比性",
            _as_dict(conclusion_dimensions.get("comparison_readiness")).get("label"),
        ),
        (
            "结论权威性",
            _as_dict(conclusion_dimensions.get("authority")).get("label"),
        ),
        ("目标设备", _as_dict(report.get("coverage")).get("target_device_count")),
        ("计划检查", _as_dict(report.get("coverage")).get("expected_check_count")),
        ("已评估", _as_dict(report.get("coverage")).get("assessed_check_count")),
        (
            "采集路径韧性降低",
            _as_dict(report.get("collection_resilience")).get(
                "reduced_count",
                0,
            ),
        ),
        (
            risk_label,
            _as_dict(report.get("risk_issue_summary")).get(
                "problem_count",
                len(_as_rows(report.get("risks"))),
            ),
        ),
        (
            "风险Episode",
            _as_dict(report.get("risk_issue_summary")).get(
                "risk_episode_count",
                len(_as_rows(report.get("risks"))),
            ),
        ),
        (
            "关联症状",
            _as_dict(report.get("risk_issue_summary")).get(
                "associated_symptom_count",
                0,
            ),
        ),
        (
            "影响实体",
            _as_dict(report.get("risk_issue_summary")).get(
                "affected_entity_count",
                0,
            ),
        ),
        ("变化", len(_as_rows(report.get("changes")))),
        ("事件", len(_as_rows(report.get("events")))),
        ("数据指纹", _sha256_label(report.get("source_digest"))),
        (
            "目标清单指纹",
            _sha256_label(report.get("target_manifest_digest")),
        ),
        (
            "运行资产清单指纹",
            _sha256_label(report.get("runtime_manifest_digest")),
        ),
        (
            "结论权限证明指纹",
            _sha256_label(
                _as_dict(report.get("report_authority_context")).get("digest"),
            ),
        ),
        (
            "历史风险待复核",
            len(_as_rows(report.get("historical_not_reverified_risks"))),
        ),
        (
            "恢复待确认",
            len(_as_rows(report.get("recovery_pending_risks"))),
        ),
        ("报告用途", purpose_label),
        ("用途声明", purpose_disclosure),
        (
            "原始证据集合指纹",
            _sha256_label(report.get("raw_evidence_set_digest")),
        ),
        (
            "事实声明账本指纹",
            _sha256_label(report.get("claim_ledger_digest")),
        ),
        (
            "报告投影合同指纹",
            _sha256_label(report.get("report_projection_digest")),
        ),
        ("事实声明数量", len(_as_rows(report.get("claims")))),
    ]
    for row_index, (label, value) in enumerate(metrics, start=3):
        overview.cell(row_index, 1, label)
        _set_excel_cell_value(
            overview.cell(row_index, 2),
            value,
        )
        overview.cell(row_index, 1).font = Font(bold=True)
    overview.column_dimensions["A"].width = 24
    overview.column_dimensions["B"].width = 70
    overview.freeze_panes = "A3"

    _append_sheet(
        workbook,
        name="检查项结论",
        headers=(
            "设备名称",
            "管理IP",
            "检查项名称",
            "执行状态",
            "评估结论",
            "完整性",
            "可比性",
            "身份",
            "输出质量",
            "时效性",
            "空结果分类",
            "覆盖状态",
            "覆盖基准",
            "风险验证",
            "期望实体数",
            "观测实体数",
            "评估实体数",
            "缺失实体样例",
            "额外实体样例",
            "覆盖判定原因",
            "判断依据来源",
            "判断依据状态",
            "策略快照摘要",
            "判断依据摘要",
            "阻断原因",
            "采集路径韧性",
            "可用命令/声明命令",
            "不可用路径及原因",
            "设备ID",
            "检查项ID",
            "证据摘要",
            "声明编号",
        ),
        rows=(
            (
                item.get("device_name"),
                item.get("management_ip"),
                item.get("check_item_display_name"),
                _localized_label(
                    item.get("execution_status"),
                    _EXECUTION_STATUS_LABELS,
                ),
                _localized_label(
                    item.get("assessment_status"),
                    _STATUS_LABELS,
                ),
                _localized_label(
                    item.get("completeness_status"),
                    _COMPLETENESS_LABELS,
                ),
                _localized_label(
                    item.get("comparison_status"),
                    _COMPARISON_LABELS,
                ),
                _localized_label(
                    item.get("source_identity_status"),
                    _SOURCE_IDENTITY_LABELS,
                ),
                _localized_label(
                    item.get("output_quality_status"),
                    _OUTPUT_QUALITY_LABELS,
                ),
                _localized_label(
                    item.get("freshness_status"),
                    _FRESHNESS_LABELS,
                ),
                _localized_label(
                    item.get("empty_result_classification"),
                    _EMPTY_RESULT_LABELS,
                ),
                _localized_label(
                    item.get("coverage_status"),
                    _COVERAGE_STATUS_LABELS,
                ),
                _localized_label(
                    item.get("coverage_basis"),
                    _COVERAGE_BASIS_LABELS,
                ),
                _localized_label(
                    item.get("risk_verification_status"),
                    _RISK_VERIFICATION_LABELS,
                ),
                item.get("expected_entity_count"),
                item.get("observed_entity_count"),
                item.get("evaluated_entity_count"),
                item.get("missing_entity_sample"),
                item.get("unexpected_entity_sample"),
                _reason_codes_label(
                    item.get("coverage_reason_codes"),
                ),
                _count_map_label(
                    item.get("assessment_basis_source_counts"),
                    _ASSESSMENT_SOURCE_LABELS,
                ),
                _count_map_label(
                    item.get("assessment_basis_status_counts"),
                    _ASSESSMENT_BASIS_STATUS_LABELS,
                ),
                _sha256_label(item.get("assessment_policy_snapshot_digest")),
                _sha256_label(item.get("assessment_basis_digest")),
                _reason_codes_label(
                    item.get("blocking_reason_codes"),
                ),
                _localized_label(
                    item.get("fallback_resilience"),
                    _FALLBACK_RESILIENCE_LABELS,
                ),
                _collection_route_count_label(item),
                _unavailable_command_label(item),
                item.get("device_id"),
                item.get("check_item_id"),
                _sha256_label(item.get("evidence_digest")),
                item.get("claim_id"),
            )
            for item in _as_rows(report.get("matrix"))
        ),
    )
    for table in project_report_tables(report):
        _append_sheet(
            workbook,
            name=table.sheet_name,
            headers=table.headers,
            rows=(
                tuple(
                    _projection_display_value(column_id, value)
                    for column_id, value in zip(
                        table.column_ids,
                        row,
                        strict=True,
                    )
                )
                for row in table.rows
            ),
            freeze_columns=4,
        )
    _append_sheet(
        workbook,
        name="巡检主问题",
        headers=(
            "等级",
            "主问题类型",
            "主RiskCase",
            "主Episode",
            "涉及检查项",
            "技术风险数",
            "关联症状数",
            "影响设备数",
            "影响实体数",
            "事件发生数",
            "关联状态",
            "关联ID",
            "关联证据",
            "声明编号",
        ),
        rows=(
            (
                _localized_label(
                    item.get("severity"),
                    _SEVERITY_LABELS,
                ),
                item.get("risk_type"),
                item.get("primary_risk_case_id"),
                item.get("primary_risk_episode_id"),
                item.get("check_item_ids"),
                item.get("technical_risk_count", item.get("symptom_count")),
                item.get("associated_symptom_count", 0),
                item.get("affected_device_count"),
                item.get("affected_entity_count"),
                item.get("event_occurrence_count"),
                "已批准关联" if item.get("correlation_applied") else "独立问题",
                item.get("association_ids"),
                item.get("association_evidence_refs"),
                item.get("claim_id"),
            )
            for item in _as_rows(report.get("risk_issues"))
        ),
    )
    _append_sheet(
        workbook,
        name=("当前风险" if current_state_authoritative else "风险结果"),
        headers=(
            "等级",
            "设备名称",
            "检查项",
            "风险类型",
            "实体类型",
            "实体",
            "当前值",
            "单位",
            "判据参考值",
            "判据来源",
            "状态",
            "首次发现",
            "最近发现",
            "整改状态",
            "关闭状态",
            "正式关闭时间",
            "关闭人",
            "责任人",
            "期限",
            "CurrentStateRule版本",
            "规则摘要",
            "设备ID",
            "证据",
            "声明编号",
        ),
        rows=(
            (
                _localized_label(
                    item.get("severity"),
                    _SEVERITY_LABELS,
                ),
                item.get("device_name"),
                item.get("check_item_label"),
                item.get("risk_type"),
                item.get("entity_type"),
                item.get("entity_key"),
                _risk_actual_value(item),
                item.get("unit"),
                (
                    item.get("reference_value_number")
                    if item.get("reference_value_number") is not None
                    else item.get("reference_value_text")
                    if item.get("reference_value_text") not in (None, "")
                    else item.get("reference_value")
                ),
                _reference_source_label(item.get("reference_source")),
                _localized_label(
                    item.get("status"),
                    _RISK_STATUS_LABELS,
                ),
                item.get("first_seen_at"),
                item.get("last_seen_at"),
                _localized_label(
                    item.get("remediation_status"),
                    _REMEDIATION_STATUS_LABELS,
                ),
                _localized_label(
                    item.get("closure_status"),
                    _RISK_CLOSURE_LABELS,
                ),
                item.get("closed_at"),
                item.get("closed_by"),
                item.get("owner"),
                item.get("due_at"),
                item.get("event_rule_version_id"),
                _sha256_label(item.get("rule_digest")),
                item.get("device_id"),
                item.get("source_finding_ids_json") or item.get("source_event_ids_json"),
                item.get("claim_id"),
            )
            for item in _as_rows(report.get("risks"))
        ),
    )
    _append_sheet(
        workbook,
        name="历史风险待复核",
        headers=(
            "等级",
            "设备名称",
            "检查项",
            "风险类型",
            "实体类型",
            "实体",
            "历史状态",
            "首次发现",
            "最近发现",
            "最近验证",
            "整改状态",
            "设备ID",
            "声明编号",
        ),
        rows=(
            (
                _localized_label(
                    item.get("severity"),
                    _SEVERITY_LABELS,
                ),
                item.get("device_name"),
                item.get("check_item_label"),
                item.get("risk_type"),
                item.get("entity_type"),
                item.get("entity_key"),
                _localized_label(
                    item.get("status"),
                    _RISK_STATUS_LABELS,
                ),
                item.get("first_seen_at"),
                item.get("last_seen_at"),
                item.get("last_verified_at"),
                _localized_label(
                    item.get("remediation_status"),
                    _REMEDIATION_STATUS_LABELS,
                ),
                item.get("device_id"),
                item.get("claim_id"),
            )
            for item in _as_rows(report.get("historical_not_reverified_risks"))
        ),
    )
    _append_sheet(
        workbook,
        name="恢复待确认",
        headers=(
            "等级",
            "设备名称",
            "检查项",
            "风险类型",
            "实体类型",
            "实体",
            "状态",
            "最近发现",
            "最近验证",
            "整改状态",
            "设备ID",
            "声明编号",
        ),
        rows=(
            (
                _localized_label(
                    item.get("severity"),
                    _SEVERITY_LABELS,
                ),
                item.get("device_name"),
                item.get("check_item_label"),
                item.get("risk_type"),
                item.get("entity_type"),
                item.get("entity_key"),
                _localized_label(
                    item.get("status"),
                    _RISK_STATUS_LABELS,
                ),
                item.get("last_seen_at"),
                item.get("last_verified_at"),
                _localized_label(
                    item.get("remediation_status"),
                    _REMEDIATION_STATUS_LABELS,
                ),
                item.get("device_id"),
                item.get("claim_id"),
            )
            for item in _as_rows(report.get("recovery_pending_risks"))
        ),
    )
    _append_sheet(
        workbook,
        name="技术恢复与关闭",
        headers=(
            "设备名称",
            "检查项",
            "风险类型",
            "实体类型",
            "实体",
            "技术状态",
            "技术恢复确认时间",
            "整改状态",
            "关闭状态",
            "正式关闭时间",
            "关闭人",
            "关闭依据",
            "设备ID",
            "恢复证据",
            "关闭证据",
            "声明编号",
        ),
        rows=(
            (
                item.get("device_name"),
                item.get("check_item_label"),
                item.get("risk_type"),
                item.get("entity_type"),
                item.get("entity_key"),
                _localized_label(
                    item.get("status"),
                    _RISK_STATUS_LABELS,
                ),
                item.get("resolved_at"),
                _localized_label(
                    item.get("remediation_status"),
                    _REMEDIATION_STATUS_LABELS,
                ),
                _localized_label(
                    item.get("closure_status"),
                    _RISK_CLOSURE_LABELS,
                ),
                item.get("closed_at"),
                item.get("closed_by"),
                item.get("closure_reason"),
                item.get("device_id"),
                item.get("resolution_evidence_refs_json"),
                item.get("closure_evidence_refs_json"),
                item.get("claim_id"),
            )
            for item in _as_rows(report.get("resolved_risks"))
        ),
    )
    _append_sheet(
        workbook,
        name="整改跟踪",
        headers=(
            "设备名称",
            "风险类型",
            "等级",
            "整改状态",
            "责任人",
            "期限",
            "整改动作",
            "创建人",
            "完成人",
            "完成时间",
            "验收人",
            "验收时间",
            "验收意见",
            "设备ID",
            "完成证据",
            "验收证据",
            "声明编号",
        ),
        rows=(
            (
                item.get("device_name"),
                item.get("risk_type"),
                _localized_label(
                    item.get("severity"),
                    _SEVERITY_LABELS,
                ),
                _localized_label(
                    item.get("status"),
                    _REMEDIATION_STATUS_LABELS,
                ),
                item.get("owner"),
                item.get("due_at"),
                item.get("action_text"),
                item.get("created_by"),
                item.get("completed_by"),
                item.get("completed_at"),
                item.get("accepted_by"),
                item.get("accepted_at"),
                item.get("acceptance_reason"),
                item.get("device_id"),
                item.get("evidence_refs_json"),
                item.get("acceptance_evidence_refs_json"),
                item.get("claim_id"),
            )
            for item in _as_rows(report.get("remediations"))
        ),
    )
    _append_sheet(
        workbook,
        name="事件",
        headers=(
            "时间",
            "等级",
            "设备名称",
            "检查项",
            "事件类型",
            "状态",
            "实体类型",
            "实体",
            "字段",
            "原值",
            "当前值",
            "单位",
            "触发判据",
            "阈值字段",
            "阈值",
            "阈值来源",
            "迟滞",
            "恢复边界",
            "EventRule",
            "EventRule版本",
            "事件ID",
            "执行ID",
            "声明编号",
        ),
        rows=(
            (
                item.get("occurred_at"),
                _localized_label(
                    item.get("severity"),
                    _SEVERITY_LABELS,
                ),
                item.get("device_name"),
                item.get("check_item_label"),
                item.get("event_type"),
                _localized_label(
                    item.get("status"),
                    _RISK_STATUS_LABELS,
                ),
                item.get("entity_type"),
                item.get("entity_key"),
                item.get("field"),
                item.get("previous_value"),
                item.get("current_value"),
                item.get("unit"),
                item.get("criterion_label"),
                item.get("reference_field"),
                item.get("reference_value"),
                _reference_source_label(item.get("reference_source")),
                item.get("hysteresis"),
                item.get("recovery_boundary"),
                item.get("event_rule_id"),
                item.get("event_rule_version"),
                item.get("id"),
                item.get("execution_id"),
                item.get("claim_id"),
            )
            for item in _as_rows(report.get("events"))
        ),
        freeze_columns=4,
    )
    _append_sheet(
        workbook,
        name="变化",
        headers=(
            "时间",
            "设备名称",
            "检查项",
            "标准观测类型",
            "变化类型",
            "实体类型",
            "实体",
            "字段",
            "原值",
            "现值",
            "变化ID",
            "执行ID",
            "声明编号",
        ),
        rows=(
            (
                item.get("observed_at"),
                item.get("device_name"),
                item.get("check_item_label"),
                item.get("observation_type"),
                item.get("change_type"),
                item.get("entity_type"),
                item.get("entity_key"),
                item.get("field"),
                item.get("old"),
                item.get("new"),
                item.get("id"),
                item.get("execution_id"),
                item.get("claim_id"),
            )
            for item in _as_rows(report.get("changes"))
        ),
        freeze_columns=3,
    )
    _append_sheet(
        workbook,
        name="当前观测",
        headers=(
            "设备名称",
            "管理IP",
            "检查项",
            "标准观测类型",
            "时间",
            "实体类型",
            "实体",
            "字段",
            "值",
            "单位",
            "维度",
            "设备ID",
            "执行ID",
            "观测ID",
            "快照ID",
            "Parser模板",
            "Parser版本",
            "Contract版本",
        ),
        rows=(
            (
                item.get("device_name"),
                item.get("management_ip"),
                _joined_display_values(item.get("check_item_labels")),
                item.get("observation_type"),
                item.get("observed_at"),
                item.get("entity_type"),
                item.get("entity_key"),
                item.get("field"),
                _display_value(item),
                item.get("unit"),
                item.get("dimensions"),
                item.get("device_id"),
                item.get("execution_ids"),
                item.get("id"),
                item.get("snapshot_id"),
                item.get("parser_template_id"),
                item.get("parser_version"),
                item.get("contract_version"),
            )
            for item in _as_rows(report.get("observations"))
        ),
        freeze_columns=4,
    )
    _append_sheet(
        workbook,
        name="事实声明",
        headers=(
            "声明编号",
            "声明代码",
            "范围",
            "断言级别",
            "允许表述",
            "声明对象",
            "事实引用",
            "证据引用",
            "限制条件",
            "权威上下文SHA256",
            "事实基准SHA256",
            "声明SHA256",
        ),
        rows=(
            (
                item.get("claim_id"),
                item.get("claim_code"),
                item.get("claim_scope"),
                item.get("assertion_level"),
                item.get("allowed_phrase_code"),
                item.get("subject"),
                item.get("fact_refs"),
                item.get("evidence_refs"),
                item.get("limitation_refs"),
                _sha256_label(item.get("authority_context_digest")),
                _sha256_label(item.get("source_facts_digest")),
                _sha256_label(item.get("claim_digest")),
            )
            for item in _as_rows(report.get("claims"))
        ),
        freeze_columns=5,
    )
    _append_sheet(
        workbook,
        name="证据索引",
        headers=(
            "设备名称",
            "管理IP",
            "厂商",
            "型号",
            "检查项",
            "时间",
            "命令",
            "输出分类",
            "解析状态",
            "Parser版本",
            "Contract版本",
            "设备ID",
            "快照ID",
            "运行ID",
            "Parser模板",
            "采集动作",
            "原文SHA256",
        ),
        rows=(
            (
                item.get("device_name"),
                item.get("management_ip"),
                item.get("vendor"),
                item.get("model"),
                _joined_display_values(item.get("check_item_labels")),
                item.get("observed_at"),
                item.get("canonical_command"),
                _localized_label(
                    item.get("output_status"),
                    _COLLECTION_OUTPUT_STATUS_LABELS,
                ),
                _localized_label(
                    item.get("parse_status"),
                    _PARSE_STATUS_LABELS,
                ),
                item.get("parser_version"),
                item.get("contract_version"),
                item.get("device_id"),
                item.get("snapshot_id"),
                item.get("run_id"),
                item.get("parser_template_id"),
                item.get("collection_action_id"),
                _sha256_label(item.get("raw_text_hash")),
            )
            for item in _as_rows(report.get("evidence_index"))
        ),
        freeze_columns=5,
    )
    stream = io.BytesIO()
    workbook.save(stream)
    return _normalize_zip_archive(stream.getvalue())


def render_check_item_xlsx(report: Mapping[str, Any]) -> bytes:
    """Render the immutable technical annex used by report delivery.

    Preserve frozen claims, coverage, authority, risk and evidence ledgers in
    every delivered workbook.
    """

    return _render_audit_check_item_xlsx(report)


def _csv_bytes(
    headers: Sequence[str],
    rows: Iterable[Sequence[Any]],
) -> bytes:
    stream = io.StringIO(newline="")
    writer = csv.writer(stream, lineterminator="\n")
    writer.writerow(headers)
    for row in rows:
        writer.writerow([_safe_csv_value(value) for value in row])
    return ("\ufeff" + stream.getvalue()).encode("utf-8")


def render_check_item_evidence_zip(
    manifest: Mapping[str, Any],
    report: Mapping[str, Any],
) -> bytes:
    """Package the evidence index without embedding protected raw CLI."""

    validate_report_manifest(manifest)
    files: dict[str, bytes] = {
        "README.txt": (
            "网络设备巡检证据包\n"
            f"巡检批次: {report.get('patrol_batch_id')}\n"
            f"报告清单: {report.get('report_manifest_id')}\n"
            f"事实指纹: {report.get('source_digest')}\n\n"
            f"事实声明账本: {report.get('claim_ledger_digest') or '-'}\n"
            f"报告投影合同: {report.get('report_projection_digest') or '-'}\n"
            "本证据包包含冻结事实、检查项结论和命令快照索引，"
            "不包含原始 CLI 正文。原始输出需通过受权限控制的"
            "“原始输出”交付物获取，并以 raw_text_hash 校验一致性。\n"
        ).encode(),
        "report-manifest.json": json.dumps(
            _jsonable(dict(manifest)),
            ensure_ascii=False,
            sort_keys=True,
            indent=2,
        ).encode("utf-8"),
        "check-assessments.csv": _csv_bytes(
            (
                "device_id",
                "device_name",
                "check_item_id",
                "execution_status",
                "assessment_status",
                "completeness_status",
                "comparison_status",
                "blocking_reason_codes",
                "assessment_policy_snapshot_digest",
                "assessment_basis_digest",
                "assessment_basis_source_counts",
                "assessment_basis_status_counts",
                "evidence_digest",
                "claim_id",
            ),
            (
                (
                    item.get("device_id"),
                    item.get("device_name"),
                    item.get("check_item_id"),
                    item.get("execution_status"),
                    item.get("assessment_status"),
                    item.get("completeness_status"),
                    item.get("comparison_status"),
                    item.get("blocking_reason_codes"),
                    item.get("assessment_policy_snapshot_digest"),
                    item.get("assessment_basis_digest"),
                    item.get("assessment_basis_source_counts"),
                    item.get("assessment_basis_status_counts"),
                    item.get("evidence_digest"),
                    item.get("claim_id"),
                )
                for item in _as_rows(report.get("matrix"))
            ),
        ),
        "evidence-index.csv": _csv_bytes(
            (
                "snapshot_id",
                "run_id",
                "device_id",
                "device_name",
                "observed_at",
                "canonical_command",
                "raw_text_hash",
                "parse_status",
                "output_status",
                "parser_template_id",
                "parser_version",
                "contract_version",
            ),
            (
                (
                    item.get("snapshot_id"),
                    item.get("run_id"),
                    item.get("device_id"),
                    item.get("device_name"),
                    item.get("observed_at"),
                    item.get("canonical_command"),
                    item.get("raw_text_hash"),
                    item.get("parse_status"),
                    item.get("output_status"),
                    item.get("parser_template_id"),
                    item.get("parser_version"),
                    item.get("contract_version"),
                )
                for item in _as_rows(report.get("evidence_index"))
            ),
        ),
        "risk-lifecycle.csv": _csv_bytes(
            (
                "risk_case_id",
                "occurrence_kind",
                "observed_at",
                "status_after",
                "actor",
                "reason",
                "evidence_refs",
                "claim_id",
            ),
            (
                (
                    item.get("risk_case_id"),
                    item.get("occurrence_kind"),
                    item.get("observed_at"),
                    item.get("status_after"),
                    item.get("actor"),
                    item.get("reason"),
                    item.get("evidence_refs_json"),
                    item.get("claim_id"),
                )
                for item in _as_rows(report.get("risk_occurrences"))
            ),
        ),
        "report-claims.json": json.dumps(
            _jsonable(
                {
                    "claim_ledger_digest": report.get("claim_ledger_digest"),
                    "claims": _as_rows(report.get("claims")),
                }
            ),
            ensure_ascii=False,
            sort_keys=True,
            indent=2,
        ).encode("utf-8"),
        "report-claims.csv": _csv_bytes(
            (
                "claim_id",
                "claim_code",
                "claim_scope",
                "assertion_level",
                "allowed_phrase_code",
                "fact_refs",
                "evidence_refs",
                "limitation_refs",
                "claim_digest",
            ),
            (
                (
                    item.get("claim_id"),
                    item.get("claim_code"),
                    item.get("claim_scope"),
                    item.get("assertion_level"),
                    item.get("allowed_phrase_code"),
                    item.get("fact_refs"),
                    item.get("evidence_refs"),
                    item.get("limitation_refs"),
                    item.get("claim_digest"),
                )
                for item in _as_rows(report.get("claims"))
            ),
        ),
    }
    output = io.BytesIO()
    with zipfile.ZipFile(
        output,
        "w",
        compression=zipfile.ZIP_DEFLATED,
        compresslevel=9,
    ) as archive:
        for name, content in sorted(files.items()):
            info = zipfile.ZipInfo(name, _ARCHIVE_TIMESTAMP)
            info.compress_type = zipfile.ZIP_DEFLATED
            archive.writestr(info, content)
    return output.getvalue()


def render_check_item_json(report: Mapping[str, Any]) -> bytes:
    return json.dumps(
        _jsonable(dict(report)),
        ensure_ascii=False,
        sort_keys=True,
        indent=2,
    ).encode("utf-8")


def render_check_item_artifacts(
    manifest: Mapping[str, Any],
) -> dict[str, ReportArtifact]:
    """Render all final deliverables from one source-fact manifest."""

    report = build_check_item_report_view(manifest)
    batch_id = str(report.get("patrol_batch_id") or "patrol")
    markdown = render_check_item_markdown(report).encode("utf-8")
    artifacts = {
        "markdown": ReportArtifact(
            name="markdown",
            filename=f"{batch_id}-report.md",
            content_type="text/markdown; charset=utf-8",
            content=markdown,
        ),
        "docx": ReportArtifact(
            name="docx",
            filename=f"{batch_id}-leadership-report.docx",
            content_type=(
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
            ),
            content=render_check_item_docx(report),
        ),
        "excel": ReportArtifact(
            name="excel",
            filename=f"{batch_id}-technical-report.xlsx",
            content_type=("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
            content=render_check_item_xlsx(report),
        ),
        "json": ReportArtifact(
            name="json",
            filename=f"{batch_id}-report.json",
            content_type="application/json; charset=utf-8",
            content=render_check_item_json(report),
        ),
        "evidence": ReportArtifact(
            name="evidence",
            filename=f"{batch_id}-evidence.zip",
            content_type="application/zip",
            content=render_check_item_evidence_zip(manifest, report),
        ),
    }
    return artifacts


def render_check_item_artifact(
    manifest: Mapping[str, Any],
    name: str,
) -> ReportArtifact:
    """Render one requested format without producing unrelated artifacts."""

    report = build_check_item_report_view(manifest)
    batch_id = str(report.get("patrol_batch_id") or "patrol")
    factories = {
        "markdown": lambda: ReportArtifact(
            name="markdown",
            filename=f"{batch_id}-report.md",
            content_type="text/markdown; charset=utf-8",
            content=render_check_item_markdown(report).encode("utf-8"),
        ),
        "docx": lambda: ReportArtifact(
            name="docx",
            filename=f"{batch_id}-leadership-report.docx",
            content_type=(
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
            ),
            content=render_check_item_docx(report),
        ),
        "excel": lambda: ReportArtifact(
            name="excel",
            filename=f"{batch_id}-technical-report.xlsx",
            content_type=("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
            content=render_check_item_xlsx(report),
        ),
        "json": lambda: ReportArtifact(
            name="json",
            filename=f"{batch_id}-report.json",
            content_type="application/json; charset=utf-8",
            content=render_check_item_json(report),
        ),
        "evidence": lambda: ReportArtifact(
            name="evidence",
            filename=f"{batch_id}-evidence.zip",
            content_type="application/zip",
            content=render_check_item_evidence_zip(manifest, report),
        ),
    }
    try:
        return factories[str(name)]()
    except KeyError as exc:
        raise ValueError(f"unsupported report artifact: {name}") from exc


__all__ = [
    "CheckItemReportMaterializer",
    "ReportArtifact",
    "build_check_item_report_view",
    "load_attached_report_artifact",
    "render_check_item_artifacts",
    "render_check_item_artifact",
    "render_check_item_docx",
    "render_check_item_evidence_zip",
    "render_check_item_json",
    "render_check_item_markdown",
    "render_check_item_xlsx",
    "validate_report_manifest",
    "verify_attached_report_artifacts",
]
